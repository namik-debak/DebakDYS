"""
Y03 F06 Excel → ProcessKPI tanım + dönemsel ölçüm aktarımı.
"""

from __future__ import annotations

import os
import re
from datetime import date, datetime

from openpyxl import load_workbook

from kpi_engine import parse_hedef, evaluate_kpi, _to_float

DEFAULT_KPI_XLSX = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    "static", "branding", "y03-f06-kpi-2026.xlsx",
)


def _sheet_surec_kod(sheet_name: str) -> str:
    m = re.match(r"^([YMD]\d{2})", (sheet_name or "").strip(), re.I)
    return (m.group(1) if m else sheet_name[:3]).upper()


def _is_header_row(values):
    joined = " ".join(str(v) for v in values if v is not None).lower()
    return "performans" in joined and "parametre" in joined


def _is_aksiyon_section(values):
    joined = " ".join(str(v) for v in values if v is not None).upper()
    return "DÜZELTİCİ" in joined or "DUZELTICI" in joined


def _parse_year(v):
    if isinstance(v, (datetime, date)):
        return v.year
    if isinstance(v, bool):
        return None
    if isinstance(v, int) and 2000 <= v <= 2100:
        return v
    if isinstance(v, float) and 2000 <= v <= 2100:
        return int(v)
    if isinstance(v, str):
        m = re.match(r"^\s*(20\d{2})\s*$", v)
        if m:
            return int(m.group(1))
    return None


def _build_period_columns(ws, header_r):
    """
    Yıl / ay başlıklarından (kolon, yıl, ay|None) listesi üretir.
    Ay None ise yıllık sütun.
    """
    max_c = ws.max_column or 6
    years_at = {}
    for c in range(7, max_c + 1):
        y = _parse_year(ws.cell(header_r, c).value)
        if y:
            years_at[c] = y

    months_at = {}
    for c in range(7, max_c + 1):
        v = ws.cell(header_r + 1, c).value
        if isinstance(v, bool):
            continue
        if isinstance(v, (int, float)) and 1 <= int(v) <= 12:
            months_at[c] = int(v)

    # Eksik ara yıl başlığı (ör. 2024 ile 2026 arası 2025 sütunu)
    if years_at:
        sorted_cols = sorted(years_at.keys())
        for i in range(len(sorted_cols) - 1):
            c1, c2 = sorted_cols[i], sorted_cols[i + 1]
            y1, y2 = years_at[c1], years_at[c2]
            gap_cols = [c for c in range(c1 + 1, c2) if c not in months_at]
            expected = list(range(y1 + 1, y2))
            if gap_cols and len(gap_cols) == len(expected):
                for gc, ey in zip(gap_cols, expected):
                    years_at[gc] = ey

    cols = []
    cur_y = None
    for c in range(7, max_c + 1):
        if c in years_at:
            cur_y = years_at[c]
        if cur_y is None:
            continue
        cols.append((c, cur_y, months_at.get(c)))
    return cols


def donem_from_cell(year, month, gg_periyot):
    """Hücre konumunu dönem koduna çevirir (2026, 2026-H1, 2026-Q1, 2026-01)."""
    if month is None:
        return str(year)
    p = (gg_periyot or "").strip().upper().replace(" ", "")
    if p in ("1A", "1AY", "A", "AYLIK"):
        return f"{year}-{month:02d}"
    if p in ("3A", "3AY", "Q"):
        return f"{year}-Q{(month - 1) // 3 + 1}"
    if p in ("6A", "6AY", "H"):
        return f"{year}-H1" if month <= 6 else f"{year}-H2"
    if p in ("1Y", "Y", "YILLIK"):
        return str(year)
    return f"{year}-{month:02d}"


def _cell_measurement(raw):
    """Sayısal ölçüm değerini döner; metin notlarını atlar."""
    if raw is None or isinstance(raw, bool):
        return None
    if isinstance(raw, (int, float)):
        return float(raw)
    if isinstance(raw, str):
        s = raw.strip().replace(" ", "").replace(",", ".")
        if not s or not re.match(r"^-?\d+(\.\d+)?%?$", s):
            return None
        return _to_float(s)
    return None


def parse_kpi_workbook(path=None):
    """
    Excel'den KPI tanım + ölçüm listesi döner.
    Her kayıt: ... + olcumler: [{donem, gerceklesen}, ...]
    """
    path = path or DEFAULT_KPI_XLSX
    if not os.path.isfile(path):
        raise FileNotFoundError(f"KPI Excel bulunamadı: {path}")

    wb = load_workbook(path, data_only=True)
    rows_out = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        kod = _sheet_surec_kod(sheet_name)
        header_r = None
        last_prosedure = ""
        period_cols = []

        for r in range(1, (ws.max_row or 0) + 1):
            vals = [ws.cell(r, c).value for c in range(1, 8)]
            if not any(v is not None and str(v).strip() for v in vals):
                continue
            if _is_aksiyon_section(vals):
                break
            if header_r is None:
                if _is_header_row(vals):
                    header_r = r
                    period_cols = _build_period_columns(ws, header_r)
                continue
            # Ay numarası satırı
            if r == header_r + 1 and vals[0] is None and vals[2] is None:
                continue
            if vals[0] is None and vals[2] is None:
                continue

            parametre = vals[2]
            if parametre is None or not str(parametre).strip():
                continue

            sira = vals[0]
            prosedure = vals[1]
            if prosedure and str(prosedure).strip():
                last_prosedure = str(prosedure).strip()
            else:
                prosedure = last_prosedure
            try:
                sira_no = int(sira) if sira is not None and str(sira).strip().isdigit() else None
            except (TypeError, ValueError):
                sira_no = None

            hedef_raw = vals[4]
            birim = (str(vals[5]).strip() if vals[5] is not None else None) or None
            gg = (str(vals[3]).strip() if vals[3] is not None else None) or None
            kural = parse_hedef(hedef_raw, birim)

            # Dönemsel ölçümler (aynı dönem birden fazla hücredeyse son geçerli kazanır)
            olcum_map = {}
            for c, year, month in period_cols:
                val = _cell_measurement(ws.cell(r, c).value)
                if val is None:
                    continue
                donem = donem_from_cell(year, month, gg)
                olcum_map[donem] = val

            rows_out.append({
                "surec_kod": kod,
                "sira_no": sira_no,
                "prosedure": (str(prosedure).strip() if prosedure else None) or None,
                "parametre": str(parametre).strip(),
                "gg_periyot": gg,
                "hedef_metin": kural["metin"],
                "hedef_tip": kural["tip"],
                "hedef_deger": kural["deger"],
                "hedef_deger_ust": kural["deger_ust"],
                "birim": birim,
                "olcumler": [
                    {"donem": d, "gerceklesen": v} for d, v in sorted(olcum_map.items())
                ],
            })
    return rows_out


def import_kpis_to_db(db, path=None, replace=False, import_measurements=True):
    """
    ProcessKPI (+ isteğe bağlı ölçüm) kayıtlarını Excel'den yükler.
    Dönüş: (kpi_eklenen, kpi_atlanan, olcum_eklenen, olcum_guncellenen)
    Tarihsel ölçüm aktarımında DÖF açılmaz.
    """
    from models import Process, ProcessKPI, ProcessKPIMeasurement

    items = parse_kpi_workbook(path)
    procs = {p.kod: p for p in db.query(Process).all()}
    if replace:
        for k in db.query(ProcessKPI).filter_by(aktif=True).all():
            k.aktif = False

    existing = {
        (k.surec_kod, (k.parametre or "").strip().lower()): k
        for k in db.query(ProcessKPI).filter_by(aktif=True).all()
    }
    added = 0
    skipped = 0
    for item in items:
        key = (item["surec_kod"], item["parametre"].lower())
        if key in existing and not replace:
            skipped += 1
            continue
        proc = procs.get(item["surec_kod"])
        kpi = ProcessKPI(
            surec_id=proc.id if proc else None,
            surec_kod=item["surec_kod"],
            sira_no=item["sira_no"],
            prosedure=item["prosedure"],
            parametre=item["parametre"],
            gg_periyot=item["gg_periyot"],
            hedef_metin=item["hedef_metin"],
            hedef_tip=item["hedef_tip"],
            hedef_deger=item["hedef_deger"],
            hedef_deger_ust=item["hedef_deger_ust"],
            birim=item["birim"],
            aktif=True,
        )
        db.add(kpi)
        db.flush()
        existing[key] = kpi
        added += 1

    meas_added = 0
    meas_updated = 0
    if import_measurements:
        # Tüm aktif KPI'ları yeniden map'le
        kpi_map = {
            (k.surec_kod, (k.parametre or "").strip().lower()): k
            for k in db.query(ProcessKPI).filter_by(aktif=True).all()
        }
        for item in items:
            kpi = kpi_map.get((item["surec_kod"], item["parametre"].lower()))
            if not kpi:
                continue
            for om in item.get("olcumler") or []:
                donem = om["donem"]
                gercek = om["gerceklesen"]
                durum = evaluate_kpi(
                    kpi.hedef_tip, kpi.hedef_deger, kpi.hedef_deger_ust, gercek, kpi.birim
                )
                m = (
                    db.query(ProcessKPIMeasurement)
                    .filter_by(kpi_id=kpi.id, donem=donem)
                    .first()
                )
                if not m:
                    m = ProcessKPIMeasurement(kpi_id=kpi.id, donem=donem)
                    db.add(m)
                    meas_added += 1
                else:
                    meas_updated += 1
                m.gerceklesen = gercek
                m.gerceklesen_metin = None
                m.durum = durum
                if not m.notlar:
                    m.notlar = "Y03 F06 Excel aktarımı"
        db.flush()

    db.commit()
    return added, skipped, meas_added, meas_updated
