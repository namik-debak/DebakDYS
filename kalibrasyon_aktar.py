"""
DYS — D04 F06 İzleme & Ölçme Cihazları ve Kalibrasyon Listesi aktarımı (2026)
  excel_aktar(db, yol)            listeyi CalibrationEquipment kayıtlarına aktarır / günceller (tekrar çalıştırılabilir)
  sertifika_klasoru_esle(db, k)   klasördeki PDF sertifikaları dosya adındaki sertifika / seri numarasına göre bağlar
Cihaz kodu benzersiz değilse (GO, NO GO …) ekipman_no = 'kod · seri no' (ya da 'kod #n'); listedeki kod
'cihaz_kodu' alanında aynen tutulur. Eşleştirme anahtarı: kod + seri no + ad. Listede olmayan cihaz silinmez.
CLI:  python kalibrasyon_aktar.py --excel "<D04 F06 ….xlsm>" [--sertifika "<klasör>"] [--uygula]
"""
import argparse
import json
import os
import re
import shutil
from datetime import date, datetime

AY = {"yıl": 12, "yil": 12, "ay": 1}


def _s(v):
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, datetime):
        return v.strftime("%d.%m.%Y")
    t = re.sub(r"\s+", " ", str(v)).strip()
    return t or None


def _tarih(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    t = _s(v) or ""
    for f in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(t, f).date()
        except ValueError:
            continue
    return None


def _periyot_ay(normal, mak):
    """'2 yıl' / '6 ay' → ay; yalnız sayı → yıl kabul edilir; '_' → None."""
    for t in ((_s(normal) or "").lower(), (_s(mak) or "").lower()):
        m = re.match(r"(\d+(?:[.,]\d+)?)\s*(yıl|yil|ay)?", t)
        if m:
            return int(round(float(m.group(1).replace(",", ".")) * AY.get(m.group(2) or "yıl", 12)))
    return None


def _norm(t):
    t = str(t or "").replace("İ", "i").replace("I", "ı").lower().replace("ı", "i")
    return re.sub(r"[^0-9a-zçğöşü]", "", t)


def _num(t):
    """Sertifika numarasının firma adı dışındaki kısmı: 'Egemet-0005K-1225-14869' → '0005k122514869'."""
    return re.sub(r"^[a-zçğöşü]+", "", _norm(t))


ALANLAR = {  # alan → başlık önekleri (küçük harf)
    "tur": ("cihazın türü",), "tip": ("ölçüm tipi",), "kod": ("cihaz kodu",),
    "tanim": ("izleme ve ölçme cihazının tanımı",), "ad": ("izleme ve ölçme cihazının adı",), "aktif": ("aktif/pasif",),
    "yer": ("kullanım yeri",), "sert": ("sertifika no",), "seri": ("seri no",), "imalatci": ("imalatçı",),
    "bolum": ("kullanıcı bölüm",), "hassas": ("ölçme hassas",), "aralik": ("ölçüm aralığı",), "birim": ("birimi",),
    "ozel": ("özel kar",), "karalik": ("kullanım aralığı",), "per": ("normal kalib",), "mper": ("mak. kalib",),
    "son": ("kalib. tarihi",), "sonraki": ("gelecek kalibrasyon",), "dper": ("normal doğrulama",),
    "uygun": ("kalibrasyon sonucu uygunluk",), "hata": ("izin verilen toplam",), "sapma": ("kalibrasyon sonucu (sapma",),
    "dyon": ("doğrulama yöntemi",), "msa": ("msa tipi",), "gerekce": ("karar gerekçesi",), "onay": ("kullanıma onay",),
}


def satirlari_oku(yol, sayfa=None):
    import openpyxl
    import warnings
    warnings.filterwarnings("ignore")
    wb = openpyxl.load_workbook(yol, data_only=True)
    ws = wb[sayfa] if sayfa else next((w for w in wb.worksheets if "kalibrasyon" in w.title.lower() or w.title.lower().startswith("liste")),
                                      wb.worksheets[0])
    kucuk = lambda v: (_s(v) or "").replace("İ", "i").replace("I", "ı").lower()   # 'İzleme'.lower() → 'i̇zleme' olmasın
    hr = None
    for r in range(1, 15):
        vals = [kucuk(ws.cell(r, c).value) for c in range(1, 60)]
        if "cihaz kodu" in vals:
            hr = r
            break
    if not hr:
        raise ValueError("Başlık satırı ('Cihaz Kodu') bulunamadı")
    K = {a: next((i + 1 for i, v in enumerate(vals) if v and any(v.startswith(o) for o in on)), None) for a, on in ALANLAR.items()}
    K["karar"] = next((i + 1 for i, v in enumerate(vals) if v == "karar"), None)
    if not K.get("mper") and K.get("karalik") and K["karalik"] < len(vals) and not vals[K["karalik"]]:
        K["mper"] = K["karalik"] + 1   # başlıksız makine kalibrasyon periyodu sütunu (yıl)
    son_baslik = max(i + 1 for i, v in enumerate(vals) if v)
    satirlar = []
    for r in range(hr + 1, ws.max_row + 1):
        g = lambda k: ws.cell(r, K[k]).value if K.get(k) else None
        if not any(_s(g(k)) for k in ("kod", "tanim", "ad", "seri")):
            continue
        notlar = [_s(ws.cell(r, c).value) for c in range(son_baslik + 1, son_baslik + 4) if _s(ws.cell(r, c).value)]
        satirlar.append({"satir": r, **{k: g(k) for k in list(ALANLAR) + ["karar"]}, "not": "; ".join(notlar) or None})
    return satirlar, ws.title


def _ekipman_no(kod, seri, kullanilan):
    taban = _s(kod) or "F06"
    aday = taban[:30]
    if aday in kullanilan and seri:
        aday = f"{taban} · {seri}"[:30]
    n = 2
    while aday in kullanilan:
        aday = f"{taban[:24]} #{n}"
        n += 1
    return aday


def excel_aktar(db, yol, sayfa=None, bugun=None):
    """Dönüş: {'yeni', 'guncellenen', 'toplam', 'ortak_kod', 'listede_yok', 'uyari', 'sayfa'}."""
    from models import CalibrationEquipment, User
    bugun = bugun or date.today()
    satirlar, sayfa_adi = satirlari_oku(yol, sayfa)
    mevcut = db.query(CalibrationEquipment).all()
    anahtar = lambda kod, seri, ad: (_norm(kod), _norm(seri), _norm(ad))
    # listeden gelen kayıtta kod 'cihaz_kodu' (kodsuz cihazda boş); elle açılmış kayıtta ekipman_no
    by_key = {anahtar(e.cihaz_kodu if (e.kaynak or "").startswith("D04 F06") else e.ekipman_no, e.seri_no,
                      e.ek.get("liste_ad") or e.ad): e for e in mevcut}
    kullanilan = {e.ekipman_no for e in mevcut}
    kullanicilar = {_norm(u.ad_soyad): u.id for u in db.query(User).all()}
    R = {"yeni": 0, "guncellenen": 0, "toplam": len(satirlar), "ortak_kod": 0, "listede_yok": 0, "uyari": [], "sayfa": sayfa_adi}
    gorulen, islenen = set(), set()
    for x in satirlar:
        kod, seri = _s(x["kod"]), _s(x["seri"])
        ad = _s(x["ad"]) or _s(x["tanim"]) or kod or "Cihaz"
        k = anahtar(kod, seri, ad)          # eşleştirme anahtarı listedeki ad ile (tekrar aktarımda aynı kalır)
        if _norm(ad) == _norm(kod) and _s(x["tanim"]):
            ad = f"{_s(x['tanim'])}"         # 'Adı' sütununa kod yazılmışsa görünen ad = tanım (ör. MI-05 → Dış Çap Mikrometresi)
        if k in gorulen:
            R["uyari"].append(f"satır {x['satir']}: aynı kod + seri no + ad tekrar ediyor ({kod} / {seri}) — atlandı")
            continue
        gorulen.add(k)
        e = by_key.get(k)
        if not e:
            eno = _ekipman_no(kod, seri, kullanilan)
            if eno != (kod or "F06")[:30]:
                R["ortak_kod"] += 1
            e = CalibrationEquipment(ekipman_no=eno, ad=ad[:200])
            db.add(e)
            kullanilan.add(eno)
            by_key[k] = e
            R["yeni"] += 1
        else:
            R["guncellenen"] += 1
        islenen.add(id(e))
        kisalt = lambda a, n: ((_s(x[a]) or "")[:n] or None)
        pasif = (_s(x["aktif"]) or "").upper().startswith("PAS")
        karar = " — ".join(t for t in (_s(x["karar"]), _s(x["gerekce"])) if t)
        ariza = bool(re.search(r"arızal|kırık|kırıldı|hurda", karar.lower())) or (_s(x["karar"]) or "").upper() == "X"
        son, sonraki = _tarih(x["son"]), _tarih(x["sonraki"])
        not_ = x.get("not") or ""
        if pasif:
            durum = "Pasif"
        elif ariza:
            durum = "Arızalı"
        elif "kalibrasyona" in not_.lower():
            durum = "Kalibrasyonda"
        elif sonraki and sonraki < bugun:
            durum = "Süresi Doldu"
        else:
            durum = "Geçerli"
        ozel = (_s(x["ozel"]) or "").lower()
        e.ad = ad[:200]
        e.cihaz_kodu = (kod or "")[:60] or None
        e.tip = kisalt("tanim", 100)
        e.cihaz_turu = kisalt("tur", 60)
        e.olcum_tipi = kisalt("tip", 40)
        e.konum = kisalt("yer", 150)
        e.seri_no = (seri or "")[:100] or None
        e.imalatci = kisalt("imalatci", 100) or e.imalatci
        e.kullanici_bolum = kisalt("bolum", 100) or e.kullanici_bolum
        e.hassasiyet = kisalt("hassas", 60) or e.hassasiyet
        e.olcum_araligi = kisalt("aralik", 100) or e.olcum_araligi
        e.kullanim_araligi = kisalt("karalik", 100)
        e.birim = kisalt("birim", 30) or e.birim
        if ozel:
            e.ozel_karakteristik = ozel.startswith("evet")
        e.kalibrasyon_araligi_ay = _periyot_ay(x["per"], x["mper"])
        e.son_kalibrasyon, e.sonraki_kalibrasyon = son, sonraki
        e.sertifika_no = kisalt("sert", 120)
        e.kalibrasyon_sapma = kisalt("sapma", 150)
        e.izin_verilen_hata = kisalt("hata", 150)
        e.dogrulama_periyodu = kisalt("dper", 40)
        e.dogrulama_yontemi = kisalt("dyon", 100)
        e.msa_tipi = kisalt("msa", 40) or e.msa_tipi
        e.karar = karar[:200] or None
        e.kullanima_onay_veren = kisalt("onay", 100) or e.kullanima_onay_veren
        if e.kullanima_onay_veren and not e.sorumlu_id:
            e.sorumlu_id = kullanicilar.get(_norm(e.kullanima_onay_veren))
        e.durum = durum
        ek = {k2: v for k2, v in {"uygunluk": _s(x["uygun"]), "not": not_ or None,
                                  "liste_ad": _s(x["ad"]) or _s(x["tanim"]) or kod}.items() if v}
        e.ek_json = json.dumps(ek, ensure_ascii=False) if ek else None
        e.kaynak = f"D04 F06 '{sayfa_adi}' satır {x['satir']}"[:120]
    R["listede_yok"] = len([e for e in mevcut if e.kaynak and e.kaynak.startswith("D04 F06") and id(e) not in islenen])
    db.commit()
    return R


def _dosya_anahtarlari(ad):
    govde = os.path.splitext(ad)[0]
    tokenler = set()
    for p in (p for p in re.split(r"[\s_]+", govde) if p):
        # '19-41013486' → '41013486' (sıra no öneki); 'Mikrometre-76191904', 'B18347869-07439' → tire ile ayrılan parçalar
        for t in (p, re.sub(r"^\d+(?:,\d+)?-", "", p), *p.split("-")):
            n = _norm(t)
            if len(n) >= 4:
                tokenler.update({n, n.lstrip("0")})
    return _norm(govde), tokenler


def sertifika_adaylari(db, ad):
    """Dosya adına göre aday cihazlar (önce sertifika no, sonra seri no)."""
    from models import CalibrationEquipment
    ekip = db.query(CalibrationEquipment).all()
    tam, tok = _dosya_anahtarlari(ad)
    guclu = [e for e in ekip if len(_num(e.sertifika_no)) >= 8 and _num(e.sertifika_no) in tam]
    return guclu or [e for e in ekip if e.seri_no and len(_norm(e.seri_no).lstrip("0")) >= 4 and _norm(e.seri_no).lstrip("0") in tok]


def sertifika_ekle(db, e, kaynak_yol, ad=None, yukleyen_id=None, aciklama=None, **alanlar):
    """Sertifika dosyasını DYS yükleme klasörüne kopyalar, kaydı açar, önceki sertifikaları arşivler (commit etmez)."""
    from config import Config
    from models import KalibrasyonSertifika
    ad = ad or os.path.basename(kaynak_yol)
    d = os.path.join(Config.UPLOAD_FOLDER, "kalibrasyon", str(e.id))
    os.makedirs(d, exist_ok=True)
    hedef = os.path.join(d, ad)
    if os.path.exists(hedef) and os.path.abspath(hedef) != os.path.abspath(kaynak_yol):
        kok, uz = os.path.splitext(ad)
        hedef = os.path.join(d, f"{kok}_{datetime.now():%Y%m%d%H%M%S}{uz}")
    if os.path.abspath(hedef) != os.path.abspath(kaynak_yol):
        shutil.copy2(kaynak_yol, hedef)
    for c in e.sertifikalar:
        c.gecerli = False
    s = KalibrasyonSertifika(ekipman_id=e.id, dosya_adi=os.path.basename(hedef), dosya_yolu=hedef, gecerli=True,
                             yukleyen_id=yukleyen_id, aciklama=aciklama, **alanlar)
    db.add(s)
    return s


def sertifika_klasoru_esle(db, klasor, yukleyen_id=None):
    """PDF'leri cihazlara bağlar; yalnız TEK cihaza eşleşen dosya bağlanır. Kaynak klasör değişmez."""
    from models import KalibrasyonSertifika
    R = {"dosya": 0, "baglanan": 0, "zaten": 0, "eslesmeyen": [], "cok_eslesen": []}
    mevcut = {(c.ekipman_id, (c.aciklama or "").split(" | ")[-1]) for c in db.query(KalibrasyonSertifika).all()}
    for ad in sorted(os.listdir(klasor)):
        if not ad.lower().endswith(".pdf"):
            continue
        R["dosya"] += 1
        adaylar = sertifika_adaylari(db, ad)
        if len(adaylar) != 1:
            if adaylar:
                R["cok_eslesen"].append(f"{ad} → {', '.join(e.ekipman_no for e in adaylar[:5])}")
            else:
                R["eslesmeyen"].append(ad)
            continue
        e = adaylar[0]
        if (e.id, ad) in mevcut:
            R["zaten"] += 1
            continue
        sertifika_ekle(db, e, os.path.join(klasor, ad), yukleyen_id=yukleyen_id,
                       aciklama=f"'{os.path.basename(klasor)}' klasöründen otomatik eşleştirildi | {ad}",
                       sertifika_no=e.sertifika_no, kalibrasyon_tarihi=e.son_kalibrasyon,
                       sonraki_kalibrasyon=e.sonraki_kalibrasyon, sapma=e.kalibrasyon_sapma, sonuc="Uygun")
        db.flush()
        mevcut.add((e.id, ad))
        R["baglanan"] += 1
    db.commit()
    return R


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--excel", required=True)
    ap.add_argument("--sayfa")
    ap.add_argument("--sertifika")
    ap.add_argument("--uygula", action="store_true")
    a = ap.parse_args()
    from models import SessionLocal, init_db
    init_db()
    db = SessionLocal()
    try:
        if not a.uygula:
            s, sayfa = satirlari_oku(a.excel, a.sayfa)
            print(f"DENEME TURU: '{sayfa}' sayfasında {len(s)} cihaz satırı okundu. Aktarmak için --uygula ekleyin.")
        else:
            print(json.dumps(excel_aktar(db, a.excel, a.sayfa), ensure_ascii=False, default=str)[:3000])
            if a.sertifika:
                print(json.dumps(sertifika_klasoru_esle(db, a.sertifika), ensure_ascii=False)[:3000])
    finally:
        db.close()
