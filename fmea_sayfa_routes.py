"""
DYS — FMEA çalışma sayfaları (2026): D-FMEA / P-FMEA / FMEA-MSR, GSI-RD-370 Rev. C02 kuralları.
  /iatf/fmea/sayfalar                      sayfa listesi (tür / arama), yeni sayfa
  /iatf/fmea/sayfa/<id>                    çalışma sayfası: satır tablosu, otomatik AP, puan dayanakları, uyarılar, İlk / Güncel AP sayaçları
  /iatf/fmea/sayfa/<id>/satir[/<rid>]      satır ekle / güncelle (JSON, otomatik kayıt) / sil
  /iatf/fmea/sayfa/<id>/excel              Excel; /iatf/fmea/sayfa/<id>/yazdir  A4 yatay çıktı
Kurallar fmea_core.py'de (saf fonksiyonlar). AP yalnız matristen; RPN yalnız bilgi. Puanlar yönlendirme amaçlıdır; nihai karar FMEA ekibinindir.
Yetki: görüntüleme herkes; satır / sayfa yazma Sadece Görüntüleme hariç; sayfa silme Admin / Doküman Kontrol.
"""
import io
from datetime import date, datetime

from flask import render_template, request, redirect, url_for, flash, abort, session, jsonify, send_file

from app_runtime import host as _host
_h = _host()
app = _h.app
login_required = _h.login_required
get_db = _h.get_db
log_action = _h.log_action

from models import FmeaSayfa, FmeaSatir  # noqa: E402
import fmea_core as FC  # noqa: E402

METIN = ("item", "fn", "eff", "mode", "cause", "prev", "det", "act")
PUAN = ("S", "O", "D", "O2", "D2")
KISA = {"resp": 150, "date": 60, "st": 30}
ALANLAR = METIN + PUAN + tuple(KISA)


def _yazabilir():
    return session.get("rol") != "Sadece Görüntüleme"


def _yonetici():
    return session.get("rol") in ("Admin", "Doküman Kontrol")


def _sira_no(db):
    prefix = f"FS-{date.today().year}-"
    son = 0
    for (n,) in db.query(FmeaSayfa.sayfa_no).filter(FmeaSayfa.sayfa_no.like(prefix + "%")).all():
        try:
            son = max(son, int(n.rsplit("-", 1)[1]))
        except (ValueError, IndexError):
            pass
    return f"{prefix}{son + 1:03d}"


def _dict(s):
    return {k: getattr(s, k) for k in ALANLAR}


def _satir_json(s, tip):
    d = _dict(s)
    h = FC.hesapla(d, tip)
    return {"id": s.id, **{k: ("" if v is None else v) for k, v in d.items()}, "ap": h["ap"], "ap_ad": FC.AP_AD[h["ap"]], "ap2": h["ap2"],
            "ap2_ad": FC.AP_AD[h["ap2"]], "rpn": h["rpn"] if h["rpn"] is not None else "", "uyari": h["uyari"]}


def _ozet(db, sayfa):
    return FC.ozet([_dict(s) for s in db.query(FmeaSatir).filter_by(sayfa_id=sayfa.id).all()], sayfa.tip)


# ─────────────────────────── liste / yeni ───────────────────────────
@app.route("/iatf/fmea/sayfalar", methods=["GET", "POST"])
@login_required
def fmea_sayfalar():
    db = get_db()
    try:
        if request.method == "POST":
            if not _yazabilir():
                abort(403)
            f = request.form
            tip = f.get("tip")
            baslik = (f.get("baslik") or "").strip()
            if tip not in FC.TIPLER or not baslik:
                flash("Tür ve başlık zorunludur.", "error")
                return redirect(url_for("fmea_sayfalar"))
            try:
                tarih = datetime.strptime(f.get("tarih"), "%Y-%m-%d").date() if f.get("tarih") else None
            except ValueError:
                tarih = None
            s = FmeaSayfa(sayfa_no=_sira_no(db), tip=tip, baslik=baslik[:200], parca_proses=(f.get("parca_proses") or "").strip()[:200] or None,
                          musteri=(f.get("musteri") or "").strip()[:150] or None, ekip=(f.get("ekip") or "").strip()[:300] or None,
                          revizyon=(f.get("revizyon") or "").strip()[:10] or "00", tarih=tarih or date.today(),
                          aciklama=(f.get("aciklama") or "").strip()[:1000] or None, olusturan_id=session.get("user_id"))
            db.add(s); db.commit()
            log_action(db, "Oluşturma", detay=f"FMEA çalışma sayfası {s.sayfa_no}: {FC.TIP_AD[tip]} — {s.baslik}")
            return redirect(url_for("fmea_sayfa", sid=s.id))
        tur, q = request.args.get("tur") or "", (request.args.get("q") or "").strip().lower()
        L = db.query(FmeaSayfa).order_by(FmeaSayfa.guncelleme_tarihi.desc()).all()
        L = [s for s in L if (not tur or s.tip == tur) and (not q or q in f"{s.sayfa_no} {s.baslik} {s.parca_proses} {s.musteri}".lower())]
        satirlar = {}
        for r in db.query(FmeaSatir).filter(FmeaSatir.sayfa_id.in_([s.id for s in L] or [0])).all():
            satirlar.setdefault(r.sayfa_id, []).append(_dict(r))
        liste = [{"s": s, "n": len(satirlar.get(s.id, [])), "ozet": FC.ozet(satirlar.get(s.id, []), s.tip)} for s in L]
        return render_template("fmea_sayfalar.html", liste=liste, tur=tur, q=request.args.get("q", ""), tip_ad=FC.TIP_AD,
                               yazabilir=_yazabilir(), today=date.today().isoformat())
    finally:
        db.close()


# ─────────────────────────── çalışma sayfası ───────────────────────────
def _sayfa_veri(db, sid):
    s = db.get(FmeaSayfa, sid)
    if not s:
        abort(404)
    satirlar = db.query(FmeaSatir).filter_by(sayfa_id=s.id).order_by(FmeaSatir.sira, FmeaSatir.id).all()
    return s, satirlar


@app.route("/iatf/fmea/sayfa/<int:sid>", methods=["GET", "POST"])
@login_required
def fmea_sayfa(sid):
    db = get_db()
    try:
        s, satirlar = _sayfa_veri(db, sid)
        if request.method == "POST":
            if not _yazabilir():
                abort(403)
            a = request.form.get("action")
            if a == "bilgi":
                f = request.form
                if (f.get("baslik") or "").strip():
                    s.baslik = f.get("baslik").strip()[:200]
                s.parca_proses = (f.get("parca_proses") or "").strip()[:200] or None
                s.musteri = (f.get("musteri") or "").strip()[:150] or None
                s.ekip = (f.get("ekip") or "").strip()[:300] or None
                s.revizyon = (f.get("revizyon") or "").strip()[:10] or "00"
                try:
                    s.tarih = datetime.strptime(f.get("tarih"), "%Y-%m-%d").date() if f.get("tarih") else s.tarih
                except ValueError:
                    pass
                db.commit(); log_action(db, "Güncelleme", detay=f"FMEA sayfa bilgisi: {s.sayfa_no}")
                flash("Sayfa bilgileri kaydedildi.", "success")
            elif a == "sil":
                if not _yonetici():
                    abort(403)
                no = s.sayfa_no
                for r in satirlar:
                    db.delete(r)
                db.delete(s); db.commit(); log_action(db, "Silme", detay=f"FMEA çalışma sayfası silindi: {no}")
                flash(f"{no} silindi.", "success")
                return redirect(url_for("fmea_sayfalar"))
            return redirect(url_for("fmea_sayfa", sid=s.id))
        K = FC.kriterler()
        tip = s.tip
        kri = K["kriterler"].get(tip, {})
        return render_template("fmea_sayfa.html", s=s, satirlar=[_satir_json(r, tip) for r in satirlar], ozet=FC.ozet([_dict(r) for r in satirlar], tip),
                               tipbilgi=K["tipler"].get(tip, {}), kriterler=kri, kriter_notu=K["notlar"].get(tip, ""), ipuclari=K["ipuclari"].get(tip, []),
                               tip_ad=FC.TIP_AD, durumlar=FC.DURUMLAR[tip], yazabilir=_yazabilir(), yonetici=_yonetici(),
                               puanlar={"S": FC.S_PUANLAR, "O": FC.FM_PUANLAR if tip == "M" else FC.O_PUANLAR, "D": FC.FM_PUANLAR if tip == "M" else FC.D_PUANLAR},
                               kriter_var=bool(kri), today=date.today())
    finally:
        db.close()


@app.route("/iatf/fmea/sayfa/<int:sid>/satir", methods=["POST"])
@login_required
def fmea_satir_ekle(sid):
    if not _yazabilir():
        abort(403)
    db = get_db()
    try:
        s, satirlar = _sayfa_veri(db, sid)
        r = FmeaSatir(sayfa_id=s.id, sira=(max((x.sira or 0) for x in satirlar) + 1) if satirlar else 1)
        # önceki satırın ortak alanlarını kopyala (aynı hata modu için yeni neden satırı): isteğe bağlı
        if request.form.get("kopya") and satirlar:
            ref = db.get(FmeaSatir, request.form.get("kopya", type=int) or 0)
            if ref and ref.sayfa_id == s.id:
                for k in ("item", "fn", "eff", "S", "mode"):
                    setattr(r, k, getattr(ref, k))
        db.add(r); s.guncelleme_tarihi = datetime.now(); db.commit()
        return jsonify(ok=True, satir=_satir_json(r, s.tip), ozet=_ozet(db, s))
    finally:
        db.close()


@app.route("/iatf/fmea/sayfa/<int:sid>/satir/<int:rid>", methods=["POST"])
@login_required
def fmea_satir_guncelle(sid, rid):
    if not _yazabilir():
        abort(403)
    db = get_db()
    try:
        s, _ = _sayfa_veri(db, sid)
        r = db.get(FmeaSatir, rid)
        if not r or r.sayfa_id != s.id:
            abort(404)
        if request.form.get("action") == "sil":
            db.delete(r); s.guncelleme_tarihi = datetime.now(); db.commit()
            return jsonify(ok=True, silindi=True, ozet=_ozet(db, s))
        alan, deger = request.form.get("alan"), (request.form.get("deger") or "").strip()
        if alan not in ALANLAR:
            abort(400)
        if alan in PUAN:
            if not FC.dogrula_puan(s.tip, alan[0], deger):
                return jsonify(ok=False, hata="Bu puan seçilemez."), 400
            setattr(r, alan, int(deger) if deger else None)
        elif alan == "st":
            if deger and deger not in FC.DURUMLAR[s.tip]:
                return jsonify(ok=False, hata="Geçersiz durum."), 400
            r.st = deger or None
        else:
            setattr(r, alan, (deger[:KISA[alan]] if alan in KISA else deger[:4000]) or None)
        s.guncelleme_tarihi = datetime.now(); db.commit()
        return jsonify(ok=True, satir=_satir_json(r, s.tip), ozet=_ozet(db, s))
    finally:
        db.close()


# ─────────────────────────── çıktılar ───────────────────────────
@app.route("/iatf/fmea/sayfa/<int:sid>/yazdir")
@login_required
def fmea_sayfa_yazdir(sid):
    db = get_db()
    try:
        s, satirlar = _sayfa_veri(db, sid)
        K = FC.kriterler()
        return render_template("fmea_sayfa_yazdir.html", s=s, satirlar=[_satir_json(r, s.tip) for r in satirlar], tipbilgi=K["tipler"].get(s.tip, {}),
                               ozet=FC.ozet([_dict(r) for r in satirlar], s.tip), tip_ad=FC.TIP_AD, simdi=datetime.now())
    finally:
        db.close()


@app.route("/iatf/fmea/sayfa/<int:sid>/excel")
@login_required
def fmea_sayfa_excel(sid):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    db = get_db()
    try:
        s, satirlar = _sayfa_veri(db, sid)
        T = FC.kriterler()["tipler"].get(s.tip, {}) or {"o": "O", "d": "D", "ol": "O", "dl": "D", "prev": "Önleyici aksiyon", "det": "Tespit aksiyonu",
                                                          "item": "Eleman", "fn": "Fonksiyon"}
        wb = Workbook(); ws = wb.active; ws.title = FC.TIP_AD[s.tip]
        ws.append([f"{s.sayfa_no} — {FC.TIP_AD[s.tip]} — {s.baslik}"]); ws["A1"].font = Font(bold=True, size=13)
        ws.append([f"Parça / proses: {s.parca_proses or ''}", f"Müşteri: {s.musteri or ''}", f"Ekip: {s.ekip or ''}", f"Rev: {s.revizyon or ''}",
                   f"Tarih: {s.tarih.strftime('%d.%m.%Y') if s.tarih else ''}"])
        bas = ["#", T["item"], T["fn"], "Hata etkisi", "S", "Hata modu", "Hata nedeni", T["prev"], T["ol"], T["det"], T["dl"], "AP", "RPN (bilgi)",
               "Optimizasyon aksiyonu", "Sorumlu", "Hedef tarih", "Durum", T["o"] + "′", T["d"] + "′", "AP′"]
        ws.append(bas)
        for c in ws[3]:
            c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F4E78"); c.alignment = Alignment(wrap_text=True, vertical="center")
        renk = {"H": "F4B6B6", "M": "FFE699", "L": "C6E0B4"}
        for i, r in enumerate(satirlar, 1):
            j = _satir_json(r, s.tip)
            ws.append([i, j["item"], j["fn"], j["eff"], j["S"], j["mode"], j["cause"], j["prev"], j["O"], j["det"], j["D"], j["ap_ad"], j["rpn"], j["act"],
                       j["resp"], j["date"], j["st"], j["O2"], j["D2"], j["ap2_ad"]])
            for col, ap in ((12, j["ap"]), (20, j["ap2"])):
                if ap:
                    ws.cell(ws.max_row, col).fill = PatternFill("solid", fgColor=renk[ap])
        for i, w in enumerate([4, 22, 26, 26, 5, 22, 24, 26, 6, 26, 6, 9, 9, 28, 16, 14, 14, 6, 6, 9], 1):
            ws.column_dimensions[ws.cell(3, i).column_letter].width = w
        for row in ws.iter_rows(min_row=4):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "C4"; ws.auto_filter.ref = f"A3:T{ws.max_row}"
        ws.append([]); ws.append(["Derecelendirmeler yönlendirme amaçlıdır; nihai karar FMEA ekibine aittir. Kaynak: GSI-RD-370 Rev. C02. AP kararı matristen; RPN yalnız bilgidir."])
        bio = io.BytesIO(); wb.save(bio); bio.seek(0)
        log_action(db, "Görüntüleme", detay=f"FMEA sayfası Excel: {s.sayfa_no}")
        return send_file(bio, as_attachment=True, download_name=f"{s.sayfa_no} {FC.TIP_AD[s.tip]} {s.baslik}.xlsx"[:150],
                         mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    finally:
        db.close()
