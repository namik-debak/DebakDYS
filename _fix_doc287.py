# -*- coding: utf-8 -*-
"""Fix D01 T03 footer to company format: D01 T03 14.02.2023"""
import re
import shutil
import zipfile
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

PATH = Path(
    r"D:\UYGULAMA PROJELER\debak dijital\dys\uploads\documents\287\rev0\D01 T03 El Yıkama Talimatı.docx"
)
NET = Path(
    r"\\192.168.0.249\ortak\HARUN\dys\uploads\documents\287\rev0\D01 T03 El Yıkama Talimatı.docx"
)
TARGET_LINE = "D01 T03 14.02.2023"


def xml_all_text(docx_path):
    out = []
    with zipfile.ZipFile(docx_path) as z:
        for name in z.namelist():
            if not name.endswith(".xml"):
                continue
            data = z.read(name).decode("utf-8", errors="ignore")
            for m in re.finditer(r"<w:t[^>]*>([^<]*)</w:t>", data):
                if m.group(1).strip():
                    out.append((name, m.group(1)))
    return out


def set_footer_compact(path: Path):
    doc = Document(str(path))
    for sec in doc.sections:
        footer = sec.footer
        # Mevcut alt bilgi paragrafını temizle, orijinal formatta yaz
        if footer.paragraphs:
            p = footer.paragraphs[0]
            p.clear()
        else:
            p = footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = p.add_run(TARGET_LINE)
        run.font.size = Pt(9)
    doc.save(str(path))


def replace_in_body(path: Path):
    doc = Document(str(path))
    changed = False
    pattern = re.compile(r"D2\s*T44\s*\(\s*1\s*\)[_\s]*14\.02\.2023", re.I)
    alt = re.compile(r"D2\s*T44[^\d]*14\.02\.2023", re.I)

    def fix_paragraphs(paragraphs):
        nonlocal changed
        for p in paragraphs:
            txt = p.text
            if not txt:
                continue
            new = pattern.sub(TARGET_LINE, txt)
            new = alt.sub(TARGET_LINE, new)
            if "D2 T44" in new or "D2 T44" in txt:
                new = re.sub(r"D2\s*T44.*", TARGET_LINE, txt)
            if new != txt:
                p.text = new
                changed = True

    fix_paragraphs(doc.paragraphs)
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                fix_paragraphs(cell.paragraphs)

    for sec in doc.sections:
        for part in (sec.header, sec.footer):
            fix_paragraphs(part.paragraphs)

    if changed:
        doc.save(str(path))
    return changed


if __name__ == "__main__":
    # yedek al
    bak2 = PATH.with_suffix(".docx.bak2")
    if not bak2.exists():
        shutil.copy2(PATH, bak2)

    print("BEFORE fix:")
    for name, t in xml_all_text(PATH):
        if any(x in t for x in ("D2", "T44", "D01", "T03", "14.02", "Dok")):
            print(f"  {name}: {t!r}")

    set_footer_compact(PATH)
    replace_in_body(PATH)

    print("\nAFTER fix:")
    for name, t in xml_all_text(PATH):
        if any(x in t for x in ("D2", "T44", "D01", "T03", "14.02", "Dok")):
            print(f"  {name}: {t!r}")

    # ağ kopyası farklıysa senkronize et
    if NET.exists() and NET.resolve() != PATH.resolve():
        shutil.copy2(PATH, NET)
        print("\nNetwork copy updated.")
    elif NET.exists():
        print("\nLocal and network path appear same or network synced.")

    print(f"\nFooter set to: {TARGET_LINE}")
