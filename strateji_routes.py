"""
DYS — Stratejik Planlama modülü (2026) — Y01 F02 Stratejik Planlama Tabloları + Y01 F03 İSG ve Çevre Stratejik Planlama
Tabloları yerine: Çevre (bağlam) analizi → Paydaş analizi → SWOT → Stratejiler → Öncelikli stratejiler (uygulama ve
gözden geçirme). Her yıl bir dönemdir; kalemler 'Kurumsal' ya da 'İSG & Çevre' kapsamındadır.
  /strateji                         ?yil=2026&sekme=baglam|paydas|swot|strateji|eylem&kapsam=
  /strateji/kalem/<id>              kalem düzenleme / silme (POST)
  /strateji/kalem/yeni              yeni kalem (POST)
  /strateji/donem                   dönem gözden geçirme tarihleri / notlar (POST); yeni yıl = önceki yıldan kopya
  /strateji/export.xlsx             yılın tüm tabloları Excel olarak
  /strateji/yazdir                  yönetim gözden geçirmesi için çıktı
  /strateji/aktar                   'DYS Aktarım/Strateji/*.json' içe aktarma (Admin)
app.py sonunda import edilir.
"""
import io
import json
from datetime import date, datetime

from flask import render_template, request, redirect, url_for, flash, abort, session, send_file

from app_runtime import host as _host
_h = _host()
app = _h.app
login_required = _h.login_required
get_db = _h.get_db
log_action = _h.log_action

from models import (StratejiDonem, StratejiKalem, User, STRATEJI_BOLUMLERI, STRATEJI_KAPSAMLARI,  # noqa: E402
                    STRATEJI_BAGLAM_BILESENLERI, STRATEJI_SWOT_TURLERI, STRATEJI_EYLEM_DURUMLARI)

SEKMELER = [("baglam", "1 · Çevre (bağlam) analizi"), ("paydas", "2 · Paydaş analizi"), ("swot", "3 · SWOT analizi"),
            ("strateji", "4 · Stratejiler"), ("eylem", "5 · Öncelikli stratejiler ve gözden geçirme")]
ALANLAR = {   # bölüm → (form alanı, etiket, tip) — veri_json'a yazılanlar
    "baglam": [("kaynak", "Kurumsal bilginin kaynağı", "metin"), ("analiz", "İlk analiz / değişiklikler", "uzun")],
    "paydas": [("statu", "Statü", "metin"), ("etki", "Etki (Direkt / Dolaylı)", "secim:Direkt,Dolaylı"),
               ("faaliyet_etkisi", "Kuruluş faaliyetlerine etkisi", "uzun"), ("beklenti", "İhtiyaç ve beklentiler", "uzun"),
               ("isg_cevre_beklenti", "İSG ve çevre ihtiyaç ve beklentileri", "uzun")],
    "swot": [],
    "strateji": [("ilgili_paydas", "İlgili paydaş", "metin"), ("swot", "Dayandığı SWOT maddeleri (ör. Z1, T2)", "metin")],
    "eylem": [("kbf", "Kritik başarı faktörü", "uzun"), ("kaynak", "Kaynak", "metin"), ("gg_notu", "Gözden geçirme notu", "uzun"),
              ("gg_tarihi", "Gözden geçirme tarihi", "tarih"), ("gerceklesen", "Gerçekleşen tarih", "tarih")],
}


def _can_edit():
    return session.get("rol") in ("Admin", "Doküman Kontrol")


def _yillar(db):
    return sorted({y for (y,) in db.query(StratejiDonem.yil).all()} | {y for (y,) in db.query(StratejiKalem.yil).distinct().all()}, reverse=True)


def _kalemler(db, yil, bolum, kapsam=None):
    q = db.query(StratejiKalem).filter(StratejiKalem.yil == yil, StratejiKalem.bolum == bolum)
    if kapsam:
        q = q.filter(StratejiKalem.kapsam == kapsam)
    return q.order_by(StratejiKalem.kapsam, StratejiKalem.grup, StratejiKalem.sira, StratejiKalem.id).all()


def _tarih(s):
    try:
        return datetime.strptime(s, "%Y-%m-%d").date() if s else None
    except ValueError:
        return None


@app.route("/strateji")
@login_required
def strateji():
    db = get_db()
    try:
        yillar = _yillar(db) or [date.today().year]
        yil = request.args.get("yil", type=int) or (date.today().year if date.today().year in yillar else yillar[0])
        sekme = request.args.get("sekme") if request.args.get("sekme") in STRATEJI_BOLUMLERI else "baglam"
        kapsam = request.args.get("kapsam") if request.args.get("kapsam") in STRATEJI_KAPSAMLARI else None
        donem = db.query(StratejiDonem).filter(StratejiDonem.yil == yil).first()
        say = {b: db.query(StratejiKalem).filter(StratejiKalem.yil == yil, StratejiKalem.bolum == b).count() for b in STRATEJI_BOLUMLERI}
        kalemler = _kalemler(db, yil, sekme, kapsam)
        gruplar = {}
        if sekme in ("baglam", "swot"):
            sira = STRATEJI_BAGLAM_BILESENLERI if sekme == "baglam" else STRATEJI_SWOT_TURLERI
            for k in kalemler:
                gruplar.setdefault(k.grup or "Diğer", []).append(k)
            gruplar = dict(sorted(gruplar.items(), key=lambda kv: sira.index(kv[0]) if kv[0] in sira else 99))
        eylem_ozet = None
        if sekme == "eylem":
            bugun = date.today()
            eylem_ozet = {d: sum(1 for k in kalemler if (k.durum or "Planlandı") == d) for d in STRATEJI_EYLEM_DURUMLARI}
            eylem_ozet["geciken"] = sum(1 for k in kalemler if k.termin and k.termin < bugun and k.durum not in ("Tamamlandı", "İptal"))
        return render_template("strateji.html", yil=yil, yillar=yillar, sekme=sekme, sekmeler=SEKMELER, kapsam=kapsam,
                               kapsamlar=STRATEJI_KAPSAMLARI, donem=donem, say=say, kalemler=kalemler, gruplar=gruplar,
                               alanlar=ALANLAR[sekme], bilesenler=STRATEJI_BAGLAM_BILESENLERI, swot_turleri=STRATEJI_SWOT_TURLERI,
                               durumlar=STRATEJI_EYLEM_DURUMLARI, eylem_ozet=eylem_ozet, can_edit=_can_edit(), bugun=date.today(),
                               users=db.query(User).filter_by(aktif=True).order_by(User.ad_soyad).all() if _can_edit() else [])
    finally:
        db.close()


def _formdan(k, form):
    k.kapsam = form.get("kapsam") if form.get("kapsam") in STRATEJI_KAPSAMLARI else (k.kapsam or "Kurumsal")
    k.grup = (form.get("grup") or "").strip() or None
    k.kod = (form.get("kod") or "").strip() or None
    k.sira = form.get("sira", type=int) or k.sira
    k.metin = (form.get("metin") or "").strip()
    veri = k.veri
    for ad, _, tip in ALANLAR.get(k.bolum, []):
        if ad in form:
            veri[ad] = (form.get(ad) or "").strip()
    k.veri_json = json.dumps(veri, ensure_ascii=False)
    if k.bolum == "eylem":
        k.termin = _tarih(form.get("termin"))
        k.durum = form.get("durum") if form.get("durum") in STRATEJI_EYLEM_DURUMLARI else (k.durum or "Planlandı")
        k.sorumlu_id = form.get("sorumlu_id", type=int) or None


@app.route("/strateji/kalem/yeni", methods=["POST"])
@login_required
def strateji_kalem_yeni():
    if not _can_edit():
        abort(403)
    bolum = request.form.get("bolum")
    if bolum not in STRATEJI_BOLUMLERI:
        abort(400)
    db = get_db()
    try:
        yil = request.form.get("yil", type=int)
        k = StratejiKalem(yil=yil, bolum=bolum, veri_json="{}")
        _formdan(k, request.form)
        if not k.metin:
            flash("Metin boş olamaz.", "error")
        else:
            if not k.sira:
                k.sira = (db.query(StratejiKalem).filter(StratejiKalem.yil == yil, StratejiKalem.bolum == bolum, StratejiKalem.grup == k.grup).count() + 1)
            db.add(k); db.commit()
            log_action(db, "Oluşturma", detay=f"Stratejik planlama {yil} · {bolum}: {k.metin[:80]}")
            flash("Kalem eklendi.", "success")
        return redirect(url_for("strateji", yil=yil, sekme=bolum))
    finally:
        db.close()


@app.route("/strateji/kalem/<int:kid>", methods=["GET", "POST"])
@login_required
def strateji_kalem(kid):
    db = get_db()
    try:
        k = db.get(StratejiKalem, kid)
        if not k:
            abort(404)
        if request.method == "POST":
            if not _can_edit():
                abort(403)
            if request.form.get("islem") == "sil":
                yil, bolum, metin = k.yil, k.bolum, k.metin
                db.delete(k); db.commit()
                log_action(db, "Silme", detay=f"Stratejik planlama {yil} · {bolum}: {metin[:80]}")
                flash("Kalem silindi.", "success")
                return redirect(url_for("strateji", yil=yil, sekme=bolum))
            _formdan(k, request.form)
            if not k.metin:
                flash("Metin boş olamaz.", "error")
                return redirect(url_for("strateji_kalem", kid=k.id))
            db.commit()
            log_action(db, "Düzenleme", detay=f"Stratejik planlama {k.yil} · {k.bolum}: {k.metin[:80]}")
            flash("Kalem güncellendi.", "success")
            return redirect(url_for("strateji", yil=k.yil, sekme=k.bolum))
        return render_template("strateji_kalem.html", k=k, alanlar=ALANLAR[k.bolum], kapsamlar=STRATEJI_KAPSAMLARI,
                               bilesenler=STRATEJI_BAGLAM_BILESENLERI, swot_turleri=STRATEJI_SWOT_TURLERI, durumlar=STRATEJI_EYLEM_DURUMLARI,
                               sekme_adi=dict(SEKMELER)[k.bolum], can_edit=_can_edit(),
                               users=db.query(User).filter_by(aktif=True).order_by(User.ad_soyad).all())
    finally:
        db.close()


@app.route("/strateji/donem", methods=["POST"])
@login_required
def strateji_donem():
    if not _can_edit():
        abort(403)
    db = get_db()
    try:
        yil = request.form.get("yil", type=int)
        if not yil:
            abort(400)
        d = db.query(StratejiDonem).filter(StratejiDonem.yil == yil).first()
        kaynak_yil = request.form.get("kopyala", type=int)
        if not d:
            d = StratejiDonem(yil=yil, gozden_gecirme_json="{}"); db.add(d); db.flush()
            if kaynak_yil:   # yeni yıl: önceki yılın bağlam, paydaş, SWOT ve stratejileri taslak olarak kopyalanır
                n = 0
                for k in db.query(StratejiKalem).filter(StratejiKalem.yil == kaynak_yil).all():
                    veri = k.veri
                    if k.bolum == "eylem":
                        veri.update(gg_notu="", gg_tarihi="", gerceklesen="")
                    db.add(StratejiKalem(yil=yil, bolum=k.bolum, kapsam=k.kapsam, grup=k.grup, sira=k.sira, kod=k.kod, metin=k.metin,
                                         veri_json=json.dumps(veri, ensure_ascii=False), termin=None,
                                         durum="Planlandı" if k.bolum == "eylem" else None, sorumlu_id=k.sorumlu_id))
                    n += 1
                d.notlar = f"{kaynak_yil} döneminden kopyalandı ({n} kalem); gözden geçirilecek."
            db.commit()
            log_action(db, "Oluşturma", detay=f"Stratejik planlama dönemi {yil}" + (f" ({kaynak_yil}'den kopya)" if kaynak_yil else ""))
            flash(f"{yil} dönemi açıldı.", "success")
        else:
            gg = d.gozden_gecirme
            for b in STRATEJI_BOLUMLERI:
                if f"gg_{b}" in request.form:
                    gg[b] = request.form.get(f"gg_{b}") or None
            d.gozden_gecirme_json = json.dumps({k: v for k, v in gg.items() if v}, ensure_ascii=False)
            if "notlar" in request.form:
                d.notlar = request.form.get("notlar")
            db.commit()
            log_action(db, "Düzenleme", detay=f"Stratejik planlama dönemi {yil}: gözden geçirme bilgileri")
            flash("Dönem bilgileri kaydedildi.", "success")
        return redirect(url_for("strateji", yil=yil, sekme=request.form.get("sekme") or "baglam"))
    finally:
        db.close()


@app.route("/strateji/export.xlsx")
@login_required
def strateji_export():
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment
    db = get_db()
    try:
        yil = request.args.get("yil", type=int) or date.today().year
        wb = Workbook(); wb.remove(wb.active)
        for b, ad in SEKMELER:
            ws = wb.create_sheet(ad.split("·")[1].strip()[:31])
            bas = ["Kapsam"] + (["Bileşen"] if b == "baglam" else ["Tür"] if b == "swot" else []) + ["No", "Metin"] + [e for _, e, _ in ALANLAR[b]]
            if b == "eylem":
                bas += ["Planlanan termin", "Durum", "Sorumlu"]
            ws.append(bas)
            for c in ws[1]:
                c.font = Font(bold=True)
            for k in _kalemler(db, yil, b):
                v = k.veri
                satir = [k.kapsam] + ([k.grup] if b in ("baglam", "swot") else []) + [k.kod or k.sira, k.metin] + [v.get(a, "") for a, _, _ in ALANLAR[b]]
                if b == "eylem":
                    satir += [k.termin, k.durum, k.sorumlu.ad_soyad if k.sorumlu else ""]
                ws.append(satir)
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = 60 if col[0].value in ("Metin",) or len(str(col[0].value)) > 20 else 16
                for c in col:
                    c.alignment = Alignment(wrap_text=True, vertical="top")
        buf = io.BytesIO(); wb.save(buf); buf.seek(0)
        return send_file(buf, as_attachment=True, download_name=f"Stratejik_Planlama_{yil}.xlsx",
                         mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    finally:
        db.close()


@app.route("/strateji/yazdir")
@login_required
def strateji_yazdir():
    """Yazdırılabilir rapor: antet, özet (sayılar, durumlar, terminini geçenler), seçilen bölümler, imza alanı.
    ?yil=2026&bolum=baglam&bolum=swot…&kapsam=&yon=yatay|dikey"""
    from models import Process
    db = get_db()
    try:
        yil = request.args.get("yil", type=int) or date.today().year
        secili = [b for b in request.args.getlist("bolum") if b in STRATEJI_BOLUMLERI] or list(STRATEJI_BOLUMLERI)
        kapsam = request.args.get("kapsam") if request.args.get("kapsam") in STRATEJI_KAPSAMLARI else None
        veriler = {b: _kalemler(db, yil, b, kapsam) for b in STRATEJI_BOLUMLERI}
        bugun = date.today()
        eylem = veriler["eylem"]
        ozet = {"durum": {d: sum(1 for k in eylem if (k.durum or "Planlandı") == d) for d in STRATEJI_EYLEM_DURUMLARI},
                "geciken": [k for k in eylem if k.termin and k.termin < bugun and k.durum not in ("Tamamlandı", "İptal")],
                "swot": {t: sum(1 for k in veriler["swot"] if k.grup == t) for t in STRATEJI_SWOT_TURLERI}}
        donem = db.query(StratejiDonem).filter(StratejiDonem.yil == yil).first()
        tarihler = [v for k, v in (donem.gozden_gecirme if donem else {}).items() if not k.endswith("sayfa") and v]
        log_action(db, "İndirme", detay=f"Stratejik planlama {yil} raporu yazdırıldı ({', '.join(secili)}{' · ' + kapsam if kapsam else ''})")
        return render_template("strateji_yazdir.html", yil=yil, donem=donem, son_gg=max(tarihler) if tarihler else None,
                               sekmeler=SEKMELER, secili=secili, kapsam=kapsam, kapsamlar=STRATEJI_KAPSAMLARI, yon=request.args.get("yon") or "yatay",
                               alanlar=ALANLAR, veriler=veriler, ozet=ozet, swot_turleri=STRATEJI_SWOT_TURLERI, bugun=bugun,
                               simdi=datetime.now(), kullanici=session.get("ad_soyad") or "",
                               surec=db.query(Process).filter(Process.kod == "Y01").first(), yillar=_yillar(db))
    finally:
        db.close()

@app.route("/strateji/aktar", methods=["POST"])
@login_required
def strateji_aktar():
    if session.get("rol") != "Admin":
        abort(403)
    import glob
    import os
    import strateji_veri
    klasor = request.form.get("klasor") or os.environ.get("DYS_STRATEJI_PAKETI") or ""
    db = get_db()
    try:
        sonuc = [strateji_veri.paketten_yukle(db, p) for p in sorted(glob.glob(os.path.join(klasor, "strateji_*.json")))]
        log_action(db, "Oluşturma", detay=f"Stratejik planlama paketi: {sonuc}")
        flash(f"İçe aktarıldı: {sonuc}" if sonuc else "Aktarılacak paket bulunamadı.", "success" if sonuc else "warning")
        return redirect(url_for("strateji"))
    finally:
        db.close()
