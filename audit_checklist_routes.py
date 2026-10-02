"""
DYS — Denetim soru listeleri, önizleme, cevaplama ve kanıt yükleme
"""

from datetime import datetime, date
import os

from flask import (
    render_template, request, redirect, url_for, flash, abort, session,
    send_file, jsonify,
)

from app_runtime import host as _host
_h = _host()
app = _h.app
login_required = _h.login_required
admin_required = _h.admin_required
get_db = _h.get_db
log_action = _h.log_action

from models import (
    Process, InternalAudit, AuditChecklist, AuditChecklistQuestion,
    AuditAnswer, AuditEvidence, AuditEkKanit, Document, EK_KANIT_TIPLERI,
    AUDIT_CEVAPLARI, AUDIT_KANIT_TIPLERI, AUDIT_TIPLERI, BULGU_TIPLERI,
)
import denetim_otomasyon as OT
from models import AuditProgramItem as OT_Item
from helpers import allowed_file, save_generic_file, resolve_file_path
from audit_checklist_import import (
    import_all_checklists, find_checklist_for_audit, CHECKLIST_DIR,
)


def _can_edit(audit=None):
    """Admin / Doküman Kontrol ya da (2026) tetkike atanmış denetim ekibi üyesi."""
    if session.get("rol") in ("Admin", "Doküman Kontrol"):
        return True
    uid = session.get("user_id")
    return bool(audit and uid and (audit.tetkik_eden_id == uid or any(e.get("user_id") == uid for e in OT.ekip_oku(audit))))


def _audit_checklist(db, audit):
    if getattr(audit, "checklist_id", None):
        cl = db.get(AuditChecklist, audit.checklist_id)
        if cl:
            return cl
    surec_kod = audit.denetlenen_surec.kod if audit.denetlenen_surec else None
    if audit.denetim_tipi == "Sistem Denetimi" and surec_kod:
        return find_checklist_for_audit(db, "Sistem Denetimi", OT.ust_kod(surec_kod)) if db.query(AuditChecklist).filter_by(aktif=True, surec_kod=OT.ust_kod(surec_kod)).first() \
            else find_checklist_for_audit(db, "Sistem Denetimi", surec_kod)
    return find_checklist_for_audit(db, audit.denetim_tipi or "Sistem Denetimi", surec_kod)


def _get_or_create_answer(db, audit_id, question_id):
    ans = (
        db.query(AuditAnswer)
        .filter_by(audit_id=audit_id, question_id=question_id)
        .first()
    )
    if not ans:
        ans = AuditAnswer(
            audit_id=audit_id, question_id=question_id, sonuc="Değerlendirilmedi",
        )
        db.add(ans)
        db.flush()
    return ans


def save_answers_from_form(db, audit_id, form):
    """Form alanları sonuc_<qid> / gozlem_<qid> → AuditAnswer. Dönüş: kayıt sayısı."""
    n = 0
    for key in form:
        if not key.startswith("sonuc_"):
            continue
        try:
            qid = int(key.split("_", 1)[1])
        except (ValueError, IndexError):
            continue
        sonuc = form.get(key) or "Değerlendirilmedi"
        if sonuc not in AUDIT_CEVAPLARI:
            sonuc = "Değerlendirilmedi"
        gozlem = (form.get(f"gozlem_{qid}") or "").strip() or None
        derece = form.get(f"derece_{qid}") or None
        if derece not in BULGU_TIPLERI:
            derece = None
        mevcut = db.query(AuditAnswer).filter_by(audit_id=audit_id, question_id=qid).first()
        if sonuc == "Değerlendirilmedi" and not gozlem and not mevcut:
            continue
        if mevcut and mevcut.sonuc == sonuc and (mevcut.gozlem or None) == gozlem and (mevcut.bulgu_derecesi or None) == derece:
            continue
        ans = mevcut or _get_or_create_answer(db, audit_id, qid)
        ans.sonuc = sonuc
        ans.gozlem = gozlem
        ans.bulgu_derecesi = derece
        ans.guncelleme_tarihi = datetime.now()
        n += 1
    return n


@app.route("/api/audits/checklist-preview")
@login_required
def audit_checklist_preview_api():
    """Yeni tetkik formunda süreç/tip seçimine göre soru listesi JSON."""
    tip = (request.args.get("tip") or "Sistem Denetimi").strip()
    if tip not in AUDIT_TIPLERI:
        tip = "Sistem Denetimi"
    surec_id = request.args.get("surec_id", type=int)
    db = get_db()
    try:
        surec_kod = None
        if surec_id:
            p = db.get(Process, surec_id)
            surec_kod = p.kod if p else None
        cl = find_checklist_for_audit(db, tip, surec_kod)
        if not cl:
            return jsonify({
                "ok": False,
                "message": "Bu tip/süreç için soru listesi bulunamadı. Önce Soru Listeleri ekranından aktarın.",
                "checklist": None,
                "questions": [],
                "cevaplar": list(AUDIT_CEVAPLARI),
            })
        questions = sorted(cl.questions, key=lambda q: (q.sira or 0, q.id))
        return jsonify({
            "ok": True,
            "message": None,
            "checklist": {
                "id": cl.id,
                "kod": cl.kod,
                "ad": cl.ad,
                "denetim_tipi": cl.denetim_tipi,
                "surec_kod": cl.surec_kod,
                "soru_sayisi": len(questions),
            },
            "questions": [
                {
                    "id": q.id,
                    "sira": q.sira,
                    "bolum": q.bolum or "Genel",
                    "soru_no": q.soru_no,
                    "soru_metin": q.soru_metin,
                    "sart_no": q.sart_no,
                }
                for q in questions
            ],
            "cevaplar": list(AUDIT_CEVAPLARI),
        })
    finally:
        db.close()


@app.route("/audits/checklists")
@login_required
def audit_checklist_catalog():
    db = get_db()
    try:
        listeler = (
            db.query(AuditChecklist)
            .filter_by(aktif=True)
            .order_by(AuditChecklist.denetim_tipi, AuditChecklist.surec_kod, AuditChecklist.kod)
            .all()
        )
        return render_template(
            "audit_checklists.html",
            listeler=listeler,
            excel_var=os.path.isdir(CHECKLIST_DIR),
            can_edit=_can_edit(),
        )
    finally:
        db.close()


def _liste_yonetici():
    return session.get("rol") in ("Admin", "Doküman Kontrol")


@app.route("/audits/checklists/<int:cid>", methods=["GET", "POST"])
@login_required
def audit_checklist_detail(cid):
    """2026 — soru listesinin içi: bölümlere göre sorular (şart, kime, ipucu), kullanıldığı denetimler;
    Admin / Doküman Kontrol: liste bilgisi, soru ekle / düzenle / sil (cevap almış soru silinmez)."""
    db = get_db()
    try:
        cl = db.get(AuditChecklist, cid)
        if not cl:
            abort(404)
        qids = [q.id for q in cl.questions]
        cevap = {}
        if qids:
            from sqlalchemy import func
            cevap = dict(db.query(AuditAnswer.question_id, func.count(AuditAnswer.id)).filter(AuditAnswer.question_id.in_(qids))
                         .group_by(AuditAnswer.question_id).all())
        if request.method == "POST":
            if not _liste_yonetici():
                abort(403)
            f = request.form
            a = f.get("action")
            git = url_for("audit_checklist_detail", cid=cl.id)
            s = lambda k, n: ((f.get(k) or "").strip()[:n] or None)
            if a == "liste":
                if not (f.get("ad") or "").strip():
                    flash("Liste adı zorunludur.", "error"); return redirect(git)
                cl.ad = f.get("ad").strip()[:200]
                if f.get("denetim_tipi") in (*AUDIT_TIPLERI, "Saha Denetimi"):
                    cl.denetim_tipi = f.get("denetim_tipi")
                cl.surec_kod = s("surec_kod", 10)
                cl.aktif = f.get("aktif") == "1"
                db.commit(); log_action(db, "Güncelleme", detay=f"Soru listesi bilgisi: {cl.kod}")
                flash("Liste bilgileri kaydedildi.", "success")
                return redirect(git)
            if a in ("soru_kaydet", "soru_ekle"):
                metin = (f.get("soru_metin") or "").strip()
                if not metin:
                    flash("Soru metni zorunludur.", "error"); return redirect(git)
                if a == "soru_ekle":
                    q = AuditChecklistQuestion(checklist_id=cl.id, sira=0, soru_metin=metin)
                    db.add(q)
                else:
                    q = db.get(AuditChecklistQuestion, f.get("qid", type=int) or 0)
                    if not q or q.checklist_id != cl.id:
                        abort(404)
                    q.soru_metin = metin
                q.bolum = s("bolum", 120) or "Genel"; q.soru_no = s("soru_no", 20); q.sart_no = s("sart_no", 40)
                q.kime = s("kime", 200); q.ipucu = (f.get("ipucu") or "").strip() or None
                sira = f.get("sira", type=int)
                db.flush()
                L = sorted((x for x in cl.questions if x.id != q.id), key=lambda x: (x.sira or 0, x.id))
                konum = max(0, min(len(L), (sira - 1) if sira else len(L)))
                L.insert(konum, q)
                for i, x in enumerate(L, 1):
                    x.sira = i
                db.commit()
                log_action(db, "Oluşturma" if a == "soru_ekle" else "Güncelleme", detay=f"Soru listesi {cl.kod}: soru {q.soru_no or q.sira}"
                           + (f" ({cevap.get(q.id)} cevaplı)" if cevap.get(q.id) else ""))
                flash("Soru eklendi." if a == "soru_ekle" else "Soru kaydedildi." + (" Bu soru daha önce cevaplanmıştı — eski denetimlerde yeni metin görünür." if cevap.get(q.id) else ""), "success")
                return redirect(git + f"#q-{q.id}")
            if a == "soru_sil":
                q = db.get(AuditChecklistQuestion, f.get("qid", type=int) or 0)
                if not q or q.checklist_id != cl.id:
                    abort(404)
                if cevap.get(q.id):
                    flash(f"Bu soru {cevap[q.id]} denetimde cevaplanmış; silinemez.", "error"); return redirect(git + f"#q-{q.id}")
                cl.questions.remove(q); db.delete(q); db.flush()
                for i, x in enumerate(sorted(cl.questions, key=lambda x: (x.sira or 0, x.id)), 1):
                    x.sira = i
                db.commit(); log_action(db, "Silme", detay=f"Soru listesi {cl.kod}: soru silindi")
                flash("Soru silindi.", "success")
                return redirect(git)
            abort(400)
        sorular = sorted(cl.questions, key=lambda x: (x.sira or 0, x.id))
        bolumler = []
        for q in sorular:
            b = q.bolum or "Genel"
            if not bolumler or bolumler[-1][0] != b:
                bolumler.append((b, []))
            bolumler[-1][1].append(q)
        denetimler = db.query(InternalAudit).filter_by(checklist_id=cl.id).order_by(InternalAudit.planlanan_tarih.desc()).all()
        return render_template("audit_checklist_detail.html", cl=cl, bolumler=bolumler, toplam=len(sorular), cevap=cevap,
                               denetimler=denetimler, yonetici=_liste_yonetici(), tipler=(*AUDIT_TIPLERI, "Saha Denetimi"),
                               bolum_adlari=sorted({q.bolum or "Genel" for q in sorular}), cikti=request.args.get("cikti") == "1")
    finally:
        db.close()


@app.route("/audits/checklists/<int:cid>/excel")
@login_required
def audit_checklist_excel(cid):
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    db = get_db()
    try:
        cl = db.get(AuditChecklist, cid)
        if not cl:
            abort(404)
        wb = Workbook(); ws = wb.active; ws.title = "Soru Listesi"
        ws.append([f"{cl.kod} — {cl.ad}"]); ws["A1"].font = Font(bold=True, size=13)
        ws.append(["Sıra", "Bölüm", "Soru No", "Soru", "Şart No", "Kime", "İpucu / ilgili doküman", "Gözlem / Delil", "Sonuç"])
        for c in ws[2]:
            c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F4E78")
        for q in sorted(cl.questions, key=lambda x: (x.sira or 0, x.id)):
            ws.append([q.sira, q.bolum or "", q.soru_no or "", q.soru_metin, q.sart_no or "", q.kime or "", q.ipucu or "", "", ""])
        for i, w in enumerate([6, 22, 9, 70, 14, 26, 40, 30, 12], 1):
            ws.column_dimensions[ws.cell(2, i).column_letter].width = w
        for row in ws.iter_rows(min_row=3):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A3"
        bio = io.BytesIO(); wb.save(bio); bio.seek(0)
        return send_file(bio, as_attachment=True, download_name=f"{cl.kod} {cl.ad}.xlsx"[:150],
                         mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    finally:
        db.close()


@app.route("/audits/checklists/import", methods=["POST"])
@login_required
@admin_required
def audit_checklist_import():
    db = get_db()
    try:
        replace = bool(request.form.get("replace"))
        n_cl, n_q = import_all_checklists(db, replace=replace)
        log_action(db, "Oluşturma", detay=f"Denetim soru listesi: {n_cl} liste, {n_q} soru")
        flash(f"Soru listeleri aktarıldı: {n_cl} checklist, {n_q} soru.", "success")
        return redirect(url_for("audit_checklist_catalog"))
    except Exception as exc:
        db.rollback()
        flash(f"Aktarım başarısız: {exc}", "error")
        return redirect(url_for("audit_checklist_catalog"))
    finally:
        db.close()


@app.route("/audits/<int:audit_id>/checklist", methods=["GET", "POST"])
@login_required
def audit_checklist_run(audit_id):
    db = get_db()
    try:
        audit = db.get(InternalAudit, audit_id)
        if not audit:
            flash("Tetkik bulunamadı.", "error")
            return redirect(url_for("audit_list"))

        checklist = None if (audit.denetim_tipi == "Sistem Denetimi" and not audit.denetlenen_surec and not getattr(audit, "checklist_id", None)) \
            else _audit_checklist(db, audit)
        can_edit = _can_edit(audit)

        if request.method == "POST":
            if not can_edit:
                abort(403)
            action = request.form.get("action") or "save"
            if action == "liste_sec":
                cid = request.form.get("checklist_id", type=int)
                if cid and db.get(AuditChecklist, cid):
                    audit.checklist_id = cid
                    db.commit()
                    flash("Soru listesi bağlandı.", "success")
                return redirect(url_for("audit_checklist_run", audit_id=audit.id))
            if action == "save" and checklist:
                n = save_answers_from_form(db, audit.id, request.form)
                db.flush()
                db.expire(audit)
                yeni_bulgu = OT.bulgulari_esitle(db, audit)
                db.commit()
                log_action(db, "Güncelleme", detay=f"Tetkik soru listesi: {audit.tetkik_no} — {n} cevap, {yeni_bulgu} yeni bulgu")
                flash((f"{n} cevap kaydedildi." if n else "Değişiklik yok.") + (f" {yeni_bulgu} yeni bulgu oluştu." if yeni_bulgu else ""), "success")
                return redirect(url_for("audit_checklist_run", audit_id=audit.id))
            if action == "tamamla":
                ozet = (request.form.get("sonuc_ozeti") or "").strip()
                if not ozet:
                    flash("Sonuç özeti zorunludur.", "error")
                    return redirect(url_for("audit_checklist_run", audit_id=audit.id))
                try:
                    audit.gerceklesen_tarih = datetime.strptime(request.form.get("gerceklesen_tarih") or "", "%Y-%m-%d").date()
                except ValueError:
                    audit.gerceklesen_tarih = date.today()
                audit.sonuc_ozeti = ozet
                audit.durum = "Tamamlandı"
                OT.bulgulari_esitle(db, audit)
                db.flush()
                db.expire(audit)
                dof = OT.dof_ac(db, audit, session.get("user_id")) if OT.ayar(db, "denetim_dof_otomatik") == "1" else []
                for it in db.query(OT_Item).filter_by(audit_id=audit.id).all():
                    it.durum = "Tamamlandı"
                db.commit()
                log_action(db, "Güncelleme", detay=f"Tetkik tamamlandı: {audit.tetkik_no}; DÖF: {', '.join(c.dof_no for c in dof) or '—'}")
                flash(f"Denetim tamamlandı." + (f" Açılan DÖF: {', '.join(c.dof_no for c in dof)}" if dof else ""), "success")
                return redirect(url_for("audit_detail", audit_id=audit.id))

            if action == "evidence":
                qid = request.form.get("question_id", type=int)
                f = request.files.get("dosya")
                if not qid or not f or not f.filename:
                    flash("Soru ve dosya gerekli.", "error")
                    return redirect(url_for("audit_checklist_run", audit_id=audit.id))
                if not allowed_file(f.filename):
                    flash("Dosya türüne izin verilmiyor.", "error")
                    return redirect(url_for("audit_checklist_run", audit_id=audit.id))
                ans = _get_or_create_answer(db, audit.id, qid)
                ad, yol = save_generic_file(f, "audit_evidence", audit.id, ans.id)
                db.add(AuditEvidence(
                    answer_id=ans.id,
                    kanit_tipi=request.form.get("kanit_tipi") or "Diğer",
                    aciklama=(request.form.get("aciklama") or "").strip() or None,
                    dosya_adi=ad, dosya_yolu=yol,
                    yukleyen_id=session.get("user_id"),
                ))
                db.commit()
                flash(f"Kanıt yüklendi: {ad}", "success")
                return redirect(url_for("audit_checklist_run", audit_id=audit.id) + f"#q-{qid}")

            if action == "ek_kanit":
                if OT.denetim_kategorisi(db, audit) not in OT.EK_KANIT_KATEGORILERI:
                    abort(400)
                n, hatali = _ek_kanit_ekle(db, audit, request)
                db.commit()
                if n:
                    log_action(db, "Oluşturma", detay=f"Tetkik ek kanıt: {audit.tetkik_no} — {n} kayıt")
                    flash(f"{n} ek kanıt eklendi.", "success")
                for h in hatali:
                    flash(h, "error")
                if not n and not hatali:
                    flash("Dosya seçin ya da bir DYS doküman numarası girin.", "error")
                return redirect(url_for("audit_checklist_run", audit_id=audit.id) + "#ek-kanit")
            if action == "ek_kanit_sil":
                k = db.get(AuditEkKanit, request.form.get("kanit_id", type=int) or 0)
                if not k or k.audit_id != audit.id:
                    abort(404)
                if audit.durum in ("Tamamlandı", "Kapatıldı") and session.get("rol") != "Admin":
                    flash("Tamamlanan denetimin kanıtını yalnız Admin kaldırabilir.", "error")
                    return redirect(url_for("audit_checklist_run", audit_id=audit.id) + "#ek-kanit")
                ad = k.dosya_adi or (f"doküman #{k.dokuman_id}")
                if k.dosya_yolu:
                    yol = resolve_file_path(k.dosya_yolu, k.dosya_adi)
                    if yol and os.path.isfile(yol):
                        try:
                            os.remove(yol)
                        except OSError:
                            pass
                db.delete(k); db.commit()
                log_action(db, "Silme", detay=f"Tetkik ek kanıt kaldırıldı: {audit.tetkik_no} — {ad}")
                flash(f"Kanıt kaldırıldı: {ad}", "success")
                return redirect(url_for("audit_checklist_run", audit_id=audit.id) + "#ek-kanit")

        answers = {a.question_id: a for a in db.query(AuditAnswer).filter_by(audit_id=audit.id).all()}
        kategori = OT.denetim_kategorisi(db, audit)
        ek_goster = kategori in OT.EK_KANIT_KATEGORILERI
        ek_kanitlar = _ek_kanit_listesi(db, audit) if ek_goster else []
        bolumler = []
        if checklist:
            cur, grup = None, []
            for q in sorted(checklist.questions, key=lambda x: (x.sira or 0, x.id)):
                b = q.bolum or "Genel"
                if b != cur:
                    if grup:
                        bolumler.append((cur, grup))
                    cur, grup = b, []
                grup.append(q)
            if grup:
                bolumler.append((cur, grup))

        stats = {"toplam": 0, "cevapli": 0, "uygunsuz": 0, "kismen": 0, "kanit": 0}
        if checklist:
            stats["toplam"] = len(checklist.questions)
            for q in checklist.questions:
                a = answers.get(q.id)
                if a and a.sonuc != "Değerlendirilmedi":
                    stats["cevapli"] += 1
                if a and a.sonuc == "Uygunsuz":
                    stats["uygunsuz"] += 1
                if a and a.sonuc == "Kısmen Uygun":
                    stats["kismen"] += 1
                if a:
                    stats["kanit"] += len(a.evidences or [])
        stats["kanit"] += len(ek_kanitlar)

        return render_template(
            "audit_checklist_run.html",
            audit=audit, checklist=checklist, bolumler=bolumler, answers=answers,
            stats=stats, cevaplar=AUDIT_CEVAPLARI, kanit_tipleri=AUDIT_KANIT_TIPLERI,
            can_edit=can_edit, ekip=OT.ekip_oku(audit), bulgu_tipleri=BULGU_TIPLERI, today=date.today(),
            dof_otomatik=OT.ayar(db, "denetim_dof_otomatik") == "1",
            kategori=kategori, ek_goster=ek_goster, ek_kanitlar=ek_kanitlar, ek_kanit_tipleri=EK_KANIT_TIPLERI,
            dok_secenek=_dok_secenekleri(db) if ek_goster and can_edit else [],
            listeler=db.query(AuditChecklist).filter_by(aktif=True).order_by(AuditChecklist.kod).all() if not checklist else [],
        )
    finally:
        db.close()


@app.route("/audits/<int:audit_id>/checklist/print")
@login_required
def audit_checklist_print(audit_id):
    """Sahada kâğıt üzerinde doldurmak için soru listesi çıktısı (sonra DYS'ye girilir)."""
    db = get_db()
    try:
        audit = db.get(InternalAudit, audit_id)
        if not audit:
            abort(404)
        checklist = _audit_checklist(db, audit)
        answers = {a.question_id: a for a in db.query(AuditAnswer).filter_by(audit_id=audit.id).all()}
        sorular = sorted(checklist.questions, key=lambda x: (x.sira or 0, x.id)) if checklist else []
        cikti_no = f"ÇKT-{audit.tetkik_no}-{datetime.now():%y%m%d%H%M%S}"
        log_action(db, "Görüntüleme", detay=f"Tetkik soru listesi çıktısı: {cikti_no}")
        return render_template("audit_checklist_print.html", audit=audit, checklist=checklist, sorular=sorular,
                               answers=answers, ekip=OT.ekip_oku(audit), cikti_no=cikti_no, simdi=datetime.now())
    finally:
        db.close()


# ─────────────── 2026: ek kanıt dokümanları (proses / ürün denetimi) ───────────────
_GORSEL = {"png", "jpg", "jpeg", "gif", "webp", "bmp"}


def _ek_kanit_ekle(db, audit, req):
    """Formdaki dosyalar (çoklu) ve / veya DYS doküman numarası → AuditEkKanit. (eklenen sayısı, hata mesajları)"""
    tip = (req.form.get("kanit_tipi") or "Diğer").strip()
    tip = tip if tip in EK_KANIT_TIPLERI else "Diğer"
    aciklama = (req.form.get("aciklama") or "").strip()[:300] or None
    qid = req.form.get("question_id", type=int)
    if qid:
        q = db.get(AuditChecklistQuestion, qid)
        qid = q.id if q and audit.checklist_id and q.checklist_id == audit.checklist_id else None
    n, hatali = 0, []
    for f in req.files.getlist("ek_dosyalar"):
        if not f or not f.filename:
            continue
        if not allowed_file(f.filename):
            hatali.append(f"{f.filename}: dosya türüne izin verilmiyor."); continue
        ad, yol = save_generic_file(f, "audit_ek_kanit", audit.id)
        db.add(AuditEkKanit(audit_id=audit.id, question_id=qid, kanit_tipi=tip, aciklama=aciklama, dosya_adi=ad, dosya_yolu=yol,
                            yukleyen_id=session.get("user_id")))
        n += 1
    dno = (req.form.get("dokuman_no") or "").split(" — ")[0].strip()
    if dno:
        d = db.query(Document).filter(Document.dokuman_no == dno).order_by(Document.id.desc()).first()
        if not d:
            hatali.append(f"DYS'de '{dno}' numaralı doküman bulunamadı.")
        else:
            db.add(AuditEkKanit(audit_id=audit.id, question_id=qid, kanit_tipi=tip, aciklama=aciklama, dokuman_id=d.id,
                                yukleyen_id=session.get("user_id")))
            n += 1
    return n, hatali


def _ek_kanit_listesi(db, audit):
    from models import User
    out = []
    for k in db.query(AuditEkKanit).filter_by(audit_id=audit.id).order_by(AuditEkKanit.id).all():
        d = db.get(Document, k.dokuman_id) if k.dokuman_id else None
        q = db.get(AuditChecklistQuestion, k.question_id) if k.question_id else None
        u = db.get(User, k.yukleyen_id) if k.yukleyen_id else None
        ext = (k.dosya_adi or "").rsplit(".", 1)[-1].lower() if k.dosya_adi and "." in k.dosya_adi else ""
        out.append({"k": k, "dokuman": d, "soru": q, "yukleyen": u.ad_soyad if u else "—", "gorsel": ext in _GORSEL, "pdf": ext == "pdf"})
    return out


def _dok_secenekleri(db):
    """Bağlanabilecek yürürlükteki DYS dokümanları (doküman no — başlık)."""
    return [f"{n} — {b}" for n, b in db.query(Document.dokuman_no, Document.baslik)
            .filter(~Document.durum.in_(["Eskimiş", "İptal"])).order_by(Document.dokuman_no).all() if n]


def _ek_kanitlari_sablon(audit):
    """Şablon yardımcısı (tetkik detayı, çıktı): proses / ürün denetiminin ek kanıtları."""
    db = get_db()
    try:
        return _ek_kanit_listesi(db, audit) if OT.denetim_kategorisi(db, audit) in OT.EK_KANIT_KATEGORILERI else []
    finally:
        db.close()


app.jinja_env.globals["denetim_ek_kanitlari"] = _ek_kanitlari_sablon


@app.route("/audits/ek-kanit/<int:kanit_id>/dosya")
@login_required
def audit_ek_kanit_dosya(kanit_id):
    """Ek kanıt dosyası — görsel / PDF tarayıcıda açılır, diğerleri indirilir."""
    db = get_db()
    try:
        k = db.get(AuditEkKanit, kanit_id)
        if not k or not k.dosya_yolu:
            abort(404)
        path = resolve_file_path(k.dosya_yolu, k.dosya_adi)
        if not path or not os.path.isfile(path):
            flash("Kanıt dosyası bulunamadı.", "error")
            return redirect(url_for("audit_checklist_run", audit_id=k.audit_id))
        ext = (k.dosya_adi or "").rsplit(".", 1)[-1].lower()
        return send_file(path, as_attachment=not (ext in _GORSEL or ext == "pdf") or request.args.get("indir") == "1",
                         download_name=k.dosya_adi or "kanit")
    finally:
        db.close()


@app.route("/audits/evidence/<int:ev_id>/download")
@login_required
def audit_evidence_download(ev_id):
    db = get_db()
    try:
        ev = db.get(AuditEvidence, ev_id)
        if not ev or not ev.dosya_yolu:
            abort(404)
        path = resolve_file_path(ev.dosya_yolu, ev.dosya_adi)
        if not path or not os.path.isfile(path):
            flash("Kanıt dosyası bulunamadı.", "error")
            return redirect(url_for("audit_list"))
        return send_file(path, as_attachment=True, download_name=ev.dosya_adi or "kanit")
    finally:
        db.close()
