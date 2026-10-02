# -*- coding: utf-8 -*-
import sqlite3
import re
import openpyxl
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from datetime import datetime
from pathlib import Path
import shutil

DB = r"\\192.168.0.249\ortak\HARUN\dys\dys.db"
XLSX = r"c:\Users\labaratuvar03\Downloads\ana_dokuman_listesi2-.xlsx"
BASE = Path(r"\\192.168.0.249\ortak\HARUN\dys")


def norm(k):
    if not k:
        return ""
    return re.sub(r"[\s\-]+", "", str(k).upper())


def fmt_date(v):
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%d.%m.%Y")
    return str(v)


def load_excel_talimatlar():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb["Ana Doküman Listesi"]
    out = {}
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 3).value == "T":
            kod = ws.cell(r, 2).value
            out[norm(kod)] = {
                "kod": kod,
                "tanim": ws.cell(r, 4).value,
                "rev": ws.cell(r, 5).value,
                "tarih": ws.cell(r, 6).value,
            }
    return out


def resolve_path(p):
    if not p:
        return None
    path = Path(p)
    if path.exists():
        return path
    alt = BASE / "uploads" / Path(*Path(p).parts[Path(p).parts.index("uploads") + 1 :]) if "uploads" in p.replace("\\", "/") else None
    if alt and alt.exists():
        return alt
    # try replacing local prefix with network base
    marker = "uploads\\documents"
    if marker in p.replace("/", "\\"):
        rel = p.split(marker, 1)[1].lstrip("\\/")
        candidate = BASE / "uploads" / "documents" / rel
        if candidate.exists():
            return candidate
    return path if path.exists() else None


def inspect_docx(path):
    doc = Document(str(path))
    info = {"path": str(path), "sections": len(doc.sections)}
    for i, sec in enumerate(doc.sections):
        hf = {}
        for name in ("header", "footer", "first_page_header", "first_page_footer"):
            part = getattr(sec, name, None)
            if part is None:
                continue
            texts = []
            for p in part.paragraphs:
                t = p.text.strip()
                if t:
                    texts.append(t)
            for tbl in part.tables:
                for row in tbl.rows:
                    for cell in row.cells:
                        t = cell.text.strip()
                        if t:
                            texts.append(t)
            hf[name] = texts
        info[f"section_{i}"] = hf
    body_sample = [p.text.strip() for p in doc.paragraphs[:8] if p.text.strip()]
    info["body_sample"] = body_sample
    return info


def compare():
    conn = sqlite3.connect(DB)
    cur = conn.cursor()
    cur.execute(
        "SELECT dokuman_no, baslik, dosya_yolu, revizyon_no FROM documents WHERE dokuman_tipi='Talimat'"
    )
    db = {
        norm(r[0]): {"no": r[0], "baslik": r[1], "path": r[2], "rev": r[3]}
        for r in cur.fetchall()
    }
    ex = load_excel_talimatlar()

    common = set(db) & set(ex)
    only_db = set(db) - set(ex)
    only_ex = set(ex) - set(db)

    print(f"Common: {len(common)} | Portal only: {len(only_db)} | Excel only: {len(only_ex)}")

    mismatches = []
    for k in sorted(common):
        db_t = (db[k]["baslik"] or "").strip().lower()
        ex_t = (ex[k]["tanim"] or "").strip().lower()
        if db_t != ex_t:
            mismatches.append((db[k]["no"], db[k]["baslik"], ex[k]["tanim"]))

    print(f"Title mismatches: {len(mismatches)}")
    for no, db_t, ex_t in mismatches[:25]:
        print(f"  {no}")
        print(f"    Portal: {db_t}")
        print(f"    Excel : {ex_t}")

    if only_db:
        print("\nPortal only:")
        for k in sorted(only_db):
            print(f"  {db[k]['no']} -> {db[k]['baslik']}")

    if only_ex:
        print("\nExcel only:")
        for k in sorted(only_ex):
            print(f"  {ex[k]['kod']} -> {ex[k]['tanim']}")

    conn.close()
    return db, ex, mismatches


def update_docx_meta(path, kod, rev, tarih, backup=True):
    path = Path(path)
    if backup:
        bak = path.with_suffix(path.suffix + ".bak")
        shutil.copy2(path, bak)

    doc = Document(str(path))
    rev_s = str(rev) if rev is not None else "0"
    tarih_s = fmt_date(tarih)
    footer_line = f"Doküman No: {kod}    Rev: {rev_s}    Tarih: {tarih_s}"

    updated = False
    for sec in doc.sections:
        footer = sec.footer
        existing_text = "\n".join(p.text for p in footer.paragraphs).strip()

        # Mevcut alt bilgi alanında dok no / rev / tarih varsa güncelle
        if re.search(r"(dok[üu]man\s*no|rev|tarih|revizyon)", existing_text, re.I):
            for p in footer.paragraphs:
                txt = p.text
                new_txt = re.sub(
                    r"Dok[üu]man\s*No\s*[:.]?\s*[^\n\r|]*",
                    f"Doküman No: {kod}",
                    txt,
                    flags=re.I,
                )
                new_txt = re.sub(
                    r"Rev(?:\.|izyon)?(?:\s*No)?\s*[:.]?\s*[^\n\r|]*",
                    f"Rev: {rev_s}",
                    new_txt,
                    flags=re.I,
                )
                new_txt = re.sub(
                    r"Tarih\s*[:.]?\s*[^\n\r|]*",
                    f"Tarih: {tarih_s}",
                    new_txt,
                    flags=re.I,
                )
                if new_txt != txt:
                    p.text = new_txt
                    updated = True
        else:
            p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
            p.text = footer_line
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            for run in p.runs:
                run.font.size = Pt(9)
            updated = True

        # Üst bilgi
        header = sec.header
        header_text = "\n".join(p.text for p in header.paragraphs).strip()
        if re.search(r"(dok[üu]man\s*no|rev|tarih|revizyon)", header_text, re.I):
            for p in header.paragraphs:
                txt = p.text
                new_txt = re.sub(
                    r"Dok[üu]man\s*No\s*[:.]?\s*[^\n\r|]*",
                    f"Doküman No: {kod}",
                    txt,
                    flags=re.I,
                )
                new_txt = re.sub(
                    r"Rev(?:\.|izyon)?(?:\s*No)?\s*[:.]?\s*[^\n\r|]*",
                    f"Rev: {rev_s}",
                    new_txt,
                    flags=re.I,
                )
                new_txt = re.sub(
                    r"Tarih\s*[:.]?\s*[^\n\r|]*",
                    f"Tarih: {tarih_s}",
                    new_txt,
                    flags=re.I,
                )
                if new_txt != txt:
                    p.text = new_txt
                    updated = True

    doc.save(str(path))
    return updated, footer_line


if __name__ == "__main__":
    db, ex, mismatches = compare()

    # Örnek: D01 T03 El Yıkama Talimatı
    sample_key = norm("D01 T03")
    if sample_key in db and sample_key in ex:
        rec = db[sample_key]
        meta = ex[sample_key]
        fpath = resolve_path(rec["path"])
        print(f"\n--- SAMPLE: {rec['no']} / {rec['baslik']} ---")
        print(f"Excel: kod={meta['kod']} rev={meta['rev']} tarih={fmt_date(meta['tarih'])}")
        print(f"File: {fpath}")
        if fpath and fpath.suffix.lower() == ".docx":
            print("\nBEFORE:")
            import json

            print(json.dumps(inspect_docx(fpath), ensure_ascii=False, indent=2))
            ok, line = update_docx_meta(fpath, meta["kod"], meta["rev"], meta["tarih"])
            print(f"\nUPDATED footer line: {line}")
            print("\nAFTER:")
            print(json.dumps(inspect_docx(fpath), ensure_ascii=False, indent=2))
        else:
            print("Sample is not docx or file missing")
