# -*- coding: utf-8 -*-
"""
Talimat .docx dosyalarında doküman no / rev / tarih güncelleme.
Word biçimlerini korur: paragraf hizası, stil ve ilk run'un font özellikleri.
"""
from __future__ import annotations

import re
import shutil
import sqlite3
import traceback
import zipfile
from datetime import datetime
from pathlib import Path

import openpyxl
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

DB = r"\\192.168.0.249\ortak\HARUN\dys\dys.db"
XLSX = r"c:\Users\labaratuvar03\Downloads\ana_dokuman_listesi2-.xlsx"
BASE = Path(r"D:\UYGULAMA PROJELER\debak dijital\dys")
LOG = BASE / "_talimat_bulk_log.txt"

DOC_CODE_RE = re.compile(r"[A-Z]\d{2}(?:\.\d+)?\s*T\d+", re.I)
DATE_RE = re.compile(r"\d{1,2}\.\d{1,2}\.\d{4}")
REV_PAREN_RE = re.compile(r"\(\s*\d+\s*\)")
META_HINT_RE = re.compile(
    r"(dok[üu]man\s*no|rev(?:\.|izyon)?|tarih|[A-Z]\d{2}(?:\.\d+)?\s*T\d+|\(\d+\)|\d{2}\.\d{2}\.\d{4})",
    re.I,
)


def norm(k: str | None) -> str:
    if not k:
        return ""
    return re.sub(r"[\s\-]+", "", str(k).upper())


def fmt_date(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%d.%m.%Y")
    s = str(v).strip()
    if " " in s:
        s = s.split(" ")[0]
    return s


def fmt_rev(v) -> str:
    if v is None or v == "":
        return "0"
    try:
        return str(int(v))
    except (TypeError, ValueError):
        return str(v).strip()


def load_excel() -> dict:
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb["Ana Doküman Listesi"]
    out = {}
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 3).value == "T":
            kod = ws.cell(r, 2).value
            if not kod:
                continue
            out[norm(kod)] = {
                "kod": str(kod).strip(),
                "rev": ws.cell(r, 5).value,
                "tarih": ws.cell(r, 6).value,
            }
    return out


def resolve_path(p: str | None) -> Path | None:
    if not p:
        return None
    path = Path(p)
    if path.exists():
        return path
    p2 = p.replace("/", "\\")
    if "uploads\\documents" in p2:
        rel = p2.split("uploads\\documents", 1)[1].lstrip("\\/")
        c = BASE / "uploads" / "documents" / rel
        if c.exists():
            return c
    return None


def is_valid_docx(path: Path) -> bool:
    try:
        with zipfile.ZipFile(path) as z:
            return "word/document.xml" in z.namelist()
    except zipfile.BadZipFile:
        return False


def detect_format(old: str) -> str:
    old = old.strip()
    if re.search(r"dok[üu]man\s*no", old, re.I):
        return "labeled"
    if REV_PAREN_RE.search(old):
        if "_" in old:
            return "compact_rev_underscore"
        return "compact_rev_space"
    if DOC_CODE_RE.search(old) and DATE_RE.search(old):
        return "compact_no_rev"
    return "compact_no_rev"


def build_meta_line(kod: str, rev, tarih, fmt: str) -> str:
    rev_s = fmt_rev(rev)
    tarih_s = fmt_date(tarih)
    if fmt == "labeled":
        return f"Doküman No: {kod}    Rev: {rev_s}    Tarih: {tarih_s}"
    if fmt == "compact_rev_underscore":
        return f"{kod}({rev_s})_{tarih_s}"
    if fmt == "compact_rev_space":
        return f"{kod}({rev_s}) {tarih_s}"
    return f"{kod} {tarih_s}"


def paragraph_has_meta(text: str) -> bool:
    t = text.strip()
    if not t:
        return False
    if META_HINT_RE.search(t):
        if DOC_CODE_RE.search(t) or DATE_RE.search(t) or re.search(
            r"dok[üu]man|rev|tarih", t, re.I
        ):
            return True
    return False


def set_paragraph_text_keep_style(paragraph, new_text: str) -> None:
    """Metni değiştir; paragraf ve ilk run biçimini koru."""
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(new_text)


def update_labeled_paragraph(paragraph, kod: str, rev, tarih) -> bool:
    old = paragraph.text
    if not paragraph_has_meta(old):
        return False
    rev_s = fmt_rev(rev)
    tarih_s = fmt_date(tarih)
    new = old
    new = re.sub(
        r"Dok[üu]man\s*No\s*[:.]?\s*[^\n\r|]*",
        f"Doküman No: {kod}",
        new,
        flags=re.I,
    )
    new = re.sub(
        r"Rev(?:\.|izyon)?(?:\s*No)?\s*[:.]?\s*[^\n\r|]*",
        f"Rev: {rev_s}",
        new,
        flags=re.I,
    )
    new = re.sub(
        r"Tarih\s*[:.]?\s*[^\n\r|]*",
        f"Tarih: {tarih_s}",
        new,
        flags=re.I,
    )
    if new == old:
        return False
    set_paragraph_text_keep_style(paragraph, new)
    return True


def update_compact_paragraph(paragraph, kod: str, rev, tarih) -> bool:
    old = paragraph.text.strip()
    if not paragraph_has_meta(old):
        return False
    fmt = detect_format(old)
    new = build_meta_line(kod, rev, tarih, fmt)
    if old == new:
        return False
    set_paragraph_text_keep_style(paragraph, new)
    return True


def iter_meta_paragraphs(container):
    for p in container.paragraphs:
        if paragraph_has_meta(p.text):
            yield p
    for tbl in container.tables:
        for row in tbl.rows:
            for cell in row.cells:
                yield from iter_meta_paragraphs(cell)


def add_footer_meta(paragraph, kod: str, rev, tarih) -> None:
    """Boş alt bilgiye sağ hizalı kompakt satır ekle (mevcut paragraf stilini koru)."""
    new = build_meta_line(kod, rev, tarih, "compact_no_rev")
    set_paragraph_text_keep_style(paragraph, new)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if paragraph.runs:
        r = paragraph.runs[0]
        if r.font.size is None:
            r.font.size = Pt(9)


def process_part(part, kod: str, rev, tarih) -> tuple[bool, str | None]:
    updated = False
    last_line = None
    meta_paras = list(iter_meta_paragraphs(part))
    for p in meta_paras:
        old = p.text.strip()
        if re.search(r"dok[üu]man\s*no", old, re.I):
            if update_labeled_paragraph(p, kod, rev, tarih):
                updated = True
                last_line = p.text.strip()
        elif update_compact_paragraph(p, kod, rev, tarih):
            updated = True
            last_line = p.text.strip()
    return updated, last_line


def update_docx(path: Path, kod: str, rev, tarih, backup: bool = True) -> dict:
    result = {"path": str(path), "kod": kod, "status": "skip", "line": None, "error": None}
    if not is_valid_docx(path):
        result["status"] = "invalid_docx"
        return result

    bak = path.with_suffix(path.suffix + ".bak_meta")
    if backup and not bak.exists():
        shutil.copy2(path, bak)

    try:
        doc = Document(str(path))
    except Exception as e:
        result["status"] = "open_error"
        result["error"] = str(e)
        return result

    updated = False
    last_line = None

    for sec in doc.sections:
        for part_name in ("header", "footer", "first_page_header", "first_page_footer"):
            part = getattr(sec, part_name, None)
            if part is None:
                continue
            ok, line = process_part(part, kod, rev, tarih)
            if ok:
                updated = True
                last_line = line

        footer = sec.footer
        footer_text = "\n".join(p.text.strip() for p in footer.paragraphs if p.text.strip())
        if not paragraph_has_meta(footer_text):
            # Alt bilgi yoksa: boş paragrafa ekle veya yeni paragraf
            target = None
            for p in footer.paragraphs:
                if not p.text.strip():
                    target = p
                    break
            if target is None:
                target = footer.add_paragraph()
            add_footer_meta(target, kod, rev, tarih)
            updated = True
            last_line = target.text.strip()

    if updated:
        doc.save(str(path))
        result["status"] = "updated"
        result["line"] = last_line
    else:
        result["status"] = "unchanged"

    return result


def main():
    excel = load_excel()
    conn = sqlite3.connect(DB)
    cur = conn.cursor()
    cur.execute(
        """
        SELECT dokuman_no, baslik, dosya_yolu
        FROM documents
        WHERE dokuman_tipi='Talimat'
        ORDER BY dokuman_no
        """
    )
    rows = cur.fetchall()
    conn.close()

    stats = {"updated": 0, "unchanged": 0, "skipped_no_excel": 0, "skipped_not_docx": 0,
             "missing_file": 0, "invalid_docx": 0, "error": 0}
    lines = []

    for dokuman_no, baslik, dosya_yolu in rows:
        key = norm(dokuman_no)
        if key not in excel:
            stats["skipped_no_excel"] += 1
            lines.append(f"SKIP(no excel)\t{dokuman_no}\t{baslik}")
            continue

        fp = resolve_path(dosya_yolu)
        if not fp:
            stats["missing_file"] += 1
            lines.append(f"MISSING\t{dokuman_no}\t{dosya_yolu}")
            continue

        if fp.suffix.lower() != ".docx":
            stats["skipped_not_docx"] += 1
            continue

        meta = excel[key]
        try:
            res = update_docx(fp, meta["kod"], meta["rev"], meta["tarih"])
        except Exception:
            stats["error"] += 1
            lines.append(f"ERROR\t{dokuman_no}\t{traceback.format_exc(limit=1)}")
            continue

        st = res["status"]
        if st == "updated":
            stats["updated"] += 1
            lines.append(f"OK\t{dokuman_no}\t{res['line']}")
        elif st == "unchanged":
            stats["unchanged"] += 1
            lines.append(f"SAME\t{dokuman_no}\t{meta['kod']}")
        elif st == "invalid_docx":
            stats["invalid_docx"] += 1
            lines.append(f"BADZIP\t{dokuman_no}\t{fp}")
        else:
            stats["error"] += 1
            lines.append(f"ERR\t{dokuman_no}\t{res.get('error')}")

    summary = (
        f"Bitti: updated={stats['updated']} unchanged={stats['unchanged']} "
        f"invalid={stats['invalid_docx']} missing={stats['missing_file']} "
        f"no_excel={stats['skipped_no_excel']} not_docx={stats['skipped_not_docx']} "
        f"errors={stats['error']}"
    )
    LOG.write_text(summary + "\n\n" + "\n".join(lines), encoding="utf-8")
    print(summary)
    print(f"Log: {LOG}")


if __name__ == "__main__":
    main()
