# -*- coding: utf-8 -*-
"""Inspect all text locations in a docx file."""
from docx import Document
from pathlib import Path

PATH = Path(r"D:\UYGULAMA PROJELER\debak dijital\dys\uploads\documents\287\rev0\D01 T03 El Yıkama Talimatı.docx")


def iter_paragraphs(container, location):
    for i, p in enumerate(container.paragraphs):
        t = p.text.strip()
        if t:
            yield location, i, t
    for ti, tbl in enumerate(container.tables):
        for ri, row in enumerate(tbl.rows):
            for ci, cell in enumerate(row.cells):
                for pi, p in enumerate(cell.paragraphs):
                    t = p.text.strip()
                    if t:
                        yield f"{location}/tbl{ti}r{ri}c{ci}p{pi}", pi, t


def scan(path):
    doc = Document(str(path))
    hits = []
    for loc, idx, text in iter_paragraphs(doc, "body"):
        if "D2" in text or "T44" in text or "D01" in text or "T03" in text:
            hits.append(("body", loc, text))
    for si, sec in enumerate(doc.sections):
        for name in ("header", "footer", "first_page_header", "first_page_footer"):
            part = getattr(sec, name, None)
            if not part:
                continue
            for loc, idx, text in iter_paragraphs(part, f"sec{si}.{name}"):
                if "D2" in text or "T44" in text or "D01" in text or "T03" in text or text:
                    hits.append((name, loc, text))
    return hits


if __name__ == "__main__":
    print("=== CURRENT FILE ===")
    for kind, loc, text in scan(PATH):
        print(f"[{kind}] {loc}: {repr(text)}")

    bak = PATH.with_suffix(PATH.suffix + ".bak")
    if bak.exists():
        print("\n=== BACKUP FILE ===")
        for kind, loc, text in scan(bak):
            print(f"[{kind}] {loc}: {repr(text)}")
