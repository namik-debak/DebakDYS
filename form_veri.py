"""
DYS — Dinamik form verilerinin işlenebilir dışa aktarımı (2026)
- Her form için Excel: 'Kayıtlar' (geniş: bir kayıt = bir satır), 'Kontrol Listeleri' (uzun: madde başına satır),
  'Tablolar' (uzun: tablo satırı başına satır). Power Query / Power BI / pivot için sabit sütun adları.
- DYS_FORM_VERI_KLASORU tanımlıysa kayıt oluşturma / onay / red sonrasında formun dosyası otomatik yenilenir:
    <klasör>\\<süreç>\\<form kodu> <form adı>.xlsx   ve   <klasör>\\_TUM_FORM_KAYITLARI.xlsx (tüm formlar, uzun biçim)
"""
import json
import logging
import os
import re
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

logger = logging.getLogger("dys")
META = ["Kayıt No", "Form Kodu", "Form Adı", "Form Rev.", "Durum", "Kaynak", "Oluşturan", "Oluşturma Tarihi", "Onaylayan", "Onay Tarihi", "Onay Notu", "Ekler"]


def _alanlar(t):
    try:
        return [a for a in json.loads(t.alanlar_json or "[]") if isinstance(a, dict) and a.get("ad") and a.get("tip") != "baslik"]
    except ValueError:
        return []


def _veri(k):
    try:
        return json.loads(k.veriler_json or "{}") or {}
    except ValueError:
        return {}


def _meta(t, k, v):
    return [k.kayit_no, t.form_kodu, t.ad, k.form_revizyon_no, k.durum, v.get("_kaynak", "DYS"),
            k.olusturan.ad_soyad if k.olusturan else "", k.olusturma_tarihi,
            k.onaylayan.ad_soyad if k.onaylayan else "", k.onay_tarihi, k.onay_notu or "", ", ".join(v.get("_ekler") or [])]


def _deger(tip, x):
    if tip in ("sayi", "hesap"):
        try:
            return float(str(x).replace(",", ".")) if x not in (None, "") else None
        except ValueError:
            return x
    if tip == "tarih" and x:
        try:
            return datetime.strptime(x, "%Y-%m-%d").date()
        except (TypeError, ValueError):
            return x
    if isinstance(x, list):
        return ", ".join(str(i) for i in x)
    return x


_SAYI = re.compile(r"^-?\d+([.,]\d+)?$")


def _sayiya(x):
    """Tablo hücresi '12,5' / '40.1' gibi sayıysa sayı döndürür (Excel'de hesaplanabilsin)."""
    if isinstance(x, str) and _SAYI.match(x.strip()):
        try:
            return float(x.strip().replace(",", "."))
        except ValueError:
            return x
    return x


def _sayfa(wb, ad, basliklar, satirlar, ilk=False):
    ws = wb.active if ilk else wb.create_sheet()
    ws.title = ad[:31]
    ws.append(basliklar)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3864")
    for r in satirlar:
        ws.append(r)
    for i, h in enumerate(basliklar, 1):
        ws.column_dimensions[get_column_letter(i)].width = min(45, max(10, len(str(h)) + 2))
    ws.freeze_panes = "A2"
    if satirlar:
        ws.auto_filter.ref = ws.dimensions
    return ws


def form_calisma_kitabi(tanimlar_kayitlar, tek_form=True):
    """tanimlar_kayitlar: [(FormTanim, [FormKayit...])] → Workbook"""
    wb = Workbook()
    genis_sayfalar, kl, tb, uzun = [], [], [], []
    for t, kayitlar in tanimlar_kayitlar:
        alanlar = _alanlar(t)
        bas = list(META)
        for a in alanlar:
            if a.get("tip") == "kontrol_listesi":
                for j, m in enumerate(a.get("maddeler") or [], 1):
                    bas += [f"{a['ad']} | {j}. Sonuç", f"{a['ad']} | {j}. Açıklama"]
            elif a.get("tip") == "tablo":
                bas.append(f"{a['ad']} (satır sayısı)")
            else:
                bas.append(a["ad"])
        satirlar = []
        for k in kayitlar:
            v = _veri(k)
            row = _meta(t, k, v)
            for a in alanlar:
                x = v.get(a["ad"])
                if a.get("tip") == "kontrol_listesi":
                    d = x if isinstance(x, list) else []
                    for j, m in enumerate(a.get("maddeler") or []):
                        s = d[j] if j < len(d) and isinstance(d[j], dict) else {}
                        row += [s.get("sonuc", ""), s.get("aciklama", "")]
                        kl.append([k.kayit_no, t.form_kodu, t.ad, k.durum, k.olusturma_tarihi, a["ad"], j + 1, m, s.get("sonuc", ""), s.get("aciklama", "")])
                        uzun.append([k.kayit_no, t.form_kodu, t.ad, k.durum, k.olusturma_tarihi, f"{a['ad']} | {m}", s.get("sonuc", ""), s.get("aciklama", "")])
                elif a.get("tip") == "tablo":
                    d = x if isinstance(x, list) else []
                    row.append(len(d))
                    for n, s in enumerate(d, 1):
                        for kol, val in (s or {}).items():
                            tb.append([k.kayit_no, t.form_kodu, t.ad, k.durum, k.olusturma_tarihi, a["ad"], n, kol, _sayiya(val)])
                            uzun.append([k.kayit_no, t.form_kodu, t.ad, k.durum, k.olusturma_tarihi, f"{a['ad']} | {n}. satır | {kol}", _sayiya(val), ""])
                else:
                    row.append(_deger(a.get("tip"), x))
                    uzun.append([k.kayit_no, t.form_kodu, t.ad, k.durum, k.olusturma_tarihi, a["ad"], _deger(a.get("tip"), x), ""])
            satirlar.append(row)
        genis_sayfalar.append((t, bas, satirlar))
    if tek_form and len(genis_sayfalar) == 1:
        t, bas, satirlar = genis_sayfalar[0]
        _sayfa(wb, "Kayıtlar", bas, satirlar, ilk=True)
    else:
        _sayfa(wb, "Tüm Veriler (uzun)", ["Kayıt No", "Form Kodu", "Form Adı", "Durum", "Oluşturma Tarihi", "Alan", "Değer", "Açıklama"], uzun, ilk=True)
        for t, bas, satirlar in genis_sayfalar:
            _sayfa(wb, re.sub(r"[\\/*?:\[\]]", "-", t.form_kodu), bas, satirlar)
    _sayfa(wb, "Kontrol Listeleri", ["Kayıt No", "Form Kodu", "Form Adı", "Durum", "Oluşturma Tarihi", "Liste", "Madde No", "Madde", "Sonuç", "Açıklama"], kl)
    _sayfa(wb, "Tablolar", ["Kayıt No", "Form Kodu", "Form Adı", "Durum", "Oluşturma Tarihi", "Tablo", "Satır No", "Sütun", "Değer"], tb)
    bilgi = wb.create_sheet("Bilgi")
    bilgi.append(["Oluşturma", datetime.now().strftime("%d.%m.%Y %H:%M")])
    bilgi.append(["Kaynak", "DYS — Dinamik Formlar"])
    bilgi.append(["Not", "Sütun adları sabittir; Excel Veri Al / Power Query / Power BI bağlantısı için uygundur. Dosya DYS tarafından yeniden yazılır, elle değişiklik yapmayın."])
    return wb


def _guvenli_kaydet(wb, yol):
    os.makedirs(os.path.dirname(yol), exist_ok=True)
    gecici = yol + ".tmp.xlsx"
    wb.save(gecici)
    try:
        os.replace(gecici, yol)
    except PermissionError:  # dosya kullanıcıda açık: yanına tarihli kopya bırak
        alt = yol.replace(".xlsx", f" ({datetime.now():%Y%m%d-%H%M%S}).xlsx")
        os.replace(gecici, alt)
        logger.warning("Form veri dosyası açık olduğu için %s yazıldı", alt)


def _veri_yolu(klasor, tanim):
    surec = (tanim.surec.kod if tanim.surec else tanim.form_kodu.split()[0]).split(".")[0]
    ad = re.sub(r'[\\/:*?"<>|]', "-", f"{tanim.form_kodu} {tanim.ad}")[:120]
    return os.path.join(klasor, surec, ad + ".xlsx")


def kisitli_dosyayi_kaldir(tanim):
    """Form kısıtlı yapıldığında daha önce ortak klasöre yazılmış veri dosyası kaldırılır."""
    from config import Config
    klasor = getattr(Config, "FORM_VERI_KLASORU", None)
    if not klasor:
        return
    try:
        yol = _veri_yolu(klasor, tanim)
        if os.path.isfile(yol):
            os.remove(yol)
    except Exception as exc:  # noqa: BLE001
        logger.warning("Kısıtlı form veri dosyası kaldırılamadı (%s): %s", tanim.form_kodu, exc)


def otomatik_disa_aktar(db, tanim):
    """DYS_FORM_VERI_KLASORU tanımlıysa bu formun ve tüm formların veri dosyasını yeniler. Hata uygulamayı durdurmaz."""
    from config import Config
    from models import FormKayit, FormTanim
    klasor = getattr(Config, "FORM_VERI_KLASORU", None)
    if not klasor:
        return None
    try:
        if getattr(tanim, "kisitli", False):   # KVKK: kısıtlı formun kayıtları ortak klasöre yazılmaz
            kisitli_dosyayi_kaldir(tanim)
            tanim = None
        kayitlar = db.query(FormKayit).filter(FormKayit.tanim_id == tanim.id).order_by(FormKayit.kayit_no).all() if tanim else []
        yol = None
        if tanim:
            yol = _veri_yolu(klasor, tanim)
            _guvenli_kaydet(form_calisma_kitabi([(tanim, kayitlar)]), yol)
        tum = []
        for t in db.query(FormTanim).order_by(FormTanim.form_kodu).all():
            if t.kisitli:
                continue
            kk = db.query(FormKayit).filter(FormKayit.tanim_id == t.id).order_by(FormKayit.kayit_no).all()
            if kk:
                tum.append((t, kk))
        _guvenli_kaydet(form_calisma_kitabi(tum, tek_form=False), os.path.join(klasor, "_TUM_FORM_KAYITLARI.xlsx"))
        return yol
    except Exception as exc:  # noqa: BLE001
        logger.warning("Form verisi otomatik dışa aktarılamadı (%s): %s", tanim.form_kodu, exc)
        return None
