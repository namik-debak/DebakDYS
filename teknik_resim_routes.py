"""
DYS — Teknik Resim Takip ve Revizyon modülü (2026) — D04.4 F02 Teknik Resim Takip Listesi yerine
Akış: müşteri → ürün (teknik resim) → revizyonlar. Dosyalar manuel yüklenir.
  /teknik-resim                         müşteriler (kartlar) + tüm resimlerde arama + uyarılar
  /teknik-resim?musteri=<ad>            müşterinin ürünleri / resimleri (geçerli revizyon, dağıtım, değerlendirme)
  /teknik-resim/yeni                    yeni ürün resmi + ilk revizyon
  /teknik-resim/<id>                    resim kartı: geçerli revizyon önizleme, revizyon geçmişi, yeni revizyon, dağıtım, değerlendirme
  /teknik-resim/rev/<rid>/dosya         revizyon dosyası (PDF / görsel tarayıcıda açılır; ?indir=1)
  /teknik-resim/excel[?musteri=]        D04.4 F02 düzeninde Excel liste
Kurallar: yeni revizyon yüklenince önceki geçerli revizyon "Geçersiz" olur (geçmiş revizyon arşiv olarak da girilebilir).
Dağıtım: Yönetim / Laboratuvar / Kalite / Üretim — dağıtım tarihi; geçersiz revizyonda eski kopyanın toplandığı işaretlenir.
Değişiklik değerlendirmesi (IATF 16949 7.5.3.2.2): etkilenen dokümanlar (kontrol planı, PFMEA, kalıp, mastar, PPAP …), teslim almadan
itibaren DEGERLENDIRME_GUN gün içinde yapılmalı; gecikenler uyarı olarak gösterilir.
Yetki: görüntüleme herkes; kayıt / yükleme Sadece Görüntüleme hariç; revizyon silme Admin / Doküman Kontrol.
"""
import io
import json
import os
from datetime import date, datetime, timedelta

from flask import render_template, request, redirect, url_for, flash, abort, session, send_file
from werkzeug.utils import secure_filename

from app_runtime import host as _host
_h = _host()
app = _h.app
login_required = _h.login_required
get_db = _h.get_db
log_action = _h.log_action

from models import (TeknikResim, TeknikResimRevizyon, TR_DAGITIM, TR_ETKILER, TR_EK_TURLERI, User,  # noqa: E402
                    PPAPSubmission, MusteriSikayeti)
from helpers import resolve_file_path  # noqa: E402
from config import Config  # noqa: E402

DEGERLENDIRME_GUN = 14
UZANTILAR = {"pdf", "tif", "tiff", "png", "jpg", "jpeg", "dwg", "dxf", "stp", "step", "igs", "iges", "x_t", "zip", "xls", "xlsx"}
GORUNTU = {"pdf", "png", "jpg", "jpeg", "gif", "webp"}


def _yazabilir():
    return session.get("rol") != "Sadece Görüntüleme"


def _yonetici():
    return session.get("rol") in ("Admin", "Doküman Kontrol")


def _k(s):
    return " ".join((s or "").replace("İ", "i").replace("I", "ı").lower().split())


def _tarih(s):
    try:
        return datetime.strptime(s, "%Y-%m-%d").date() if s else None
    except ValueError:
        return None


def _ext(ad):
    return ad.rsplit(".", 1)[-1].lower() if ad and "." in ad else ""


def _json(s):
    try:
        return json.loads(s) if s else {}
    except ValueError:
        return {}


def _revler(db, rid):
    return sorted(db.query(TeknikResimRevizyon).filter_by(resim_id=rid).all(),
                  key=lambda r: (r.teslim_alma_tarihi or date.min, r.id))


def _gecerli(revler):
    g = [r for r in revler if r.durum == "Geçerli"]
    return g[-1] if g else None


def _rev_durum(r, bugun=None):
    """Revizyon için uyarılar: değerlendirme bekliyor / gecikti, dağıtım eksik, eski kopya toplanmadı."""
    bugun = bugun or date.today()
    out = []
    if r.kaynak and not r.dagitim_json and not r.degerlendirme_json:
        return out                      # sistem öncesi (toplu aktarım) resim: dağıtım / değerlendirme kaydı beklenmez
    dag = _json(r.dagitim_json)
    if r.durum == "Geçerli":
        if not _json(r.degerlendirme_json).get("tarih"):
            gec = r.teslim_alma_tarihi and (bugun - r.teslim_alma_tarihi).days > DEGERLENDIRME_GUN
            out.append(("kritik" if gec else "uyari", "Değerlendirme gecikti" if gec else "Değerlendirme bekliyor"))
        eksik = [b for b in TR_DAGITIM if not (dag.get(b) or {}).get("tarih")]
        if eksik:
            out.append(("uyari", "Dağıtım eksik: " + ", ".join(eksik)))
    else:
        top = [b for b in TR_DAGITIM if (dag.get(b) or {}).get("tarih") and not (dag.get(b) or {}).get("toplandi")]
        if top:
            out.append(("uyari", "Eski kopya toplanmadı: " + ", ".join(top)))
    return out


def _musteri_onerileri(db):
    s = {(m or "").strip() for (m,) in db.query(TeknikResim.musteri).distinct().all()}
    for model in (PPAPSubmission, MusteriSikayeti):
        try:
            s |= {(m or "").strip() for (m,) in db.query(model.musteri).distinct().all()}
        except Exception:  # noqa: BLE001
            pass
    tek = {}
    for m in sorted(x for x in s if x):
        tek.setdefault(_k(m), m)
    return sorted(tek.values(), key=_k)


def _musteri_adi(db, ad):
    """Aynı müşterinin farklı yazımlarını ('Valeo ', 'VALEO') mevcut yazıma bağlar."""
    ad = " ".join((ad or "").split())
    for (m,) in db.query(TeknikResim.musteri).distinct().all():
        if _k(m) == _k(ad):
            return m
    return ad


def _dosya_kaydet(f, resim_id, rev_id):
    ad = os.path.basename(f.filename or "").strip()
    if _ext(ad) not in UZANTILAR:
        return None, None, f"{ad}: dosya türü desteklenmiyor ({', '.join(sorted(UZANTILAR))})."
    hedef = os.path.join(Config.UPLOAD_FOLDER, "teknik_resim", str(resim_id), str(rev_id))
    os.makedirs(hedef, exist_ok=True)
    yol = os.path.join(hedef, secure_filename(ad) or f"resim.{_ext(ad)}")
    f.save(yol)
    return ad[:260], yol, None


def _revizyon_ekle(db, resim, form, files, arsiv=False):
    rev = (form.get("revizyon") or "").strip()[:40]
    if not rev:
        return None, "Revizyon zorunludur."
    if any(_k(r.revizyon) == _k(rev) for r in _revler(db, resim.id)):
        return None, f"{rev} revizyonu bu resimde zaten var."
    r = TeknikResimRevizyon(resim_id=resim.id, revizyon=rev, revizyon_tarihi=_tarih(form.get("revizyon_tarihi")),
                            teslim_alma_tarihi=_tarih(form.get("teslim_alma_tarihi")) or date.today(),
                            degisiklik=(form.get("degisiklik") or "").strip()[:1000] or None,
                            durum="Geçersiz" if arsiv else "Geçerli", yukleyen_id=session.get("user_id"))
    if arsiv:
        r.gecersiz_tarihi = date.today()
    db.add(r); db.flush()
    f = files.get("dosya")
    if f and f.filename:
        ad, yol, hata = _dosya_kaydet(f, resim.id, r.id)
        if hata:
            db.rollback()
            return None, hata
        r.dosya_adi, r.dosya_yolu = ad, yol
    if not arsiv:
        for eski in _revler(db, resim.id):
            if eski.id != r.id and eski.durum == "Geçerli":
                eski.durum, eski.gecersiz_tarihi = "Geçersiz", date.today()
    resim.guncelleme_tarihi = datetime.now()
    return r, None


# ─────────────────────────────── ekranlar ───────────────────────────────
@app.route("/teknik-resim")
@login_required
def teknik_resim_liste():
    db = get_db()
    try:
        musteri = (request.args.get("musteri") or "").strip()
        q = (request.args.get("q") or "").strip()
        pasif = request.args.get("pasif") == "1"
        resimler = db.query(TeknikResim).all()
        revs = {}
        for r in db.query(TeknikResimRevizyon).all():
            revs.setdefault(r.resim_id, []).append(r)
        satirlar = []
        for t in resimler:
            rl = sorted(revs.get(t.id, []), key=lambda r: (r.teslim_alma_tarihi or date.min, r.id))
            g = _gecerli(rl)
            uy = _rev_durum(g) if g else [("uyari", "Geçerli revizyon yok")]
            for r in rl:
                if r is not g:
                    uy += _rev_durum(r)
            satirlar.append({"t": t, "gecerli": g, "rev_sayisi": len(rl), "uyarilar": uy if t.aktif is not False else []})
        musteriler = {}
        for s in satirlar:
            if s["t"].aktif is False:
                continue
            m = musteriler.setdefault(_k(s["t"].musteri), {"ad": s["t"].musteri, "urun": 0, "bekleyen": 0, "kritik": 0, "son": None})
            m["urun"] += 1
            m["bekleyen"] += 1 if s["uyarilar"] else 0
            m["kritik"] += 1 if any(u[0] == "kritik" for u in s["uyarilar"]) else 0
            ta = s["gecerli"].teslim_alma_tarihi if s["gecerli"] else None
            if ta and (not m["son"] or ta > m["son"]):
                m["son"] = ta
        liste = satirlar
        if musteri:
            liste = [s for s in liste if _k(s["t"].musteri) == _k(musteri)]
        if q:
            kq = _k(q)
            liste = [s for s in liste if kq in _k(" ".join(filter(None, [s["t"].urun_adi, s["t"].urun_kodu, s["t"].resim_no, s["t"].musteri,
                                                                         s["gecerli"].revizyon if s["gecerli"] else ""])))]
        if not pasif:
            liste = [s for s in liste if s["t"].aktif is not False]
        liste.sort(key=lambda s: (_k(s["t"].musteri), _k(s["t"].urun_adi), _k(s["t"].resim_no)))
        ozet = {"urun": sum(m["urun"] for m in musteriler.values()), "musteri": len(musteriler),
                "bekleyen": sum(1 for s in satirlar if s["uyarilar"]), "kritik": sum(1 for s in satirlar if any(u[0] == "kritik" for u in s["uyarilar"]))}
        musteri_adi = next((m["ad"] for k, m in musteriler.items() if k == _k(musteri)), musteri) if musteri else ""
        return render_template("teknik_resim_liste.html", musteriler=sorted(musteriler.values(), key=lambda m: _k(m["ad"])),
                               liste=liste, musteri=musteri_adi, q=q, pasif=pasif, ozet=ozet, yazabilir=_yazabilir(),
                               gun=DEGERLENDIRME_GUN, arama=bool(musteri or q))
    finally:
        db.close()


@app.route("/teknik-resim/yeni", methods=["GET", "POST"])
@login_required
def teknik_resim_yeni():
    if not _yazabilir():
        abort(403)
    db = get_db()
    try:
        if request.method == "POST":
            f = request.form
            musteri = _musteri_adi(db, f.get("musteri"))
            urun_adi, resim_no = (f.get("urun_adi") or "").strip(), (f.get("resim_no") or "").strip()
            if not (musteri and urun_adi and resim_no):
                flash("Müşteri, ürün adı ve resim no zorunludur.", "error")
                return render_template("teknik_resim_yeni.html", form=f, musteriler=_musteri_onerileri(db))
            var = next((t for t in db.query(TeknikResim).filter(TeknikResim.musteri == musteri).all() if _k(t.resim_no) == _k(resim_no)), None)
            if var:
                flash(f"{musteri} için {resim_no} resmi zaten kayıtlı — yeni revizyonu bu karttan yükleyin.", "error")
                return redirect(url_for("teknik_resim_detay", rid=var.id))
            t = TeknikResim(musteri=musteri[:150], urun_adi=urun_adi[:200], resim_no=resim_no[:80],
                            urun_kodu=(f.get("urun_kodu") or "").strip()[:80] or None, aciklama=(f.get("aciklama") or "").strip()[:500] or None,
                            aktif=True, olusturan_id=session.get("user_id"))
            db.add(t); db.flush()
            r, hata = _revizyon_ekle(db, t, f, request.files)
            if hata:
                db.rollback()
                flash(hata, "error")
                return render_template("teknik_resim_yeni.html", form=f, musteriler=_musteri_onerileri(db))
            db.commit()
            log_action(db, "Oluşturma", detay=f"Teknik resim: {musteri} / {resim_no} {urun_adi} — rev {r.revizyon}")
            flash(f"{resim_no} kaydedildi (rev {r.revizyon}).", "success")
            return redirect(url_for("teknik_resim_detay", rid=t.id))
        return render_template("teknik_resim_yeni.html", form={"musteri": request.args.get("musteri", "")}, musteriler=_musteri_onerileri(db))
    finally:
        db.close()


@app.route("/teknik-resim/<int:rid>", methods=["GET", "POST"])
@login_required
def teknik_resim_detay(rid):
    db = get_db()
    try:
        t = db.get(TeknikResim, rid)
        if not t:
            abort(404)
        if request.method == "POST":
            if not _yazabilir():
                abort(403)
            f = request.form
            a = f.get("action")
            git = url_for("teknik_resim_detay", rid=t.id)
            if a == "bilgi":
                urun_adi, resim_no = (f.get("urun_adi") or "").strip(), (f.get("resim_no") or "").strip()
                if not (urun_adi and resim_no and (f.get("musteri") or "").strip()):
                    flash("Müşteri, ürün adı ve resim no zorunludur.", "error")
                    return redirect(git)
                t.musteri = _musteri_adi(db, f.get("musteri"))[:150]
                t.urun_adi, t.resim_no = urun_adi[:200], resim_no[:80]
                t.urun_kodu = (f.get("urun_kodu") or "").strip()[:80] or None
                t.aciklama = (f.get("aciklama") or "").strip()[:500] or None
                t.aktif = f.get("aktif") == "1"
                db.commit(); log_action(db, "Güncelleme", detay=f"Teknik resim bilgisi: {t.musteri} / {t.resim_no}")
                flash("Bilgiler kaydedildi.", "success")
            elif a in ("revizyon", "arsiv"):
                r, hata = _revizyon_ekle(db, t, f, request.files, arsiv=(a == "arsiv"))
                if hata:
                    flash(hata, "error")
                    return redirect(git + "#yeni-rev")
                db.commit()
                log_action(db, "Oluşturma", detay=f"Teknik resim revizyonu: {t.musteri} / {t.resim_no} rev {r.revizyon}" + (" (arşiv)" if a == "arsiv" else ""))
                flash(f"Rev {r.revizyon} {'arşive eklendi' if a == 'arsiv' else 'yüklendi; önceki revizyon geçersiz yapıldı. Dağıtımı ve değerlendirmeyi tamamlayın'}.", "success")
                return redirect(git + f"#rev-{r.id}")
            elif a in ("dagitim", "degerlendirme", "dosya", "rev_sil", "ek_ekle", "ek_sil"):
                r = db.get(TeknikResimRevizyon, f.get("rev_id", type=int) or 0)
                if not r or r.resim_id != t.id:
                    abort(404)
                if a == "dagitim":
                    dag = {}
                    for i, b in enumerate(TR_DAGITIM):
                        tr_ = _tarih(f.get(f"tarih_{i}"))
                        if tr_ or f.get(f"toplandi_{i}"):
                            dag[b] = {"tarih": tr_.isoformat() if tr_ else None, "toplandi": bool(f.get(f"toplandi_{i}"))}
                    r.dagitim_json = json.dumps(dag, ensure_ascii=False) if dag else None
                    db.commit(); log_action(db, "Güncelleme", detay=f"Teknik resim dağıtımı: {t.resim_no} rev {r.revizyon}")
                    flash("Dağıtım kaydedildi.", "success")
                elif a == "degerlendirme":
                    tr_ = _tarih(f.get("tarih"))
                    etk = [e for e in TR_ETKILER if f.get(f"etki_{TR_ETKILER.index(e)}")]
                    if not tr_:
                        flash("Değerlendirme tarihi zorunludur.", "error")
                        return redirect(git + f"#rev-{r.id}")
                    u = db.get(User, session.get("user_id")) if session.get("user_id") else None
                    r.degerlendirme_json = json.dumps({"tarih": tr_.isoformat(), "kim": u.ad_soyad if u else "", "etkiler": etk,
                                                       "etki_yok": not etk, "ppap": bool(f.get("ppap")),
                                                       "not": (f.get("not") or "").strip()[:1000]}, ensure_ascii=False)
                    db.commit(); log_action(db, "Güncelleme", detay=f"Teknik resim değişiklik değerlendirmesi: {t.resim_no} rev {r.revizyon}")
                    flash("Değerlendirme kaydedildi.", "success")
                elif a == "dosya":
                    fl = request.files.get("dosya")
                    if not fl or not fl.filename:
                        flash("Dosya seçin.", "error")
                        return redirect(git + f"#rev-{r.id}")
                    ad, yol, hata = _dosya_kaydet(fl, t.id, r.id)
                    if hata:
                        flash(hata, "error")
                        return redirect(git + f"#rev-{r.id}")
                    r.dosya_adi, r.dosya_yolu = ad, yol
                    db.commit(); log_action(db, "Güncelleme", detay=f"Teknik resim dosyası: {t.resim_no} rev {r.revizyon} — {ad}")
                    flash("Dosya yüklendi.", "success")
                elif a == "ek_ekle":
                    ekler = _json(r.ek_dosyalar_json) or []
                    tur = f.get("ek_tur") if f.get("ek_tur") in TR_EK_TURLERI else "Diğer"
                    n = 0
                    for fl in request.files.getlist("ek_dosyalar"):
                        if not fl or not fl.filename:
                            continue
                        ad, yol, hata = _dosya_kaydet(fl, t.id, f"{r.id}_ek")
                        if hata:
                            flash(hata, "error"); continue
                        ekler.append({"ad": ad, "yol": yol, "tur": tur}); n += 1
                    r.ek_dosyalar_json = json.dumps(ekler, ensure_ascii=False) if ekler else None
                    db.commit()
                    if n:
                        log_action(db, "Oluşturma", detay=f"Teknik resim ek dosya: {t.resim_no} rev {r.revizyon} — {n} dosya ({tur})")
                        flash(f"{n} ek dosya eklendi.", "success")
                elif a == "ek_sil":
                    if not _yonetici():
                        abort(403)
                    ekler = _json(r.ek_dosyalar_json) or []
                    i = f.get("ek_i", type=int)
                    if i is None or not 0 <= i < len(ekler):
                        abort(404)
                    e = ekler.pop(i)
                    r.ek_dosyalar_json = json.dumps(ekler, ensure_ascii=False) if ekler else None
                    db.commit(); log_action(db, "Silme", detay=f"Teknik resim ek dosyası kaldırıldı: {t.resim_no} rev {r.revizyon} — {e.get('ad')}")
                    flash(f"Ek dosya kaldırıldı: {e.get('ad')}", "success")
                elif a == "rev_sil":
                    if not _yonetici():
                        abort(403)
                    ad = r.revizyon
                    gecerliydi = r.durum == "Geçerli"
                    if r.dosya_yolu:
                        p = resolve_file_path(r.dosya_yolu, r.dosya_adi)
                        if p and os.path.isfile(p):
                            try:
                                os.remove(p)
                            except OSError:
                                pass
                    db.delete(r); db.flush()
                    if gecerliydi:      # yanlış yüklenen geçerli revizyon silinince bir öncekini geri al
                        kalan = _revler(db, t.id)
                        if kalan and not _gecerli(kalan):
                            kalan[-1].durum, kalan[-1].gecersiz_tarihi = "Geçerli", None
                    db.commit(); log_action(db, "Silme", detay=f"Teknik resim revizyonu silindi: {t.musteri} / {t.resim_no} rev {ad}")
                    flash(f"Rev {ad} silindi.", "success")
                return redirect(git + (f"#rev-{r.id}" if a != "rev_sil" else ""))
            return redirect(git)

        revler = _revler(db, t.id)
        g = _gecerli(revler)
        kullanicilar = {u.id: u.ad_soyad for u in db.query(User).all()}
        rv = [{"r": r, "dag": _json(r.dagitim_json), "deg": _json(r.degerlendirme_json), "uyarilar": _rev_durum(r),
               "ekler": _json(r.ek_dosyalar_json) or [],
               "ext": _ext(r.dosya_adi), "yukleyen": kullanicilar.get(r.yukleyen_id, "—")} for r in reversed(revler)]
        return render_template("teknik_resim_detay.html", t=t, revler=rv, gecerli=g, dagitim=TR_DAGITIM, etkiler=TR_ETKILER,
                               yazabilir=_yazabilir(), yonetici=_yonetici(), musteriler=_musteri_onerileri(db), gun=DEGERLENDIRME_GUN,
                               goruntu=GORUNTU, ek_turleri=TR_EK_TURLERI, uzantilar=", ".join(sorted(UZANTILAR)), today=date.today(),
                               olusturan=kullanicilar.get(t.olusturan_id, "—"))
    finally:
        db.close()


@app.route("/teknik-resim/rev/<int:rev_id>/dosya")
@login_required
def teknik_resim_dosya(rev_id):
    db = get_db()
    try:
        r = db.get(TeknikResimRevizyon, rev_id)
        if not r or not r.dosya_yolu:
            abort(404)
        p = resolve_file_path(r.dosya_yolu, r.dosya_adi)
        if not p or not os.path.isfile(p):
            flash("Resim dosyası bulunamadı.", "error")
            return redirect(url_for("teknik_resim_detay", rid=r.resim_id))
        log_action(db, "Görüntüleme", detay=f"Teknik resim dosyası: rev #{r.id} {r.dosya_adi}")
        return send_file(p, as_attachment=_ext(r.dosya_adi) not in GORUNTU or request.args.get("indir") == "1", download_name=r.dosya_adi or "resim")
    finally:
        db.close()


@app.route("/teknik-resim/rev/<int:rev_id>/ek/<int:i>")
@login_required
def teknik_resim_ek(rev_id, i):
    db = get_db()
    try:
        r = db.get(TeknikResimRevizyon, rev_id)
        ekler = (_json(r.ek_dosyalar_json) or []) if r else []
        if not 0 <= i < len(ekler):
            abort(404)
        e = ekler[i]
        p = resolve_file_path(e.get("yol"), e.get("ad"))
        if not p or not os.path.isfile(p):
            flash("Ek dosya bulunamadı.", "error")
            return redirect(url_for("teknik_resim_detay", rid=r.resim_id))
        return send_file(p, as_attachment=_ext(e.get("ad")) not in GORUNTU or request.args.get("indir") == "1", download_name=e.get("ad") or "ek")
    finally:
        db.close()


@app.route("/teknik-resim/excel")
@login_required
def teknik_resim_excel():
    """D04.4 F02 Teknik Resim Takip Listesi düzeninde: resim adı, no, müşteri, geçerli revizyon, teslim alma, dağıtım, önceki revizyonlar."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    db = get_db()
    try:
        musteri = (request.args.get("musteri") or "").strip()
        resimler = [t for t in db.query(TeknikResim).all() if t.aktif is not False and (not musteri or _k(t.musteri) == _k(musteri))]
        resimler.sort(key=lambda t: (_k(t.musteri), _k(t.urun_adi)))
        wb = Workbook(); ws = wb.active; ws.title = "Teknik Resim Takip"
        bas = ["Müşteri", "Resim Adı (Ürün)", "DEBAK Ürün Kodu", "Resim No", "Geçerli Revizyon", "Revizyon Tarihi", "Teslim Alma Tarihi",
               *[f"Dağıtım — {b}" for b in TR_DAGITIM], "Değerlendirme", "Etkilenenler", "Önceki Revizyonlar (teslim alma)"]
        ws.append([f"TEKNİK RESİM TAKİP LİSTESİ (D04.4 F02) — {datetime.now():%d.%m.%Y}" + (f" — {musteri}" if musteri else "")])
        ws["A1"].font = Font(bold=True, size=13)
        ws.append(bas)
        for c in ws[2]:
            c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F4E78"); c.alignment = Alignment(wrap_text=True, vertical="center")
        f = lambda d: d.strftime("%d.%m.%Y") if d else ""
        fi = lambda s: datetime.strptime(s, "%Y-%m-%d").strftime("%d.%m.%Y") if s else ""
        for t in resimler:
            rl = _revler(db, t.id); g = _gecerli(rl)
            dag = _json(g.dagitim_json) if g else {}
            deg = _json(g.degerlendirme_json) if g else {}
            ws.append([t.musteri, t.urun_adi, t.urun_kodu or "", t.resim_no, g.revizyon if g else "—", f(g.revizyon_tarihi) if g else "",
                       f(g.teslim_alma_tarihi) if g else "", *[fi((dag.get(b) or {}).get("tarih")) for b in TR_DAGITIM],
                       fi(deg.get("tarih")) or "Bekliyor", ", ".join(deg.get("etkiler") or []) or ("Etki yok" if deg.get("tarih") else ""),
                       " · ".join(f"{r.revizyon} ({f(r.teslim_alma_tarihi)})" for r in rl if r is not g)])
        for i, w in enumerate([18, 30, 16, 18, 12, 13, 13, 12, 12, 12, 12, 13, 30, 40], 1):
            ws.column_dimensions[ws.cell(2, i).column_letter].width = w
        ws.freeze_panes = "A3"; ws.auto_filter.ref = f"A2:{ws.cell(2, len(bas)).column_letter}{ws.max_row}"
        bio = io.BytesIO(); wb.save(bio); bio.seek(0)
        log_action(db, "Görüntüleme", detay=f"Teknik resim listesi Excel ({len(resimler)} resim)")
        return send_file(bio, as_attachment=True, download_name=f"D04.4 F02 Teknik Resim Takip Listesi {date.today():%Y%m%d}.xlsx",
                         mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    finally:
        db.close()
