"""
DYS — Stratejik Planlama: Excel (Y01 F02 / F03) okuma ve içe aktarma (2026)
  excel_oku(f02, f03, yil)   → {"yil", "gozden_gecirme", "kalemler": [...]} (paket JSON'u)
  paketten_yukle(db, yol)    → paket JSON'unu dönem + kalemler olarak yazar (aynı yılın kalemleri yenilenir)
Birleştirme: F02 (kurumsal bakış) ve F03 (İSG & Çevre bakışı) aynı yılın tek analizine 'kapsam' etiketiyle girer;
aynı ada sahip paydaşlar tek satırda birleşir (F03'ün statü / etki / İSG-çevre beklentisi eklenir), aynı stratejiler tekrarlanmaz.
Excel dosyalarına yazılmaz; kopyadan salt okunur açılır.
"""
import json
import os
import re
import shutil
import tempfile
from datetime import date, datetime

SWOT_BASLIK = {"GÜÇLÜ YANLAR": "Güçlü", "ZAYIF YANLAR": "Zayıf", "FIRSATLAR": "Fırsat", "TEHDİTLER": "Tehdit"}
_MADDE_RE = re.compile(r"^\s*(\d+)\s*\)\s*")
_TARIH_RE = re.compile(r"(\d{1,2})\.(\d{1,2})\.(\d{4})")


def _s(v):
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    return re.sub(r"[ \t]+", " ", str(v)).strip()


def _U(s):
    """Türkçe büyük harf (i → İ, ı → I)."""
    return (s or "").replace("i", "İ").replace("ı", "I").upper()


def _tarih(metin):
    m = _TARIH_RE.search(metin or "")
    return f"{m.group(3)}-{int(m.group(2)):02d}-{int(m.group(1)):02d}" if m else None


def _norm(ad):
    return re.sub(r"[^a-z0-9çğıöşü]", "", (ad or "").lower().replace("İ", "i").replace("I", "ı"))


def _ac(yol):
    import openpyxl
    kopya = os.path.join(tempfile.mkdtemp(), os.path.basename(yol))
    shutil.copy2(yol, kopya)
    return openpyxl.load_workbook(kopya, data_only=True)


def _gg(ws):
    for row in ws.iter_rows(min_row=1, max_row=3):
        for c in row:
            if c.value and "Gözden Geçirme" in str(c.value) or (c.value and "Gözden geçirme" in str(c.value)):
                t = _tarih(str(c.value))
                if t:
                    return t
    return None


# ---------------- F02 ----------------
def _baglam_f02(ws, kapsam):
    out, bilesen = [], None
    for r in range(4, ws.max_row + 1):
        c, konu = _s(ws.cell(r, 3).value), _s(ws.cell(r, 5).value)
        if c:
            bilesen = c
        if not konu:
            continue
        out.append({"bolum": "baglam", "kapsam": kapsam, "grup": bilesen, "metin": konu,
                    "veri": {"kaynak": _s(ws.cell(r, 6).value), "analiz": _s(ws.cell(r, 7).value)}})
    return out


def _paydas_f02(ws, kapsam):
    out = []
    for r in range(4, ws.max_row + 1):
        ad = _s(ws.cell(r, 3).value)
        if ad:
            out.append({"bolum": "paydas", "kapsam": kapsam, "metin": ad,
                        "veri": {"faaliyet_etkisi": _s(ws.cell(r, 4).value), "beklenti": _s(ws.cell(r, 5).value)}})
    return out


def _swot_ve_strateji(ws, kapsam):
    """SWOT sayfası: iki sütunlu G/Z ve F/T maddeleri (çok satıra bölünmüş maddeler birleştirilir),
    'STRATEJİ …' tablosu ve 'Önceliklerine Göre Sıralanmış Stratejiler' tablosu."""
    satirlar = [[c for c in row] for row in ws.iter_rows(min_row=1, max_row=ws.max_row)]
    def bul(metin, bas=0):
        for i in range(bas, len(satirlar)):
            for c in satirlar[i]:
                if c.value and metin in _U(_s(c.value)):
                    return i, c.column
        return None, None
    gz, _ = bul("GÜÇLÜ YANLAR"); ft, _ = bul("FIRSATLAR"); st, _ = bul("STRATEJİ", (ft or 0) + 1)
    on, _ = bul("ÖNCELİKLERİNE GÖRE")
    if gz is None or ft is None:
        return [], []
    sol = next(c.column for c in satirlar[gz] if c.value and "GÜÇLÜ" in _U(_s(c.value)))
    sag = next(c.column for c in satirlar[gz] if c.value and "ZAYIF" in _U(_s(c.value)))
    swot = []
    def topla(bas, bit, kol, tur):
        madde = None
        for i in range(bas, bit):
            v = _s(satirlar[i][kol - 1].value) if kol - 1 < len(satirlar[i]) else ""
            if not v or v == "Durum":
                continue
            m = _MADDE_RE.match(v)
            if m:
                madde = {"bolum": "swot", "kapsam": kapsam, "grup": tur, "kod": m.group(1), "metin": v[m.end():].strip()}
                swot.append(madde)
            elif madde:
                madde["metin"] = (madde["metin"] + " " + v).strip()
    son = st if st is not None else len(satirlar)
    topla(gz + 1, ft, sol, "Güçlü"); topla(gz + 1, ft, sag, "Zayıf")
    topla(ft + 1, son, sol, "Fırsat"); topla(ft + 1, son, sag, "Tehdit")
    stratejiler, eylemler = [], []
    if st is not None:
        bas_satir = satirlar[st + 1] if st + 1 < len(satirlar) else []
        pay_kol = next((c.column for c in bas_satir if c.value and "Paydaş" in _s(c.value)), None)
        for i in range(st + 2, on if on is not None else len(satirlar)):
            no, metin = _s(satirlar[i][1].value), _s(satirlar[i][2].value)
            if not metin:
                continue
            stratejiler.append({"bolum": "strateji", "kapsam": kapsam, "kod": no, "metin": metin,
                                "veri": {"ilgili_paydas": _s(satirlar[i][pay_kol - 1].value) if pay_kol else ""}})
    if on is not None:
        bas = {}
        for c in satirlar[on]:
            t = _s(c.value)
            for anahtar, ad in (("Kritik", "kbf"), ("Termin", "termin"), ("Kaynak", "kaynak"), ("Gözden", "gg"), ("Gerçekleşen", "gerceklesen")):
                if anahtar in t and ad not in bas:
                    bas[ad] = (c.column, t)
        gg_tarih = _tarih(bas.get("gg", (0, ""))[1])
        onceki = None
        for i in range(on + 1, len(satirlar)):
            al = lambda k: _s(satirlar[i][bas[k][0] - 1].value) if k in bas and bas[k][0] - 1 < len(satirlar[i]) else ""
            no, metin = _s(satirlar[i][1].value), _s(satirlar[i][2].value)
            if not metin:
                continue
            if not no and onceki:            # numarasız devam satırı → önceki stratejiye alt madde
                onceki["metin"] += " — " + metin
                continue
            termin = al("termin")
            onceki = {"bolum": "eylem", "kapsam": kapsam, "kod": no, "metin": metin,
                      "termin": termin if re.match(r"^\d{4}-\d{2}-\d{2}$", termin) else None,
                      "veri": {"kbf": al("kbf"), "kaynak": al("kaynak"), "gg_notu": al("gg"), "gg_tarihi": gg_tarih,
                               "gerceklesen": al("gerceklesen"), "termin_metin": "" if re.match(r"^\d{4}-\d{2}-\d{2}$", termin) else termin}}
            g = onceki["veri"]["gerceklesen"]
            onceki["durum"] = "Tamamlandı" if re.match(r"^\d{4}-\d{2}-\d{2}$", g) else "Devam Ediyor"
            eylemler.append(onceki)
    return swot, stratejiler + eylemler


# ---------------- F03 (İSG & Çevre) ----------------
def _baglam_f03(ws, kapsam):
    out, bilesen, madde = [], None, None
    esle = {"PAZAR": "PAZAR VE RAKİPLER"}
    for r in range(4, ws.max_row + 1):
        v = _s(ws.cell(r, 1).value)
        if not v or v.startswith("İç / Dış"):
            continue
        if v.isupper() and not _MADDE_RE.match(v):
            bilesen = esle.get(v, v); madde = None
            continue
        m = _MADDE_RE.match(v)
        if m:
            madde = {"bolum": "baglam", "kapsam": kapsam, "grup": bilesen, "metin": v[m.end():].strip(), "veri": {"kaynak": "", "analiz": ""}}
            out.append(madde)
        elif madde:
            madde["metin"] = (madde["metin"] + " " + v).strip()
    return out


def _paydas_f03(ws, kapsam):
    out = []
    for r in range(4, ws.max_row + 1):
        ad, statu = _s(ws.cell(r, 1).value), _s(ws.cell(r, 2).value)
        if not ad or not statu:
            continue
        out.append({"bolum": "paydas", "kapsam": kapsam, "metin": ad,
                    "veri": {"statu": statu, "faaliyet_etkisi": _s(ws.cell(r, 3).value), "etki": _s(ws.cell(r, 4).value),
                             "isg_cevre_beklenti": _s(ws.cell(r, 5).value)}})
    return out


def _birlestir(kalemler):
    """Aynı yılın F02 + F03 kalemleri: aynı paydaş tek satır, aynı strateji / öncelikli strateji tek satır."""
    out, paydas, strateji = [], {}, {}
    for k in kalemler:
        if k["bolum"] == "paydas":
            n = _norm(k["metin"])
            aday = paydas.get(n) or next((v for kk, v in paydas.items() if n and (n in kk or kk in n) and min(len(n), len(kk)) > 6), None)
            if aday:
                for a in ("statu", "etki", "isg_cevre_beklenti"):
                    if k["veri"].get(a) and not aday["veri"].get(a):
                        aday["veri"][a] = k["veri"][a]
                if not aday["veri"].get("faaliyet_etkisi"):
                    aday["veri"]["faaliyet_etkisi"] = k["veri"].get("faaliyet_etkisi", "")
                if aday["kapsam"] != k["kapsam"]:
                    aday["kapsam"] = "Kurumsal"; aday["veri"]["ortak"] = True
                continue
            paydas[n] = k
        if k["bolum"] in ("strateji", "eylem"):
            n = (k["bolum"], k["kapsam"], _norm(k["metin"])[:60])
            ayni = next((v for kk, v in strateji.items() if kk[0] == n[0] and kk[2] == n[2] and kk[1] != n[1]), None)
            if ayni:                       # F02 ve F03'te aynı strateji → tek satır, iki kapsamda ortak
                ayni["veri"]["ortak"] = True
                continue
            strateji[n] = k
        out.append(k)
    sira = {}
    for k in out:
        anahtar = (k["bolum"], k.get("grup"), k["kapsam"] if k["bolum"] in ("swot", "baglam") else "")
        sira[anahtar] = sira.get(anahtar, 0) + 1
        k["sira"] = sira[anahtar]
    return out


def excel_oku(f02, f03=None, yil=2026, f02_swot=None, f03_swot=None):
    """f02_swot / f03_swot: kullanılacak SWOT sayfası adı (verilmezse o yılın ya da en son yılın sayfası)."""
    kalemler, gg = [], {}
    wb = _ac(f02)
    def sayfa(wb, ad_bas, tercih):
        adlar = [s for s in wb.sheetnames if s.upper().startswith(ad_bas.upper())]
        if tercih and tercih in wb.sheetnames:
            return wb[tercih]
        yilli = [s for s in adlar if str(yil) in s] or sorted(adlar, key=lambda s: re.findall(r"\d{4}", s) or ["0"])[-1:]
        return wb[yilli[0]] if yilli else None
    ws = next((wb[s] for s in wb.sheetnames if "BAĞLAM" in s.upper()), None)
    if ws:
        kalemler += _baglam_f02(ws, "Kurumsal"); gg["baglam"] = _gg(ws)
    ws = next((wb[s] for s in wb.sheetnames if "PAYDAŞ" in s.upper()), None)
    if ws:
        kalemler += _paydas_f02(ws, "Kurumsal"); gg["paydas"] = _gg(ws)
    ws = sayfa(wb, "SWOT", f02_swot)
    if ws:
        s, st = _swot_ve_strateji(ws, "Kurumsal"); kalemler += s + st; gg["swot"] = _gg(ws); gg["swot_sayfa"] = ws.title
    if f03 and os.path.exists(f03):
        wb3 = _ac(f03)
        kba = [s for s in wb3.sheetnames if s.upper().startswith("KBA")]
        ws = wb3[sorted(kba, key=lambda s: _gg(wb3[s]) or "")[-1]] if kba else None
        if ws:
            kalemler += _baglam_f03(ws, "İSG & Çevre"); gg["baglam_isg_cevre"] = _gg(ws)
        ws = next((wb3[s] for s in wb3.sheetnames if "PAYDAŞ" in s.upper()), None)
        if ws:
            kalemler += _paydas_f03(ws, "İSG & Çevre"); gg["paydas_isg_cevre"] = _gg(ws)
        ws = sayfa(wb3, "SWOT", f03_swot)
        if ws:
            s, st = _swot_ve_strateji(ws, "İSG & Çevre"); kalemler += s + st; gg["swot_isg_cevre"] = _gg(ws); gg["swot_isg_cevre_sayfa"] = ws.title
    return {"yil": yil, "gozden_gecirme": {k: v for k, v in gg.items() if v}, "kaynak": [os.path.basename(f02)] + ([os.path.basename(f03)] if f03 else []),
            "kalemler": _birlestir(kalemler)}


def paketten_yukle(db, yol):
    """Paket JSON'unu yazar. Aynı yılın kalemleri silinip yeniden yazılır (tekrar çalıştırılabilir). Dönüş: özet sayılar."""
    from models import StratejiDonem, StratejiKalem
    p = json.load(open(yol, encoding="utf-8")) if isinstance(yol, str) else yol
    yil = int(p["yil"])
    d = db.query(StratejiDonem).filter(StratejiDonem.yil == yil).first() or StratejiDonem(yil=yil)
    d.gozden_gecirme_json = json.dumps(p.get("gozden_gecirme") or {}, ensure_ascii=False)
    d.kaynak = ", ".join(p.get("kaynak") or [])[:300]
    db.add(d)
    db.query(StratejiKalem).filter(StratejiKalem.yil == yil).delete()
    say = {}
    for k in p["kalemler"]:
        t = k.get("termin")
        db.add(StratejiKalem(yil=yil, bolum=k["bolum"], kapsam=k.get("kapsam") or "Kurumsal", grup=k.get("grup"), sira=k.get("sira"),
                             kod=(k.get("kod") or None), metin=k["metin"], veri_json=json.dumps(k.get("veri") or {}, ensure_ascii=False),
                             termin=datetime.strptime(t, "%Y-%m-%d").date() if t else None, durum=k.get("durum")))
        say[k["bolum"]] = say.get(k["bolum"], 0) + 1
    db.commit()
    return {"yil": yil, **say}
