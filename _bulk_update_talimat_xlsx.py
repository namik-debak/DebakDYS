# -*- coding: utf-8 -*-
"""Update talimat .xlsx header/footer via openpyxl (format preserved)."""
from __future__ import annotations

import re
import shutil
import sqlite3
import traceback
from datetime import datetime
from pathlib import Path

import openpyxl

from _talimat_meta_utils import (
    build_meta_line_from_old,
    fmt_date,
    norm,
    paragraph_has_meta,
    strip_excel_footer_codes,
)

DB = r"\\192.168.0.249\ortak\HARUN\dys\dys.db"
XLSX_LIST = r"c:\Users\labaratuvar03\Downloads\ana_dokuman_listesi2-.xlsx"
BASE = Path(r"D:\UYGULAMA PROJELER\debak dijital\dys")
LOG = BASE / "_talimat_xlsx_log.txt"


def load_excel():
    wb = openpyxl.load_workbook(XLSX_LIST, data_only=True)
    ws = wb["Ana Doküman Listesi"]
    out = {}
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 3).value == "T":
            kod = ws.cell(r, 2).value
            if kod:
                out[norm(kod)] = {
                    "kod": str(kod).strip(),
                    "rev": ws.cell(r, 5).value,
                    "tarih": ws.cell(r, 6).value,
                }
    return out


def resolve_path(p):
    if not p:
        return None
    path = Path(p)
    if path.exists():
        return path
    p2 = p.replace("/", "\\")
    if "uploads\\documents" in p2:
        rel = p2.split("uploads\\documents", 1)[1].lstrip("\\/")
        c = BASE / "uploads" / "documents" / rel
        return c if c.exists() else None
    return None


def update_hf_part(part, kod, rev, tarih):
    if part is None:
        return False, None
    txt = (part.text or "").strip()
    plain = strip_excel_footer_codes(txt)
    if not paragraph_has_meta(plain):
        return False, None
    new = build_meta_line_from_old(kod, rev, tarih, plain)
    # Excel kodlarını koru (&P, &F vb.)
    codes = re.findall(r"&[A-Za-z0-9+-]+", txt)
    if codes and plain != txt:
        # meta + kod birlikteyse sadece meta kısmını değiştir
        part.text = txt.replace(plain, new)
    else:
        part.text = new
    return True, new


def update_worksheet(ws, kod, rev, tarih):
    updated = False
    last = None
    for hf in (ws.oddHeader, ws.evenHeader, ws.firstHeader):
        for part in (hf.left, hf.center, hf.right):
            ok, line = update_hf_part(part, kod, rev, tarih)
            if ok:
                updated = True
                last = line
    for hf in (ws.oddFooter, ws.evenFooter, ws.firstFooter):
        for part in (hf.left, hf.center, hf.right):
            ok, line = update_hf_part(part, kod, rev, tarih)
            if ok:
                updated = True
                last = line
    if not updated:
        new = build_meta_line_from_old(kod, rev, tarih, "")
        ws.oddFooter.right.text = new
        ws.evenFooter.right.text = new
        updated = True
        last = new
    return updated, last


def process_file(path: Path, meta):
    bak = path.with_suffix(path.suffix + ".bak_meta")
    if not bak.exists():
        shutil.copy2(path, bak)
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    ok, line = update_worksheet(ws, meta["kod"], meta["rev"], meta["tarih"])
    if ok:
        wb.save(path)
    return line


def main():
    excel = load_excel()
    conn = sqlite3.connect(DB)
    cur = conn.cursor()
    cur.execute(
        "SELECT dokuman_no, baslik, dosya_yolu FROM documents WHERE dokuman_tipi='Talimat'"
    )
    stats = {"updated": 0, "skip": 0, "err": 0}
    lines = []
    for no, baslik, p in cur.fetchall():
        key = norm(no)
        if key not in excel:
            continue
        fp = resolve_path(p)
        if not fp or fp.suffix.lower() != ".xlsx":
            continue
        try:
            line = process_file(fp, excel[key])
            stats["updated"] += 1
            lines.append(f"OK\t{no}\t{line}")
        except Exception as e:
            stats["err"] += 1
            lines.append(f"ERR\t{no}\t{e}")
    conn.close()
    summary = f"xlsx updated={stats['updated']} errors={stats['err']}"
    LOG.write_text(summary + "\n" + "\n".join(lines), encoding="utf-8")
    print(summary)


if __name__ == "__main__":
    main()
