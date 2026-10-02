"""
DYS — PDF / Excel Dışa Aktarma
==============================
PDF: ReportLab (Windows'ta harici ikili gerektirmez).
Excel: openpyxl (mevcut bağımlılık).
Her fonksiyon bir BytesIO döndürür; Flask send_file ile sunulur.
"""

from __future__ import annotations

import io
import os
from datetime import datetime
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import cm
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, HRFlowable, KeepTogether,
)

_FONT_REG = "DysSans"
_FONT_BOLD = "DysSans-Bold"
_FONTS_READY = False


def _register_fonts():
    """Türkçe karakter destekli TrueType fontları kaydeder (Windows Arial/Calibri)."""
    global _FONTS_READY, _FONT_REG, _FONT_BOLD
    if _FONTS_READY:
        return
    windir = os.environ.get("WINDIR", r"C:\Windows")
    pairs = [
        (os.path.join(windir, "Fonts", "arial.ttf"),
         os.path.join(windir, "Fonts", "arialbd.ttf")),
        (os.path.join(windir, "Fonts", "calibri.ttf"),
         os.path.join(windir, "Fonts", "calibrib.ttf")),
        (os.path.join(windir, "Fonts", "segoeui.ttf"),
         os.path.join(windir, "Fonts", "segoeuib.ttf")),
    ]
    for regular, bold in pairs:
        if os.path.isfile(regular) and os.path.isfile(bold):
            pdfmetrics.registerFont(TTFont(_FONT_REG, regular))
            pdfmetrics.registerFont(TTFont(_FONT_BOLD, bold))
            _FONTS_READY = True
            return
    # Fallback: Helvetica (Türkçe karakterler sorunlu olabilir)
    _FONT_REG = "Helvetica"
    _FONT_BOLD = "Helvetica-Bold"
    _FONTS_READY = True


def _esc(text):
    return escape("" if text is None else str(text))


def _p(text, style):
    return Paragraph(_esc(text).replace("\n", "<br/>"), style)


# ── Excel ────────────────────────────────────────────────────────────────────
def excel_from_rows(baslik, kolonlar, satirlar):
    """
    Basit bir Excel çıktısı üretir.
    kolonlar: ["Sütun A", ...]
    satirlar: [[hücre, ...], ...]
    """
    wb = Workbook()
    ws = wb.active
    safe_title = (baslik or "Rapor").replace("/", "-").replace("\\", "-")[:31]
    ws.title = safe_title or "Rapor"

    ws.append([baslik])
    ws["A1"].font = Font(size=14, bold=True)
    ws.append([f"Oluşturma: {datetime.now():%d.%m.%Y %H:%M}"])
    ws.append([])

    header_row = ws.max_row + 1
    ws.append(kolonlar)
    fill = PatternFill("solid", fgColor="1F4E78")
    for col_idx in range(1, len(kolonlar) + 1):
        cell = ws.cell(row=header_row, column=col_idx)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")

    for satir in satirlar:
        ws.append(["" if h is None else h for h in satir])

    for col_idx, _kol in enumerate(kolonlar, start=1):
        max_len = len(str(kolonlar[col_idx - 1]))
        for satir in satirlar:
            if col_idx - 1 < len(satir) and satir[col_idx - 1] is not None:
                max_len = max(max_len, len(str(satir[col_idx - 1])))
        ws.column_dimensions[ws.cell(row=header_row, column=col_idx).column_letter].width = min(max_len + 4, 60)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _pdf_styles():
    _register_fonts()
    styles = getSampleStyleSheet()
    return {
        "title": ParagraphStyle(
            "DysTitle", parent=styles["Normal"],
            fontName=_FONT_BOLD, fontSize=16, leading=20,
            textColor=colors.HexColor("#0f172a"), alignment=TA_LEFT, spaceAfter=2,
        ),
        "subtitle": ParagraphStyle(
            "DysSubtitle", parent=styles["Normal"],
            fontName=_FONT_REG, fontSize=9, leading=12,
            textColor=colors.HexColor("#64748b"), spaceAfter=8,
        ),
        "hucre": ParagraphStyle(
            "DysHucre", parent=styles["Normal"],
            fontName=_FONT_REG, fontSize=9, leading=12,
            textColor=colors.HexColor("#1e293b"),
        ),
        "header": ParagraphStyle(
            "DysHeader", parent=styles["Normal"],
            fontName=_FONT_BOLD, fontSize=9, leading=11,
            textColor=colors.white,
        ),
        "label": ParagraphStyle(
            "DysLabel", parent=styles["Normal"],
            fontName=_FONT_BOLD, fontSize=9, leading=12,
            textColor=colors.HexColor("#334155"),
        ),
        "section": ParagraphStyle(
            "DysSection", parent=styles["Normal"],
            fontName=_FONT_BOLD, fontSize=11, leading=14,
            textColor=colors.HexColor("#0f172a"), spaceBefore=10, spaceAfter=4,
        ),
        "meta_label": ParagraphStyle(
            "DysMetaLabel", parent=styles["Normal"],
            fontName=_FONT_REG, fontSize=8, leading=10,
            textColor=colors.HexColor("#64748b"),
        ),
        "meta_value": ParagraphStyle(
            "DysMetaValue", parent=styles["Normal"],
            fontName=_FONT_BOLD, fontSize=10, leading=13,
            textColor=colors.HexColor("#0f172a"),
        ),
        "body": ParagraphStyle(
            "DysBody", parent=styles["Normal"],
            fontName=_FONT_REG, fontSize=9.5, leading=13,
            textColor=colors.HexColor("#1e293b"),
        ),
        "brand": ParagraphStyle(
            "DysBrand", parent=styles["Normal"],
            fontName=_FONT_BOLD, fontSize=12, leading=14,
            textColor=colors.white,
        ),
        "brand_sub": ParagraphStyle(
            "DysBrandSub", parent=styles["Normal"],
            fontName=_FONT_REG, fontSize=8, leading=10,
            textColor=colors.HexColor("#bfdbfe"), alignment=TA_RIGHT,
        ),
    }


# ── PDF (genel tablo) ──────────────────────────────────────────────────────
def pdf_from_rows(baslik, kolonlar, satirlar, alt_bilgi=None, yatay=True):
    """Başlık + tablo içeren bir PDF üretir (Unicode font)."""
    _register_fonts()
    st = _pdf_styles()
    buf = io.BytesIO()
    sayfa = landscape(A4) if yatay else A4
    doc = SimpleDocTemplate(
        buf, pagesize=sayfa,
        leftMargin=1.4 * cm, rightMargin=1.4 * cm,
        topMargin=1.4 * cm, bottomMargin=1.4 * cm,
        title=baslik or "Rapor",
    )

    elemanlar = [
        Paragraph(_esc(baslik), st["title"]),
        Paragraph(f"Oluşturma: {datetime.now():%d.%m.%Y %H:%M}", st["subtitle"]),
        Spacer(1, 0.25 * cm),
    ]

    data = [[_p(k, st["header"]) for k in kolonlar]]
    for satir in satirlar:
        data.append([_p("" if h is None else h, st["hucre"]) for h in satir])

    # 2 sütunlu alan/değer tablolarında sabit genişlik (dikey harf kırılmasını önler)
    col_widths = None
    if len(kolonlar) == 2:
        page_w = sayfa[0] - 2.8 * cm
        col_widths = [page_w * 0.32, page_w * 0.68]

    tablo = Table(data, repeatRows=1, colWidths=col_widths)
    tablo.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e3a5f")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), _FONT_BOLD),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f1f5f9")]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
    ]))
    elemanlar.append(tablo)

    if alt_bilgi:
        elemanlar.append(Spacer(1, 0.4 * cm))
        elemanlar.append(_p(alt_bilgi, st["subtitle"]))

    doc.build(elemanlar)
    buf.seek(0)
    return buf


def _section_block(title, body, st):
    """Başlıklı içerik kutusu."""
    inner = Table(
        [[_p(title, st["label"])], [_p(body or "—", st["body"])]],
        colWidths=["*"],
    )
    inner.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e2e8f0")),
        ("BACKGROUND", (0, 1), (-1, 1), colors.HexColor("#f8fafc")),
        ("BOX", (0, 0), (-1, -1), 0.6, colors.HexColor("#cbd5e1")),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    return KeepTogether([inner, Spacer(1, 0.35 * cm)])


def capa_detail_pdf(capa, kok_bloklar=None):
    """
    Tek DÖF için şık, okunaklı dikey PDF raporu.
    capa: CorrectiveAction (ilişkiler yüklü olmalı)
    kok_bloklar: [(başlık, içerik), ...]
    """
    _register_fonts()
    st = _pdf_styles()
    buf = io.BytesIO()
    page = A4
    doc = SimpleDocTemplate(
        buf, pagesize=page,
        leftMargin=1.5 * cm, rightMargin=1.5 * cm,
        topMargin=1.2 * cm, bottomMargin=1.4 * cm,
        title=f"DÖF {capa.dof_no}",
    )
    page_w = page[0] - 3 * cm

    # Üst bant
    header = Table(
        [[
            Paragraph("debak · DYS", st["brand"]),
            Paragraph(
                f"Düzeltici / Önleyici Faaliyet<br/>{datetime.now():%d.%m.%Y %H:%M}",
                st["brand_sub"],
            ),
        ]],
        colWidths=[page_w * 0.45, page_w * 0.55],
    )
    header.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#0f2744")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 12),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 12),
        ("LEFTPADDING", (0, 0), (-1, -1), 12),
        ("RIGHTPADDING", (0, 0), (-1, -1), 12),
    ]))

    title = Paragraph(_esc(f"{capa.dof_no} — {capa.baslik or ''}"), st["title"])
    durum = capa.durum or "—"
    durum_renk = {
        "Açık": "#b45309",
        "Devam Ediyor": "#1d4ed8",
        "Tamamlandı": "#047857",
        "Etkinlik Kontrolünde": "#7c3aed",
        "Kapatıldı": "#475569",
    }.get(durum, "#334155")

    badge = Table(
        [[Paragraph(
            f'<font color="white"><b>{_esc(durum)}</b></font>',
            ParagraphStyle("badge", fontName=_FONT_BOLD, fontSize=9, leading=11, alignment=TA_CENTER),
        )]],
        colWidths=[4.2 * cm],
    )
    badge.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(durum_renk)),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
    ]))

    title_row = Table([[title, badge]], colWidths=[page_w - 4.5 * cm, 4.5 * cm])
    title_row.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (1, 0), (1, 0), "RIGHT"),
    ]))

    # Meta ızgara (2x3)
    def meta_cell(label, value):
        return [
            _p(label, st["meta_label"]),
            _p(value or "—", st["meta_value"]),
        ]

    meta_data = [
        [
            meta_cell("Kaynak", capa.kaynak_tipi),
            meta_cell("Sorumlu", capa.sorumlu.ad_soyad if capa.sorumlu else None),
            meta_cell("Açan", capa.acan.ad_soyad if capa.acan else None),
        ],
        [
            meta_cell("Kök Neden Yöntemi", capa.kok_neden_yontemi),
            meta_cell("Tespit Tarihi", capa.tespit_tarihi.strftime("%d.%m.%Y") if capa.tespit_tarihi else None),
            meta_cell("Planlanan Kapanış", capa.planlanan_tarih.strftime("%d.%m.%Y") if capa.planlanan_tarih else None),
        ],
    ]
    # Flatten nested lists into table cells with Paragraphs
    meta_rows = []
    for row in meta_data:
        meta_rows.append([
            Table([[c[0]], [c[1]]], colWidths=["*"])
            for c in row
        ])

    meta = Table(meta_rows, colWidths=[page_w / 3.0] * 3)
    meta.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f8fafc")),
        ("BOX", (0, 0), (-1, -1), 0.8, colors.HexColor("#e2e8f0")),
        ("INNERGRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#e2e8f0")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 8),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
    ]))

    elemanlar = [
        header,
        Spacer(1, 0.45 * cm),
        title_row,
        Spacer(1, 0.35 * cm),
        meta,
        Spacer(1, 0.45 * cm),
        HRFlowable(width="100%", thickness=1, color=colors.HexColor("#e2e8f0"), spaceAfter=6),
    ]

    # İçerik bölümleri
    if capa.tanim:
        elemanlar.append(_section_block("Uygunsuzluk Tanımı", capa.tanim, st))

    if capa.kok_neden_yontemi == "8D":
        elemanlar.append(Paragraph("8D Adımları", st["section"]))
        eight_d = [
            ("D1 — Takım", capa.d1_team),
            ("D2 — Problem", capa.d2_problem),
            ("D3 — Geçici Önlem", capa.d3_containment),
            ("D4 — Kök Neden", capa.d4_root_cause),
            ("D5 — Kalıcı Aksiyon", capa.d5_corrective),
            ("D6 — Uygulama", capa.d6_implement),
            ("D7 — Önleme", capa.d7_prevent),
            ("D8 — Takdir", capa.d8_congratulate),
        ]
        d_rows = []
        pair = []
        for label, val in eight_d:
            cell = Table(
                [[_p(label, st["meta_label"])], [_p(val or "—", st["body"])]],
                colWidths=["*"],
            )
            cell.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f8fafc")),
                ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
            ]))
            pair.append(cell)
            if len(pair) == 2:
                d_rows.append(pair)
                pair = []
        if pair:
            pair.append("")
            d_rows.append(pair)
        d_table = Table(d_rows, colWidths=[page_w / 2 - 2, page_w / 2 - 2])
        d_table.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 2),
            ("RIGHTPADDING", (0, 0), (-1, -1), 2),
            ("TOPPADDING", (0, 0), (-1, -1), 2),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        elemanlar.append(d_table)
        elemanlar.append(Spacer(1, 0.25 * cm))
    elif kok_bloklar:
        elemanlar.append(Paragraph(
            f"Kök Neden Analizi ({capa.kok_neden_yontemi or '—'})", st["section"]
        ))
        for baslik, icerik in kok_bloklar:
            elemanlar.append(_section_block(baslik, icerik, st))

    if capa.duzeltici_faaliyet:
        elemanlar.append(_section_block("Düzeltici Faaliyet", capa.duzeltici_faaliyet, st))
    if capa.onleyici_faaliyet:
        elemanlar.append(_section_block("Önleyici Faaliyet", capa.onleyici_faaliyet, st))
    if capa.etkinlik_kontrolu:
        elemanlar.append(_section_block("Etkinlik Kontrolü", capa.etkinlik_kontrolu, st))

    # Alt bilgi
    elemanlar.append(Spacer(1, 0.5 * cm))
    elemanlar.append(HRFlowable(width="100%", thickness=0.6, color=colors.HexColor("#cbd5e1")))
    elemanlar.append(Paragraph(
        "Bu belge DYS üzerinden otomatik üretilmiştir.",
        st["subtitle"],
    ))

    doc.build(elemanlar)
    buf.seek(0)
    return buf
