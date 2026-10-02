"""
DYS — Jenerik Excel Şablon Üretimi + Doğrulama + Tek Transaction İçe Aktarım
=============================================================================
Herhangi bir modülün (Tedarikçiler, Personel, Kalibrasyon Ekipmanları vb.)
Excel ile toplu veri girişi ihtiyacını, TEK bir merkezi şema tanımından
besleyerek karşılayan yardımcı motor.

Kullanım deseni (3 adım):
  1) İlgili modül için bir ``ImportSemasi`` tanımlanır (sütunlar, tipler,
     zorunluluklar, enum değerleri, referans/benzersizlik kontrolleri).
  2) ``sablon_uret(sema)`` ile kullanıcıya indirilecek .xlsx şablonu üretilir
     ("Veri" sayfası boş, "Örnek" sayfası kilitli + dolu örnek satırlar).
  3) Kullanıcı doldurduğu dosyayı geri yüklediğinde ``dogrula(...)`` ile
     satır satır kontrol edilir; hata yoksa ``aktar(...)`` ile TEK veritabanı
     transaction'ında (ya hep ya hiç) ilgili modele yazılır.

Şablonu üreten ve doğrulayan kod AYNI ``ImportSemasi`` nesnesini okur —
alan tanımları iki ayrı yerde tekrar edilmez.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any, Callable, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

VERI_SAYFA_ADI = "Veri"
ORNEK_SAYFA_ADI = "Örnek"

# ── Görsel stiller (şablon üretiminde kullanılır) ───────────────────────────
_ZORUNLU_DOLGU = PatternFill(start_color="FFF9D9D9", end_color="FFF9D9D9", fill_type="solid")
_OPSIYONEL_DOLGU = PatternFill(start_color="FFDCE6F1", end_color="FFDCE6F1", fill_type="solid")
_BASLIK_FONT = Font(bold=True, size=11, color="FF1F2937")
_ORNEK_BASLIK_FONT = Font(bold=True, size=11, color="FFFFFFFF")
_ORNEK_BASLIK_DOLGU = PatternFill(start_color="FF6B7280", end_color="FF6B7280", fill_type="solid")
_INCE_KENARLIK = Border(*(Side(style="thin", color="FFD1D5DB") for _ in range(4)))


@dataclass
class SutunTanimi:
    """Excel içe aktarım şemasındaki tek bir sütunun tanımı.

    baslik            : Excel'de görünecek sütun başlığı.
    alan_adi          : DB kolonuna / model alanına eşlenecek iç ad.
    tip               : "metin" | "sayi" | "tarih" | "enum" | "boolean"
    zorunlu           : Boş bırakılamaz mı?
    maks_uzunluk      : Metin alanları için maksimum karakter sayısı.
    enum_degerleri    : tip="enum" ise izin verilen değerler tuple'ı.
    varsayilan        : Hücre boşsa kullanılacak varsayılan değer.
    ornek_deger       : "Örnek" sayfasında gösterilecek örnek değer.
    benzersiz         : Aynı sütunda dosya içinde tekrar eden değer hata mı?
    referans_kontrol  : (db, deger) -> hata_mesaji|None imzalı callable.
                        Örn. "bu süreç kodu processes tablosunda var mı?"
                        veya "bu kod veritabanında zaten kayıtlı mı?" gibi
                        veritabanına bakan kontroller için kullanılır.
    aciklama          : Kullanıcıya gösterilecek kısa yardım notu (opsiyonel,
                        şablonda başlık hücresinin yorumu olarak eklenmez,
                        yalnızca form/şablon açıklamalarında kullanılabilir).
    """

    baslik: str
    alan_adi: str
    tip: str = "metin"
    zorunlu: bool = False
    maks_uzunluk: Optional[int] = None
    enum_degerleri: Optional[tuple] = None
    varsayilan: Any = None
    ornek_deger: Any = None
    benzersiz: bool = False
    referans_kontrol: Optional[Callable[[Any, Any], Optional[str]]] = None
    aciklama: Optional[str] = None


@dataclass
class ImportSemasi:
    """Bir modülün toplu Excel içe aktarım şeması.

    ad             : Şemanın insan-okunur adı (örn. "Tedarikçiler").
    sutunlar       : SutunTanimi listesi (Excel'deki sütun sırasıyla aynı).
    ornek_satirlar : "Örnek" sayfasında gösterilecek dolu örnek satırlar;
                     her biri alan_adi -> deger sözlüğü.
    """

    ad: str
    sutunlar: list = field(default_factory=list)
    ornek_satirlar: list = field(default_factory=list)

    def alan_adlari(self):
        return [s.alan_adi for s in self.sutunlar]


class DogrulamaHatasi(Exception):
    """Toplu doğrulama sırasında en az bir satırda hata bulunduğunda kullanılır."""

    def __init__(self, hatalar):
        self.hatalar = hatalar  # [(satir_no, mesaj), ...]
        super().__init__(f"{len(hatalar)} satırda doğrulama hatası bulundu.")


# ═══════════════════════════════════════════════════════════════════════════
#  1) ŞABLON ÜRETİCİ
# ═══════════════════════════════════════════════════════════════════════════
def sablon_uret(sema: ImportSemasi) -> io.BytesIO:
    """Verilen şemadan iki sayfalı bir .xlsx şablonu üretir.

    - "Veri" sayfası: kullanıcının dolduracağı, sadece başlık satırı olan
      boş sayfa. Zorunlu sütunlar farklı arkaplan rengiyle işaretlenir.
    - "Örnek" sayfası: kilitli/salt-okunur, şemadaki örnek satırlarla dolu.
      Bu sayfadaki veriler asla "Veri" sayfasına karışmaz.

    Dönüş: bellek içi (BytesIO) .xlsx dosya akışı; çağıran taraf
    ``send_file(akis, ...)`` ile kullanıcıya sunabilir.
    """
    wb = Workbook()

    # ── "Veri" sayfası ──────────────────────────────────────────────────
    veri_ws = wb.active
    veri_ws.title = VERI_SAYFA_ADI
    _basliklari_yaz(veri_ws, sema, zorunlu_vurgula=True)
    _sutun_genisliklerini_ayarla(veri_ws, sema)
    _enum_dogrulama_ekle(veri_ws, sema, ilk_satir=2, son_satir=500)
    veri_ws.freeze_panes = "A2"

    # ── "Örnek" sayfası (kilitli) ───────────────────────────────────────
    ornek_ws = wb.create_sheet(ORNEK_SAYFA_ADI)
    _basliklari_yaz(ornek_ws, sema, zorunlu_vurgula=False, ornek_stili=True)
    _sutun_genisliklerini_ayarla(ornek_ws, sema)
    for r, satir in enumerate(sema.ornek_satirlar, start=2):
        for c, sutun in enumerate(sema.sutunlar, start=1):
            deger = satir.get(sutun.alan_adi, sutun.ornek_deger)
            hucre = ornek_ws.cell(row=r, column=c, value=deger)
            hucre.border = _INCE_KENARLIK
    ornek_ws.protection.sheet = True

    wb.active = 0
    akis = io.BytesIO()
    wb.save(akis)
    akis.seek(0)
    return akis


def _basliklari_yaz(ws, sema: ImportSemasi, zorunlu_vurgula: bool, ornek_stili: bool = False):
    for c, sutun in enumerate(sema.sutunlar, start=1):
        baslik_metni = sutun.baslik + (" *" if sutun.zorunlu else "")
        hucre = ws.cell(row=1, column=c, value=baslik_metni)
        hucre.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        hucre.border = _INCE_KENARLIK
        if ornek_stili:
            hucre.font = _ORNEK_BASLIK_FONT
            hucre.fill = _ORNEK_BASLIK_DOLGU
        else:
            hucre.font = _BASLIK_FONT
            hucre.fill = _ZORUNLU_DOLGU if (zorunlu_vurgula and sutun.zorunlu) else _OPSIYONEL_DOLGU
    ws.row_dimensions[1].height = 30


def _sutun_genisliklerini_ayarla(ws, sema: ImportSemasi):
    for c, sutun in enumerate(sema.sutunlar, start=1):
        genislik = max(len(sutun.baslik) + 4, (sutun.maks_uzunluk or 0) // 3, 14)
        ws.column_dimensions[get_column_letter(c)].width = min(genislik, 45)


def _enum_dogrulama_ekle(ws, sema: ImportSemasi, ilk_satir: int, son_satir: int):
    """Enum tipindeki sütunlara açılır liste (dropdown) doğrulaması ekler."""
    for c, sutun in enumerate(sema.sutunlar, start=1):
        if sutun.tip != "enum" or not sutun.enum_degerleri:
            continue
        formul = '"' + ",".join(str(v) for v in sutun.enum_degerleri) + '"'
        dv = DataValidation(type="list", formula1=formul, allow_blank=not sutun.zorunlu)
        dv.error = "Lütfen listeden geçerli bir değer seçin."
        dv.errorTitle = "Geçersiz Değer"
        harf = get_column_letter(c)
        dv.add(f"{harf}{ilk_satir}:{harf}{son_satir}")
        ws.add_data_validation(dv)


# ═══════════════════════════════════════════════════════════════════════════
#  2) DOĞRULAYICI
# ═══════════════════════════════════════════════════════════════════════════
def dogrula(dosya, sema: ImportSemasi, db=None):
    """Yüklenen .xlsx dosyasının "Veri" sayfasını şemaya göre doğrular.

    dosya : Dosya yolu, dosya nesnesi ya da bayt akışı (openpyxl.load_workbook
            ile uyumlu herhangi bir kaynak — örn. Flask'ın ``request.files``
            FileStorage nesnesi doğrudan verilebilir).
    sema  : ImportSemasi — şablonu üreten kod ile AYNI nesne.
    db    : Referans kontrolü yapan sütunlar için SQLAlchemy session
            (referans_kontrol callable'larına aynen iletilir).

    Dönüş: (hatalar, temiz_satirlar)
      hatalar        : [(satir_no, mesaj), ...] — boşsa dosya geçerlidir.
      temiz_satirlar : Hata yoksa, tip dönüştürülmüş satır sözlükleri listesi
                       (her biri alan_adi -> Python değeri).
    """
    try:
        wb = load_workbook(dosya, data_only=True)
    except Exception as exc:
        return [(0, f"Excel dosyası okunamadı: {exc}")], []

    if VERI_SAYFA_ADI not in wb.sheetnames:
        return [(0, f"'{VERI_SAYFA_ADI}' adlı sayfa bulunamadı. Lütfen resmi şablonu kullanın.")], []

    ws = wb[VERI_SAYFA_ADI]
    hatalar = []
    temiz_satirlar = []
    benzersizlik_gorulen = {s.alan_adi: {} for s in sema.sutunlar if s.benzersiz}

    max_satir = ws.max_row or 1
    for excel_satir in range(2, max_satir + 1):
        ham_degerler = [ws.cell(row=excel_satir, column=c).value for c in range(1, len(sema.sutunlar) + 1)]
        if all(v is None or str(v).strip() == "" for v in ham_degerler):
            continue  # Tamamen boş satırı atla (şablon sonundaki boş satırlar).

        satir_hatalari, temiz_satir = _satiri_dogrula(
            excel_satir, ham_degerler, sema, db, benzersizlik_gorulen
        )
        if satir_hatalari:
            hatalar.extend(satir_hatalari)
        else:
            temiz_satirlar.append(temiz_satir)

    if hatalar:
        return hatalar, []
    return [], temiz_satirlar


def _satiri_dogrula(excel_satir, ham_degerler, sema: ImportSemasi, db, benzersizlik_gorulen):
    hatalar = []
    temiz_satir = {}

    for sutun, ham in zip(sema.sutunlar, ham_degerler):
        deger = ham.strip() if isinstance(ham, str) else ham
        bos_mu = deger is None or (isinstance(deger, str) and deger == "")

        if bos_mu:
            if sutun.zorunlu:
                hatalar.append((excel_satir, f"'{sutun.baslik}' alanı zorunludur, boş bırakılamaz."))
                continue
            temiz_satir[sutun.alan_adi] = sutun.varsayilan
            continue

        donusmus, hata_mesaji = _degeri_donustur(deger, sutun)
        if hata_mesaji:
            hatalar.append((excel_satir, f"'{sutun.baslik}': {hata_mesaji}"))
            continue

        if sutun.tip == "metin" and sutun.maks_uzunluk and len(str(donusmus)) > sutun.maks_uzunluk:
            hatalar.append((
                excel_satir,
                f"'{sutun.baslik}' en fazla {sutun.maks_uzunluk} karakter olabilir "
                f"(girilen: {len(str(donusmus))})."
            ))
            continue

        if sutun.benzersiz:
            gorulen = benzersizlik_gorulen[sutun.alan_adi]
            anahtar = str(donusmus).strip().lower()
            if anahtar in gorulen:
                hatalar.append((
                    excel_satir,
                    f"'{sutun.baslik}' değeri '{donusmus}' {gorulen[anahtar]}. satırla aynı, dosya içinde tekrar edemez."
                ))
                continue
            gorulen[anahtar] = excel_satir

        if sutun.referans_kontrol is not None:
            ref_hata = sutun.referans_kontrol(db, donusmus)
            if ref_hata:
                hatalar.append((excel_satir, f"'{sutun.baslik}': {ref_hata}"))
                continue

        temiz_satir[sutun.alan_adi] = donusmus

    return hatalar, temiz_satir


def _degeri_donustur(deger, sutun: SutunTanimi):
    """Ham Excel hücre değerini şema tipine göre Python değerine çevirir.

    Dönüş: (donusturulmus_deger, hata_mesaji|None)
    """
    if sutun.tip == "metin":
        return str(deger).strip(), None

    if sutun.tip == "sayi":
        if isinstance(deger, bool):
            return None, "sayısal bir değer olmalıdır."
        if isinstance(deger, (int, float)):
            return float(deger), None
        try:
            return float(str(deger).strip().replace(",", ".")), None
        except (TypeError, ValueError):
            return None, f"sayısal bir değer olmalıdır (girilen: '{deger}')."

    if sutun.tip == "tarih":
        if isinstance(deger, datetime):
            return deger.date(), None
        if isinstance(deger, date):
            return deger, None
        if isinstance(deger, str):
            metin = deger.strip()
            for kalip in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y"):
                try:
                    return datetime.strptime(metin, kalip).date(), None
                except ValueError:
                    continue
        return None, f"tarih formatı tanınamadı (beklenen: YYYY-AA-GG, girilen: '{deger}')."

    if sutun.tip == "enum":
        metin = str(deger).strip()
        if sutun.enum_degerleri and metin not in sutun.enum_degerleri:
            secenekler = ", ".join(sutun.enum_degerleri)
            return None, f"geçersiz değer '{metin}'. İzin verilenler: {secenekler}."
        return metin, None

    if sutun.tip == "boolean":
        if isinstance(deger, bool):
            return deger, None
        metin = str(deger).strip().lower()
        if metin in ("evet", "e", "true", "1", "var"):
            return True, None
        if metin in ("hayır", "hayir", "h", "false", "0", "yok"):
            return False, None
        return None, f"Evet/Hayır bekleniyor (girilen: '{deger}')."

    return deger, None


# ═══════════════════════════════════════════════════════════════════════════
#  3) TEK TRANSACTION İÇE AKTARICI
# ═══════════════════════════════════════════════════════════════════════════
def aktar(db, model_sinifi, temiz_satirlar: list, ekstra_alanlar: Optional[dict] = None) -> int:
    """Doğrulanmış satırları TEK veritabanı transaction'ında modele yazar.

    Ya tüm satırlar başarıyla eklenir ve commit edilir, ya da herhangi bir
    hata durumunda rollback yapılıp HİÇBİR kayıt yazılmaz (all-or-nothing).

    db             : SQLAlchemy Session.
    model_sinifi   : Kayıtların oluşturulacağı ORM modeli (örn. Supplier).
    temiz_satirlar : ``dogrula()`` çıktısındaki temiz satır sözlükleri listesi.
    ekstra_alanlar : Her satıra eklenecek sabit alanlar (örn. olusturan_id).

    Dönüş: eklenen kayıt sayısı.
    Hata durumunda mevcut istisna yeniden fırlatılır (rollback sonrası).
    """
    if not temiz_satirlar:
        return 0

    nesneler = []
    for satir in temiz_satirlar:
        veriler = dict(satir)
        if ekstra_alanlar:
            veriler.update(ekstra_alanlar)
        nesneler.append(model_sinifi(**veriler))

    try:
        db.add_all(nesneler)
        db.commit()
    except Exception:
        db.rollback()
        raise

    return len(nesneler)
