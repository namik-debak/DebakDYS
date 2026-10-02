# -*- coding: utf-8 -*-
"""
Update talimat .xls and .doc files via Microsoft Office COM (single app instance).
"""
from __future__ import annotations

import sys
from pathlib import Path as _Path

_VENV_SITE = _Path(r"\\192.168.0.249\ortak\HARUN\dys\.venv\Lib\site-packages")
if _VENV_SITE.exists() and str(_VENV_SITE) not in sys.path:
    sys.path.insert(0, str(_VENV_SITE))

import re
import shutil
import sqlite3
import traceback
from pathlib import Path

import openpyxl
import win32com.client

from _talimat_meta_utils import (
    DATE_RE,
    DOC_CODE_RE,
    REV_PAREN_RE,
    build_meta_line_from_old,
    fmt_date,
    fmt_rev,
    norm,
    paragraph_has_meta,
    strip_excel_footer_codes,
)

DB = r"\\192.168.0.249\ortak\HARUN\dys\dys.db"
XLSX_LIST = r"c:\Users\labaratuvar03\Downloads\ana_dokuman_listesi2-.xlsx"
BASE = Path(r"D:\UYGULAMA PROJELER\debak dijital\dys")
LOG = BASE / "_talimat_com_log.txt"

WD_REPLACE_ALL = 2
WD_HEADER_FOOTER_PRIMARY = 1


def load_excel_map():
    wb = openpyxl.load_workbook(XLSX_LIST, data_only=True)
    ws = wb["Ana Doküman Listesi"]
    out = {}
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 3).value == "T":
            kod = ws.cell(r, 2).value
            if kod:
                out[norm(kod)] = {
                    "kod": str(kod).strip(),
                    "rev": ws.cell(r, 5).value,
                    "tarih": ws.cell(r, 6).value,
                }
    return out


def resolve_path(p):
    if not p:
        return None
    path = Path(p)
    if path.exists():
        return path
    p2 = p.replace("/", "\\")
    if "uploads\\documents" in p2:
        rel = p2.split("uploads\\documents", 1)[1].lstrip("\\/")
        c = BASE / "uploads" / "documents" / rel
        return c if c.exists() else None
    return None


def backup(path: Path):
    bak = path.with_suffix(path.suffix + ".bak_meta")
    if not bak.exists():
        shutil.copy2(path, bak)


def preserve_replace_range(rng, new_text: str):
    try:
        alignment = rng.ParagraphFormat.Alignment
    except Exception:
        alignment = None
    try:
        font = rng.Font
        name, size, bold = font.Name, font.Size, font.Bold
    except Exception:
        name = size = bold = None
    ends_cell = (rng.Text or "").endswith("\x07")
    rng.Text = new_text + ("\x07" if ends_cell else "\r")
    try:
        if alignment is not None:
            rng.ParagraphFormat.Alignment = alignment
        if name:
            rng.Font.Name = name
        if size:
            rng.Font.Size = size
        if bold is not None:
            rng.Font.Bold = bold
    except Exception:
        pass


def update_word_tables(story_range, kod, rev, tarih) -> bool:
    updated = False
    rev_s = fmt_rev(rev)
    tarih_s = fmt_date(tarih)
    try:
        tables = story_range.Tables
    except Exception:
        return False
    if tables.Count == 0:
        return False
    for ti in range(1, tables.Count + 1):
        tbl = tables(ti)
        for ri in range(1, tbl.Rows.Count + 1):
            for ci in range(1, tbl.Columns.Count + 1):
                try:
                    cell = tbl.Cell(ri, ci)
                except Exception:
                    continue
                txt = cell.Range.Text.replace("\x07", "").replace("\r", "").strip()
                if not txt:
                    continue
                if re.fullmatch(r"\d+", txt) and ci > 1:
                    left = tbl.Cell(ri, ci - 1).Range.Text.lower()
                    if "revizyon no" in left:
                        preserve_replace_range(cell.Range, rev_s)
                        updated = True
                        continue
                if DATE_RE.fullmatch(txt) and ci > 1:
                    left = tbl.Cell(ri, ci - 1).Range.Text.lower()
                    if "tarih" in left:
                        preserve_replace_range(cell.Range, tarih_s)
                        updated = True
                        continue
                if DOC_CODE_RE.search(txt) or (paragraph_has_meta(txt) and len(txt) < 80):
                    new = build_meta_line_from_old(kod, rev, tarih, txt)
                    if new != txt:
                        preserve_replace_range(cell.Range, new)
                        updated = True
    return updated


def update_word_story(story_range, kod, rev, tarih):
    raw = story_range.Text or ""
    plain = raw.replace("\x07", " ").replace("\r", " ").strip()
    if not plain:
        return False, None
    if story_range.Tables.Count > 0:
        ok = update_word_tables(story_range, kod, rev, tarih)
        return ok, None
    first_line = raw.split("\r")[0].replace("\x07", "").strip()
    if paragraph_has_meta(first_line) and len(first_line) < 120:
        new = build_meta_line_from_old(kod, rev, tarih, first_line)
        if new != first_line:
            preserve_replace_range(story_range, new)
            return True, new
    return False, None


def update_word_doc(word, path: Path, meta):
    updated = False
    last = None
    doc = word.Documents.Open(str(path), ReadOnly=False)
    try:
        for si in range(1, doc.Sections.Count + 1):
            sec = doc.Sections(si)
            for wi in range(1, 4):
                for getter in (lambda w: sec.Footers(w), lambda w: sec.Headers(w)):
                    try:
                        story = getter(wi)
                        if not story.Exists:
                            continue
                        ok, line = update_word_story(story.Range, meta["kod"], meta["rev"], meta["tarih"])
                        if ok:
                            updated = True
                            last = line or last
                    except Exception:
                        pass
        if not updated:
            footer = doc.Sections(1).Footers(WD_HEADER_FOOTER_PRIMARY)
            rng = footer.Range
            new = build_meta_line_from_old(meta["kod"], meta["rev"], meta["tarih"], "")
            rng.ParagraphFormat.Alignment = 2
            preserve_replace_range(rng, new)
            updated = True
            last = new
        doc.Save()
        return last
    finally:
        doc.Close()


def update_excel_page_part(ps, attr, kod, rev, tarih):
    val = getattr(ps, attr) or ""
    plain = strip_excel_footer_codes(val)
    if not paragraph_has_meta(plain):
        return False, None
    new_meta = build_meta_line_from_old(kod, rev, tarih, plain)
    if plain and plain in val:
        new_val = val.replace(plain, new_meta)
    else:
        new_val = new_meta
    setattr(ps, attr, new_val)
    return True, new_meta


def update_excel_xls(excel, path: Path, meta):
    updated = False
    last = None
    any_meta = False
    wb = excel.Workbooks.Open(str(path), ReadOnly=False)
    try:
        for ws in wb.Worksheets:
            ps = ws.PageSetup
            for attr in (
                "LeftHeader",
                "CenterHeader",
                "RightHeader",
                "LeftFooter",
                "CenterFooter",
                "RightFooter",
            ):
                ok, line = update_excel_page_part(ps, attr, meta["kod"], meta["rev"], meta["tarih"])
                if ok:
                    updated = True
                    any_meta = True
                    last = line
        if not any_meta:
            ps = wb.Worksheets(1).PageSetup
            new = build_meta_line_from_old(meta["kod"], meta["rev"], meta["tarih"], "")
            existing = ps.RightFooter or ""
            codes = re.findall(r"&[A-Za-z0-9+-]+", existing)
            ps.RightFooter = (existing + " " + new).strip() if codes else new
            updated = True
            last = new
        wb.Save()
        return last
    finally:
        wb.Close()


def main():
    excel_map = load_excel_map()
    conn = sqlite3.connect(DB)
    cur = conn.cursor()
    cur.execute(
        "SELECT dokuman_no, baslik, dosya_yolu, dosya_adi FROM documents WHERE dokuman_tipi='Talimat'"
    )
    xls_jobs = []
    doc_jobs = []
    for no, baslik, p, name in cur.fetchall():
        key = norm(no)
        if key not in excel_map:
            continue
        fp = resolve_path(p)
        if not fp or not fp.exists():
            continue
        ext = Path(name or fp.name).suffix.lower()
        if ext == ".xls":
            xls_jobs.append((no, fp, excel_map[key]))
        elif ext == ".doc":
            doc_jobs.append((no, fp, excel_map[key]))
    conn.close()

    lines = []
    xls_err = 0
    doc_err = 0
    xls_ok = 0
    doc_ok = 0

    print(f"XLS jobs: {len(xls_jobs)}, DOC jobs: {len(doc_jobs)}")
    excel_app = win32com.client.Dispatch("Excel.Application")
    excel_app.Visible = False
    excel_app.DisplayAlerts = False
    try:
        for i, (no, fp, meta) in enumerate(xls_jobs, 1):
            try:
                backup(fp)
                line = update_excel_xls(excel_app, fp, meta)
                lines.append(f"OK\t{no}\t{line}")
                xls_ok += 1
                if i % 25 == 0:
                    print(f"  xls {i}/{len(xls_jobs)}")
            except Exception as e:
                xls_err += 1
                lines.append(f"ERR\t{no}\t{e}")
    finally:
        excel_app.Quit()

    word_app = win32com.client.Dispatch("Word.Application")
    word_app.Visible = False
    word_app.DisplayAlerts = 0
    try:
        for i, (no, fp, meta) in enumerate(doc_jobs, 1):
            try:
                backup(fp)
                line = update_word_doc(word_app, fp, meta)
                lines.append(f"OK\t{no}\t{line}")
                doc_ok += 1
                if i % 10 == 0:
                    print(f"  doc {i}/{len(doc_jobs)}")
            except Exception as e:
                doc_err += 1
                lines.append(f"ERR\t{no}\t{e}")
    finally:
        word_app.Quit()

    summary = f"com xls_ok={xls_ok} xls_err={xls_err} doc_ok={doc_ok} doc_err={doc_err}"
    LOG.write_text(summary + "\n\n" + "\n".join(lines), encoding="utf-8")
    print(summary)
    print(f"Log: {LOG}")


if __name__ == "__main__":
    main()
