# -*- coding: utf-8 -*-
import sqlite3
from pathlib import Path
from docx import Document

DB = r"\\192.168.0.249\ortak\HARUN\dys\dys.db"
BASE = Path(r"\\192.168.0.249\ortak\HARUN\dys")


def resolve(p):
    if not p:
        return None
    path = Path(p)
    if path.exists():
        return path
    p2 = p.replace("/", "\\")
    if "uploads\\documents" in p2:
        rel = p2.split("uploads\\documents", 1)[1].lstrip("\\/")
        c = BASE / "uploads" / "documents" / rel
        return c if c.exists() else None
    return None


conn = sqlite3.connect(DB)
cur = conn.cursor()
cur.execute(
    "SELECT dokuman_no, dosya_yolu FROM documents WHERE dokuman_tipi='Talimat'"
)
patterns = {"empty": 0, "text_footer": 0, "text_header": 0, "table_footer": 0, "no_docx": 0, "missing": 0}
samples = []

for no, p in cur.fetchall():
    fp = resolve(p)
    if not fp:
        patterns["missing"] += 1
        continue
    if fp.suffix.lower() != ".docx":
        patterns["no_docx"] += 1
        continue
    doc = Document(str(fp))
    sec = doc.sections[0]
    ht = [x.text.strip() for x in sec.header.paragraphs if x.text.strip()]
    ft = [x.text.strip() for x in sec.footer.paragraphs if x.text.strip()]
    ftables = len(sec.footer.tables) > 0
    htables = len(sec.header.tables) > 0
    if ftables or htables:
        patterns["table_footer"] += 1
        if len(samples) < 4:
            samples.append((no, "tablo", ft, ht))
    elif ft:
        patterns["text_footer"] += 1
        if len(samples) < 8:
            samples.append((no, "alt_metin", ft, ht))
    elif ht:
        patterns["text_header"] += 1
    else:
        patterns["empty"] += 1

print("Format dagilimi (499 talimat):")
for k, v in patterns.items():
    print(f"  {k}: {v}")

print("\nOrnekler:")
for s in samples:
    print(f"--- {s[0]} ({s[1]})")
    print("  footer:", s[2][:3] if s[2] else "(bos)")
    print("  header:", s[3][:2] if s[3] else "(bos)")
