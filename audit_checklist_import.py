"""
Y03 denetim soru listeleri → AuditChecklist / Question aktarımı.
"""

from __future__ import annotations

import os
import re

from openpyxl import load_workbook

CHECKLIST_DIR = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    "static", "branding", "audit_checklists",
)

CHECKLIST_FILES = [
    ("Y03-F09-proses.xls", "Y03-F09", "Proses Denetim Soru Listesi", "Proses Denetimi", None),
    ("Y03-F22-urun.xlsx", "Y03-F22", "Ürün Denetim Soru Listesi", "Ürün Denetimi", None),
    ("Y03-F21-D01.xlsx", "Y03-F21-D01", "D01 İnsan Kaynakları / İSG", "Sistem Denetimi", "D01"),
    ("Y03-F24-D02.xlsx", "Y03-F24-D02", "D02 Altyapı ve Çalışma Ortamı", "Sistem Denetimi", "D02"),
    ("Y03-F25-D03.xlsx", "Y03-F25-D03", "D03 Bakım Onarım", "Sistem Denetimi", "D03"),
    ("Y03-F26-D04.xlsx", "Y03-F26-D04", "D04 Ölçme ve İzleme Ekipmanları", "Sistem Denetimi", "D04"),
    ("Y03-F27-D05.xlsx", "Y03-F27-D05", "D05 Satınalma Yönetimi", "Sistem Denetimi", "D05"),
    ("Y03-F28-D06.xlsx", "Y03-F28-D06", "D06 Kalite Kontrol", "Sistem Denetimi", "D06"),
    ("Y03-F29-D07.xlsx", "Y03-F29-D07", "D07 Çevre Yönetimi", "Sistem Denetimi", "D07"),
    ("Y03-F33-D08.xlsx", "Y03-F33-D08", "D08 Bilgi Güvenliği", "Sistem Denetimi", "D08"),
    ("Y03-F11-Y01.xlsx", "Y03-F11-Y01", "Y01 Kalite Yönetim Prosesi", "Sistem Denetimi", "Y01"),
    ("Y03-F14-Y02.xlsx", "Y03-F14-Y02", "Y02 Yönetimin Sorumluluğu", "Sistem Denetimi", "Y02"),
    ("Y03-F15-Y03.xlsx", "Y03-F15-Y03", "Y03 Sürekli İyileştirme", "Sistem Denetimi", "Y03"),
    ("Y03-F16-M01.xlsx", "Y03-F16-M01", "M01 Satış Projeleri Yönetimi", "Sistem Denetimi", "M01"),
    ("Y03-F17-M02.xlsx", "Y03-F17-M02", "M02 Yeni Ürün Devreye Alma", "Sistem Denetimi", "M02"),
    ("Y03-F18-M03.xlsx", "Y03-F18-M03", "M03 Üretim Planlama", "Sistem Denetimi", "M03"),
    ("Y03-F19-M04.xlsx", "Y03-F19-M04", "M04 Üretim Yönetimi", "Sistem Denetimi", "M04"),
    ("Y03-F20-M05.xlsx", "Y03-F20-M05", "M05 Sevkiyat Yönetimi", "Sistem Denetimi", "M05"),
]


def _col_index(header_vals, *names):
    lowered = [(i, str(v or "").strip().lower()) for i, v in enumerate(header_vals)]
    for name in names:
        n = name.lower()
        for i, v in lowered:
            if v == n or n in v:
                return i
    return None


def _parse_sistem_xlsx(path):
    wb = load_workbook(path, data_only=True)
    out = []
    sira = 0
    for sheet_name in wb.sheetnames:
        if sheet_name.strip().lower().startswith("sayfa"):
            continue
        ws = wb[sheet_name]
        header_r = None
        cols = {}
        for r in range(1, min(12, (ws.max_row or 0) + 1)):
            vals = [ws.cell(r, c).value for c in range(1, 16)]
            joined = " ".join(str(v).lower() for v in vals if v is not None)
            if ("soru" in joined) and (
                "no" in joined or "tip" in joined or "şart" in joined or "sart" in joined
                or "standart" in joined or "gözlem" in joined or "gozlem" in joined
            ):
                header_r = r
                cols["soru"] = _col_index(
                    vals, "Standart Maddeleri Göre Sorular", "Sorular", "SORU",
                )
                cols["no"] = _col_index(vals, "NO", "Soru No")
                if cols["no"] is None:
                    idx = _col_index(vals, "Soru")
                    if idx is not None and cols.get("soru") != idx:
                        cols["no"] = idx
                cols["sart"] = _col_index(
                    vals, "Şart No", "Sart No", "ŞART NO", "Şart", "Sart", "Kodu",
                )
                break
        if header_r is None or cols.get("soru") is None:
            continue
        start_r = header_r + 1
        peek = [ws.cell(start_r, c).value for c in range(1, 8)]
        peek_j = " ".join(str(v).lower() for v in peek if v is not None)
        if "kodu" in peek_j or peek_j.strip() in ("no", "evet"):
            start_r += 1
        for r in range(start_r, (ws.max_row or 0) + 1):
            soru = ws.cell(r, cols["soru"] + 1).value
            if soru is None or not str(soru).strip():
                continue
            text = str(soru).strip()
            if text.upper() in ("SORU", "EVET", "SORULAR") or len(text) < 12:
                continue
            no_val = ws.cell(r, cols["no"] + 1).value if cols.get("no") is not None else None
            sart = ws.cell(r, cols["sart"] + 1).value if cols.get("sart") is not None else None
            sira += 1
            if isinstance(no_val, (int, float)):
                soru_no = str(int(no_val))
            elif no_val:
                soru_no = str(no_val).strip()
            else:
                soru_no = str(sira)
            out.append({
                "sira": sira,
                "bolum": sheet_name.strip(),
                "soru_no": soru_no,
                "soru_metin": text,
                "sart_no": (str(sart).strip() if sart is not None else None) or None,
            })
    return out


def _parse_urun_xlsx(path):
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    out = []
    header_r = None
    for r in range(1, 15):
        vals = [ws.cell(r, c).value for c in range(1, 8)]
        joined = " ".join(str(v).lower() for v in vals if v is not None)
        if "soru" in joined and ("no" in joined or "karakteristik" in joined):
            header_r = r
            break
    if not header_r:
        return out
    last_kar = ""
    sira = 0
    for r in range(header_r + 1, (ws.max_row or 0) + 1):
        kar = ws.cell(r, 2).value
        if kar and str(kar).strip():
            last_kar = str(kar).strip()
        soru = ws.cell(r, 3).value
        if soru is None or not str(soru).strip():
            continue
        text = str(soru).strip()
        if len(text) < 5:
            continue
        no_val = ws.cell(r, 1).value
        sira += 1
        out.append({
            "sira": sira,
            "bolum": last_kar or "Genel",
            "soru_no": str(int(no_val)) if isinstance(no_val, (int, float)) else (str(no_val).strip() if no_val else str(sira)),
            "soru_metin": text,
            "sart_no": None,
        })
    return out


def _parse_proses_xls(path):
    import xlrd
    wb = xlrd.open_workbook(path)
    name = "Denetim Soru Listesi" if "Denetim Soru Listesi" in wb.sheet_names() else wb.sheet_names()[0]
    sh = wb.sheet_by_name(name)
    out = []
    bolum = "Genel"
    sira = 0
    for r in range(sh.nrows):
        c0 = sh.cell_value(r, 0)
        c1 = sh.cell_value(r, 1) if sh.ncols > 1 else ""
        s0 = str(c0).strip() if c0 is not None else ""
        s1 = str(c1).strip() if c1 is not None else ""
        if s0 and not s1 and (re.match(r"^[A-D]\)", s0) or (")" in s0[:3] and len(s0) > 3)):
            bolum = s0[:120]
            continue
        if not s1 or len(s1) < 10:
            continue
        is_no = isinstance(c0, (int, float)) or (isinstance(c0, str) and re.match(r"^\d+(\.\d+)?$", s0))
        if not is_no:
            continue
        sira += 1
        no = str(int(c0)) if isinstance(c0, float) and c0 == int(c0) else s0
        out.append({
            "sira": sira,
            "bolum": bolum,
            "soru_no": no,
            "soru_metin": s1,
            "sart_no": None,
        })
    return out


def parse_checklist_file(filename):
    path = os.path.join(CHECKLIST_DIR, filename)
    if not os.path.isfile(path):
        raise FileNotFoundError(path)
    low = filename.lower()
    if low.endswith(".xls") and not low.endswith(".xlsx"):
        return _parse_proses_xls(path)
    if "f22" in low or "urun" in low:
        return _parse_urun_xlsx(path)
    return _parse_sistem_xlsx(path)


def import_all_checklists(db, replace=False):
    from models import AuditChecklist, AuditChecklistQuestion

    cl_count = 0
    q_count = 0
    for fname, kod, ad, tip, surec in CHECKLIST_FILES:
        path = os.path.join(CHECKLIST_DIR, fname)
        if not os.path.isfile(path):
            continue
        questions = parse_checklist_file(fname)
        if not questions:
            continue
        existing = db.query(AuditChecklist).filter_by(kod=kod).first()
        if existing:
            if not replace:
                continue
            db.delete(existing)
            db.flush()
        cl = AuditChecklist(
            kod=kod, ad=ad, denetim_tipi=tip, surec_kod=surec,
            kaynak_dosya=fname, aktif=True,
        )
        db.add(cl)
        db.flush()
        for q in questions:
            db.add(AuditChecklistQuestion(
                checklist_id=cl.id, sira=q["sira"], bolum=q["bolum"],
                soru_no=q["soru_no"], soru_metin=q["soru_metin"], sart_no=q["sart_no"],
            ))
            q_count += 1
        cl_count += 1
    db.commit()
    return cl_count, q_count


def find_checklist_for_audit(db, denetim_tipi, surec_kod=None):
    from models import AuditChecklist
    q = db.query(AuditChecklist).filter_by(aktif=True, denetim_tipi=denetim_tipi)
    if denetim_tipi == "Sistem Denetimi" and surec_kod:
        cl = q.filter_by(surec_kod=surec_kod).first()
        if cl:
            return cl
    return q.filter(AuditChecklist.surec_kod.is_(None)).first() or q.first()
