# -*- coding: utf-8 -*-
import sqlite3
import openpyxl
from pathlib import Path

DB = r"\\192.168.0.249\ortak\HARUN\dys\dys.db"
XLSX = r"c:\Users\labaratuvar03\Downloads\ana_dokuman_listesi2-.xlsx"

conn = sqlite3.connect(DB)
cur = conn.cursor()

# table info
cur.execute("SELECT name FROM sqlite_master WHERE type='table'")
print("Tables:", [r[0] for r in cur.fetchall()])

# find document table columns
for tbl in ["documents", "dokumanlar", "document"]:
    cur.execute("SELECT name FROM sqlite_master WHERE type='table' AND name=?", (tbl,))
    if cur.fetchone():
        cur.execute(f"PRAGMA table_info({tbl})")
        cols = [r[1] for r in cur.fetchall()]
        print(f"\n{tbl} columns:", cols)
        cur.execute(f"SELECT COUNT(1) FROM {tbl}")
        print(f"  total rows: {cur.fetchone()[0]}")

# try documents table
cur.execute("PRAGMA table_info(documents)")
cols = [r[1] for r in cur.fetchall()]
print("\ndocuments columns:", cols)

# talimat query - adapt column names
tip_col = "dokuman_tipi" if "dokuman_tipi" in cols else "tip"
cur.execute(f"SELECT COUNT(1) FROM documents WHERE {tip_col}='Talimat'")
print(f"\nTalimat count in DB: {cur.fetchone()[0]}")

select_cols = [c for c in ["id", "dokuman_no", "baslik", "revizyon_no", "yururluk_tarihi", "dosya_yolu", "dosya_adi", "icerik_json"] if c in cols]
if select_cols:
    cur.execute(f"SELECT {', '.join(select_cols)} FROM documents WHERE {tip_col}='Talimat' ORDER BY dokuman_no LIMIT 15")
    print("\nSample talimatlar:")
    for r in cur.fetchall():
        print(r)

# all talimat titles from DB
cur.execute(f"SELECT dokuman_no, baslik FROM documents WHERE {tip_col}='Talimat' ORDER BY dokuman_no")
db_talimatlar = {row[0]: row[1] for row in cur.fetchall() if row[0]}
print(f"\nDB talimat count with kod: {len(db_talimatlar)}")

# Excel talimatlar
wb = openpyxl.load_workbook(XLSX, data_only=True)
ws = wb["Ana Doküman Listesi"]
excel_talimatlar = {}
for r in range(2, ws.max_row + 1):
    if ws.cell(r, 3).value == "T":
        kod = ws.cell(r, 2).value
        tanim = ws.cell(r, 4).value
        rev = ws.cell(r, 5).value
        tarih = ws.cell(r, 6).value
        if kod:
            excel_talimatlar[kod.strip()] = {"tanim": tanim, "rev": rev, "tarih": tarih}

print(f"\nExcel talimat count: {len(excel_talimatlar)}")

# Match comparison
db_kodlar = set(db_talimatlar.keys())
excel_kodlar = set(excel_talimatlar.keys())

only_db = sorted(db_kodlar - excel_kodlar)
only_excel = sorted(excel_kodlar - db_kodlar)
common = sorted(db_kodlar & excel_kodlar)

print(f"\nCommon codes: {len(common)}")
print(f"Only in DB (portal): {len(only_db)}")
if only_db[:20]:
    for k in only_db[:20]:
        print(f"  DB only: {k} -> {db_talimatlar[k]}")
    if len(only_db) > 20:
        print(f"  ... and {len(only_db)-20} more")

print(f"\nOnly in Excel: {len(only_excel)}")
if only_excel[:20]:
    for k in only_excel[:20]:
        print(f"  Excel only: {k} -> {excel_talimatlar[k]['tanim']}")
    if len(only_excel) > 20:
        print(f"  ... and {len(only_excel)-20} more")

# Title mismatches for common codes
mismatches = []
for k in common:
    db_title = (db_talimatlar.get(k) or "").strip()
    ex_title = (excel_talimatlar[k]["tanim"] or "").strip()
    if db_title.lower() != ex_title.lower():
        mismatches.append((k, db_title, ex_title))

print(f"\nTitle mismatches (same kod, different title): {len(mismatches)}")
for k, db_t, ex_t in mismatches[:15]:
    print(f"  {k}:")
    print(f"    Portal: {db_t}")
    print(f"    Excel:  {ex_t}")

conn.close()
