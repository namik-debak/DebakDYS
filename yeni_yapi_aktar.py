"""
DYS — 11 süreçlik yeni yapıyı veritabanına aktarma aracı (2026)
================================================================
Varsayılan: DENEME TURU (veritabanına yazmaz, planı raporlar).

  python yeni_yapi_aktar.py --paket "<YENİ YAPI\\DYS Aktarım klasörü>" [--uygula] [--eski-pasif]
        [--dosya-koku "\\\\192.168.0.249\\kalite\\YENİ YAPI"]

--uygula      : değişiklikleri yazar (öncesinde veritabanı yedeği alınmış olmalı)
--eski-pasif  : eski süreçleri pasif yapar; yerine yenisi gelen eski dokümanları 'Eskimiş' yapar, yerine_gecen_id ile bağlar
--dosya-koku  : dokümanların dosya_yolu için YENİ YAPI klasörünün DYS sunucusundan erişilen yolu

Kurallar (hiçbir kayıt SİLİNMEZ; tekrar çalıştırmak güvenlidir):
- Kodu yeni yapıda BAŞKA sürece geçen eski süreç (Y01, M03, M04, D02, D05, D06): eski kayıt 'XXX-E' koduna alınır ve
  pasif olur, eski dokümanları onun altında kalır; yeni süreç yeni kayıtla açılır.
- Numarası yeni yapıda BAŞKA dokümana geçen eski doküman: '<no> (ESKİ)' olarak yeniden numaralanır (eski no açıklamada saklanır).
- Kodu değişmeyen doküman (ör. M01 P01): kopyası açılmaz; mevcut kayıt yeni sürece bağlanır.
"""
import argparse
import csv
import glob
import json
import os
import re
import sys
from datetime import datetime

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from models import SessionLocal, init_db, Process, Document, FormTanim  # noqa: E402

YENI_SURECLER = [  # kod, ad, kategori, üst
    ("Y01", "Yönetim Sorumluluğu ve Sürekli İyileştirme", "Yönetim", None),
    ("Y01.1", "Finansal Kaynaklar", "Yönetim", "Y01"), ("Y01.2", "Strateji Yönetimi", "Yönetim", "Y01"),
    ("Y01.3", "Risk Yönetimi", "Yönetim", "Y01"), ("Y01.4", "İletişim", "Yönetim", "Y01"),
    ("Y01.5", "Yönetim Gözden Geçirmeleri", "Yönetim", "Y01"), ("Y01.6", "Denetimlerin Yönetimi", "Yönetim", "Y01"),
    ("Y01.7", "Düzeltici Faaliyetler", "Yönetim", "Y01"), ("Y01.8", "Performans Değerlendirme", "Yönetim", "Y01"),
    ("M01", "Satış Projeleri Yönetimi", "Ana", None), ("M02", "Yeni Ürün Devreye Alma Yönetimi (YÜDA)", "Ana", None),
    ("M03", "Üretim Yönetimi", "Ana", None), ("M04", "Planlama ve Sevkiyat Yönetimi", "Ana", None),
    ("M04.1", "Üretim Planlama", "Ana", "M04"), ("M04.2", "Sevkiyat Yönetimi", "Ana", "M04"),
    ("D01", "İnsan Kaynakları", "Destek", None), ("D01.1", "Yetkinlik ve Eğitim Yönetimi", "Destek", "D01"),
    ("D01.2", "İş Sağlığı ve Güvenliği Yönetimi", "Destek", "D01"),
    ("D02", "Satın Alma ve Tedarikçi Yönetimi", "Destek", None), ("D03", "Bakım ve Ekipman Yönetimi", "Destek", None),
    ("D03.1", "Yatırım Yönetimi", "Destek", "D03"), ("D03.2", "Avadanlık Yönetimi", "Destek", "D03"),
    ("D04", "Kalite Yönetimi", "Destek", None), ("D04.1", "Girdi Kontrolleri", "Destek", "D04"),
    ("D04.2", "İmalat Kontrolleri", "Destek", "D04"), ("D04.3", "Final Kontrolleri", "Destek", "D04"),
    ("D04.4", "Kalite Laboratuvarı Yönetimi", "Destek", "D04"), ("D04.5", "Kalite Kontrol (Genel)", "Destek", "D04"),
    ("D04.6", "Kalite Yönetim Sistemi", "Destek", "D04"), ("D05", "Çevre Yönetimi", "Destek", None),
    ("D06", "Bilgi Güvenliği ve Kişisel Verilerin Korunması Yönetimi", "Destek", None),
    ("D06.1", "Kişisel Verilerin Korunması", "Destek", "D06"),   # D01.3 KVKK birleşti (30.09.2026)
]
YENI_KODLAR = {k for k, *_ in YENI_SURECLER}
ANLAMI_DEGISEN = {"Y01", "M03", "M04", "D02", "D05", "D06"}  # aynı kod, yeni yapıda farklı süreç
SEVIYE = {"Prosedür": 2, "İş Akışı": 2, "Talimat": 3, "Görev Tanımı": 3, "Form": 4, "El Kitabı": 1, "Politika": 1, "Liste": 4}


def rows(path, sheet):
    ws = openpyxl.load_workbook(path, read_only=True)[sheet]
    it = ws.iter_rows(values_only=True)
    hdr = next(it)
    return [dict(zip(hdr, r)) for r in it]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--paket", required=True)
    ap.add_argument("--uygula", action="store_true")
    ap.add_argument("--eski-pasif", action="store_true")
    ap.add_argument("--dosya-koku", default="")
    a = ap.parse_args()
    init_db()
    db = SessionLocal()
    R = {"surec_eklendi": 0, "surec_arsivlendi": 0, "surec_guncellendi": 0, "dokuman_eklendi": 0, "dokuman_ayni_kod": 0,
         "dokuman_yeniden_numaralandi": 0, "eskimis": 0, "form_tanim": 0, "uyari": []}
    try:
        # ---- eski kod → yeni kod eşlemesi (tüm ana liste) ----
        eski_yeni = {}
        csvp = os.path.join(a.paket, "kod_eslestirme.csv")
        if os.path.exists(csvp):
            for r in csv.reader(open(csvp, encoding="utf-8-sig"), delimiter=";"):
                if len(r) >= 5 and r[0] != "eski_kod" and r[4]:
                    eski_yeni[r[0]] = r[4]
        for r in rows(os.path.join(a.paket, "Doküman_Yerlesim_Listesi.xlsx"), "DYS Doküman Kayıtları"):
            if r.get("Eski Kod") and r.get("dokuman_no"):
                eski_yeni.setdefault(r["Eski Kod"], r["dokuman_no"])
        akt = rows(os.path.join(a.paket, "DYS_Aktarim_Listesi.xlsx"), "DYS Aktarım")
        for r in akt:
            m = re.search(r"yerine_gecen:\s*([^·]+)", str(r.get("İlişkili doküman / yerine_gecen") or ""))
            if r.get("İşlem") == "İptal" and m and not str(r["dokuman_no"]).startswith("("):
                eski_yeni[r["dokuman_no"]] = m.group(1).strip()

        # ---- 1) Süreçler ----
        mevcut = {p.kod: p for p in db.query(Process).all()}
        eski_surec = {}  # eski kod -> eski kaydı (arşivlenmiş ya da aynı)
        for kod, ad, kat, ust in YENI_SURECLER:
            p = mevcut.get(kod)
            if p is not None and kod in ANLAMI_DEGISEN and p.ad != ad and not (p.aciklama or "").startswith("11 süreçlik yapı"):
                arsiv = f"{kod}-E"
                R["uyari"].append(f"Süreç {kod} '{p.ad}' → arşiv kodu {arsiv} (pasif); yeni {kod} '{ad}' açılıyor")
                p.kod = arsiv; p.ad = f"(Eski) {p.ad}"; p.aktif = False
                eski_surec[kod] = p; mevcut[arsiv] = p; p = None; R["surec_arsivlendi"] += 1
            if p is None:
                p = Process(kod=kod, ad=ad, kategori=kat, aktif=True, aciklama="11 süreçlik yapı (2026)"); db.add(p); mevcut[kod] = p
                R["surec_eklendi"] += 1
            elif p.ad != ad or p.kategori != kat:
                R["uyari"].append(f"Süreç {kod}: ad '{p.ad}' → '{ad}'")
                p.ad = ad; p.kategori = kat; p.aktif = True; R["surec_guncellendi"] += 1
        db.flush()
        for kod, ad, kat, ust in YENI_SURECLER:
            mevcut[kod].ust_surec_id = mevcut[ust].id if ust else None
        if a.eski_pasif:
            for p in db.query(Process).all():
                if p.kod not in YENI_KODLAR:
                    p.aktif = False
        db.flush()

        def surec_of(no):
            k = (no or "").split()[0]
            return mevcut.get(k) if k in YENI_KODLAR else (mevcut.get(k.split(".")[0]) if k.split(".")[0] in YENI_KODLAR else None)

        # ---- 2) Yeni doküman listesi ----
        yeni = []
        for r in akt:
            if r.get("İşlem") != "Yeni":
                continue
            tip = "İş Akışı" if str(r.get("dokuman_tipi") or "").startswith("İş Akışı") else "Prosedür"
            icerik = None
            if tip == "Prosedür":
                jp = os.path.join(a.paket, f"{r['dokuman_no']}.json")
                if os.path.exists(jp):
                    icerik = json.dumps(json.load(open(jp, encoding="utf-8")).get("icerik_json") or {}, ensure_ascii=False)
            else:  # İş Akışı: akış adımları DYS'de şema olarak gösterilir
                jp = os.path.join(a.paket, "İş Akışı Tanımları", f"{r['dokuman_no']}.json")
                if os.path.exists(jp):
                    s = json.load(open(jp, encoding="utf-8"))
                    icerik = json.dumps({"amac": s.get("amac", ""), "ilgili_prosedur": s.get("prosedur", ""),
                                         "akis": json.dumps(s.get("bolumler") or [], ensure_ascii=False),
                                         "notlar": "\n".join(f"{n.get('md', '')} {n['x']}".strip() for n in s.get("notlar") or [])},
                                        ensure_ascii=False)
            dosya = re.sub(r"\s+\(sayfa: .*\)$", "", str(r.get("Dosya (YENİ YAPI altında)") or ""))
            yeni.append((r["dokuman_no"], r["baslik"], tip, r.get("ilgili_standartlar") or "", dosya, icerik))
        for r in rows(os.path.join(a.paket, "Doküman_Yerlesim_Listesi.xlsx"), "DYS Doküman Kayıtları"):
            dosya = str(r.get("Dosya (YENİ YAPI)") or "").split("YENİ YAPI\\", 1)[-1]
            yeni.append((r["dokuman_no"], r["baslik"], r["dokuman_tipi"], "", dosya, None))
        # DYS'ye aktarılmayacak dokümanlar (kullanıcı kararı): haric_tutulanlar.json
        haric = {}
        hp = os.path.join(a.paket, "haric_tutulanlar.json")
        if os.path.exists(hp):
            haric = json.load(open(hp, encoding="utf-8")).get("dokumanlar") or {}
        if haric:
            yeni = [y for y in yeni if y[0] not in haric]
            R["haric"] = sorted(haric)
        yeni_nolar = {n for n, *_ in yeni}

        # ---- 3) Çakışan eski dokümanlar: aynı kod → koru; başka dokümana geçen kod → (ESKİ) ----
        docs = {d.dokuman_no: d for d in db.query(Document).all()}
        for no in list(docs):
            if no not in yeni_nolar:
                continue
            hedef = eski_yeni.get(no)
            sk = docs[no].surec.kod if docs[no].surec else ""
            eski_yapida = sk.endswith("-E") or sk not in YENI_KODLAR
            if not eski_yapida or hedef == no:
                continue  # yeni yapıdaki doküman ya da kodu değişmeyen aynı doküman → korunur
            d = docs.pop(no)
            d.dokuman_no = f"{no} (ESKİ)"[:30]
            d.iptal_nedeni = f"11 süreçlik yapıya geçiş (2026): eski numara {no}; yeni karşılığı {hedef or '—'}"
            docs[d.dokuman_no] = d; R["dokuman_yeniden_numaralandi"] += 1
        db.flush()

        # ---- 4) Yeni dokümanları ekle / aynı koddakileri yeni sürece bağla ----
        for no, baslik, tip, std, dosya, icerik in yeni:
            s = surec_of(no)
            if not s:
                R["uyari"].append(f"{no}: süreç bulunamadı"); continue
            if no in docs:
                d = docs[no]
                if d.surec_id != s.id:
                    d.surec_id = s.id
                if icerik and (not d.icerik_json or tip == "İş Akışı"):
                    d.icerik_json = icerik
                R["dokuman_ayni_kod"] += 1
                continue
            d = Document(dokuman_no=no, baslik=(baslik or no)[:200], surec_id=s.id, dokuman_tipi=tip,
                         dokuman_seviyesi=SEVIYE.get(tip, 2), guvenlik_sinifi="Genel", ilgili_standartlar=std,
                         durum="Taslak", revizyon_no=0, icerik_json=icerik,
                         dosya_adi=os.path.basename(dosya)[:200] if dosya else None,
                         dosya_yolu=os.path.join(a.dosya_koku, dosya) if (a.dosya_koku and dosya) else None,
                         olusturma_tarihi=datetime.now(), guncelleme_tarihi=datetime.now())
            db.add(d); docs[no] = d; R["dokuman_eklendi"] += 1
        db.flush()

        # ---- 5) Eskiyen dokümanlar: yerine geçen yeni dokümana bağla ----
        if a.eski_pasif:
            for no, d in list(docs.items()):
                if no in yeni_nolar:
                    continue  # yeni yapının geçerli dokümanı (anlamı değişen süreçte aynı numarayla açılmış olabilir) — eskitilmez
                eski_no = no[:-7] if no.endswith(" (ESKİ)") else no
                hedef_no = eski_yeni.get(eski_no)
                if not hedef_no or hedef_no == no:
                    continue
                hedef = docs.get(hedef_no)
                if hedef and hedef.id != d.id and d.durum not in ("Eskimiş", "İptal"):
                    d.durum = "Eskimiş"; d.yerine_gecen_id = hedef.id
                    d.iptal_nedeni = d.iptal_nedeni or f"11 süreçlik yapıya geçiş (2026): yerine {hedef_no}"
                    d.iptal_tarihi = datetime.now().date(); R["eskimis"] += 1

        # ---- 6) Dinamik form tanımları ----
        for jp in glob.glob(os.path.join(a.paket, "Dinamik Form Taslakları", "*.json")):
            j = json.load(open(jp, encoding="utf-8"))
            kod = j.get("form_kodu")
            if not kod or kod in haric or db.query(FormTanim).filter(FormTanim.form_kodu == kod).first():
                continue
            s = surec_of(kod); d = docs.get(kod)
            db.add(FormTanim(form_kodu=kod, eski_kod=j.get("eski_kod"), ad=(j.get("form_adi") or kod)[:200],
                             surec_id=s.id if s else None, document_id=d.id if d else None,
                             alanlar_json=json.dumps(j.get("alanlar") or [], ensure_ascii=False),
                             onay_akisi=",".join(j.get("onay_akisi") or []), aktif=True,
                             erisim_json=json.dumps(j["erisim"], ensure_ascii=False) if j.get("erisim") else None))   # KVKK kısıtlı form
            db.flush()
            R["form_tanim"] += 1
        if a.uygula:
            db.commit(); print("UYGULANDI.")
            # ---- 7) İç denetim: 11 süreç soru listesi + F06 planı + denetçi yetkinlikleri (kendi commit'i) ----
            den = os.path.join(a.paket, "Denetim")
            if os.path.isdir(den):
                import denetim_otomasyon
                R["denetim"] = denetim_otomasyon.paketten_yukle(db, den)
            # ---- 8) M03 F08 Ürün Enjeksiyon Parametreleri: ürünler + Excel ayar geçmişi ----
            ejp = os.path.join(a.paket, "Enjeksiyon", "enjeksiyon_parametreleri.json")
            if os.path.exists(ejp):
                from enjeksiyon_veri import paketten_yukle as enj_yukle
                R["enjeksiyon"] = enj_yukle(db, json.load(open(ejp, encoding="utf-8")))
            # ---- 9) Stratejik Planlama (Y01 F02 / F03 → modül): yıllık bağlam, paydaş, SWOT, stratejiler ----
            import glob as _glob
            stp = sorted(_glob.glob(os.path.join(a.paket, "Strateji", "strateji_*.json")))
            if stp:
                import strateji_veri
                R["strateji"] = [strateji_veri.paketten_yukle(db, p) for p in stp]
            # ---- 10) Performans Değerlendirme ve CAPA (kpi.debak.com → DYS): süreç, KPI, ölçüm, NCR (NCR'ler DÖF'e yansıtılır) ----
            kpp = os.path.join(a.paket, "Performans", "kpi_state.json")
            if os.path.exists(kpp):
                import pk_veri
                R["performans"] = pk_veri.aktar(db, kpp)
        else:
            db.rollback(); print("DENEME TURU — veritabanına yazılmadı. Uygulamak için --uygula ekleyin.")
        print(json.dumps({k: v for k, v in R.items() if k != "uyari"}, ensure_ascii=False))
        for u in R["uyari"][:60]:
            print("  !", u)
    finally:
        db.close()


if __name__ == "__main__":
    main()
